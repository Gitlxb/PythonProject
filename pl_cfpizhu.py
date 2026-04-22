# -- coding: utf-8 --
# @Time : 2025-03-22 14:33
# @Author : 贝特利
# @Email : 1356087739@qq.com
# @File : pl_cfpizhu.py
# @Software: PyCharm

import os
from openpyxl import load_workbook
from openpyxl.comments import Comment
from tkinter import Tk, messagebox
from tkinter.filedialog import askopenfilename

# 隐藏Tkinter根窗口
Tk().withdraw()

# 选择要打开的Excel文件
file_path = askopenfilename(title="选择Excel文件", filetypes=[("Excel files", "*.xlsx")])

if not file_path:
    print("未选择文件，程序退出。")
    exit()

# 加载工作簿
wb = load_workbook(file_path)
ws = wb.active

# 遍历A列的单元格
for row in ws.iter_rows(min_col=1, max_col=1):
    for cell in row:
        # 检查单元格是否有批注
        if cell.comment is not None:
            comment_text = cell.comment.text
            # 过滤掉以"Admin:"开头的行
            filtered_lines = [line for line in comment_text.split('\n') if not line.strip().startswith("Admin:")]
            # 将过滤后的内容合并为一个字符串
            filtered_text = " ".join(filtered_lines)
            # 将多个连续空格替换为单个空格，并按空格拆分成列表
            split_values = " ".join(filtered_text.split()).split(" ")
            # 将拆分后的内容写入右侧的单元格
            for i, value in enumerate(split_values):
                ws.cell(row=cell.row, column=cell.column + i + 1, value=value)

# 直接保存到源文件
wb.save(file_path)

# 提示用户操作完成
messagebox.showinfo("完成", "批量拆分批注完成！")
print("批量拆分批注完成！")
print(f"文件已保存到原路径: {file_path}")