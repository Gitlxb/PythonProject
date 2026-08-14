#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
工人预支申请 与 预支人员信息查询明细表 —— 字段匹配工具

功能：
  1. 用户分别选择「预支人员信息查询明细表」(输出数据) 与「工人预支申请」(匹配对照) 两个 xlsx；
  2. 匹配键 = 员工工号（单键）；同一工号允许多行（多对多），全部保留；
  3. 过滤“员工工号”为空白的行；
  4. 对匹配记录逐行比对 5 个字段：
     还款状态、预支状态、预支金额、部分还款金额、剩余未还金额；
  5. 生成新 Excel，列顺序：
     还款状态、预支状态、预支金额、部分还款金额、剩余未还金额、
     员工工号、姓名、匹配状态、备注
  6. 汇总弹窗 + 另存为选择保存位置。

匹配状态规则：
  - 同工号下 5 字段全部相等 -> TRUE，备注为空
  - 同工号下 5 字段无全等   -> FALSE，备注写 “字段名不一致”(多个用、连接)
  - 工号在申请表中不存在    -> 未匹配，备注写 “仅存在于预支人员信息查询明细表”
"""

import os
import re
import sys
import subprocess
try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
    _TkBase = tk.Tk
except ImportError:
    tk = None   # 核心逻辑仍可用；仅 GUI 无法启动
    _TkBase = object  # 占位基类，避免无 tkinter 时类定义报错

try:
    import openpyxl
except ImportError:
    # 尝试自动安装，失败则提示用户手动安装
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
        import openpyxl
    except Exception:
        openpyxl = None


# 用于比对的 5 个字段
MATCH_FIELDS = [
    "还款状态", "预支状态", "预支金额",
    "部分还款金额", "剩余未还金额",
]
# 输出列 = 5 个匹配字段 + 员工工号 + 姓名（不含所属月份）
OUTPUT_FIELDS = MATCH_FIELDS + ["员工工号", "姓名"]
# 需要按数值比对的字段
NUMERIC_FIELDS = {"预支金额", "部分还款金额", "剩余未还金额"}


# ------------------------- 工具函数 -------------------------
def normalize_header(h):
    """表头去全部空白（含首尾空格、全/半角空格），用于字段定位。"""
    if h is None:
        return ""
    return re.sub(r"\s+", "", str(h))


def norm_val(v):
    """单元格值去空白后再比较，避免空格导致误判。"""
    if v is None:
        return ""
    return re.sub(r"\s+", "", str(v))


def to_number(v):
    """把金额等字段归一化为 float；无法转换返回 None。"""
    if v is None:
        return None
    s = str(v).strip()
    if s == "":
        return None
    s = re.sub(r"[¥￥$€,\s]", "", s)  # 去货币符号、千分位、空白
    try:
        return float(s)
    except ValueError:
        return None


def load_table(path):
    """
    读取一个 xlsx。
    匹配键 = 员工工号(标准化) 单键；同一工号允许多行（多对多），全部保留。
    过滤“员工工号”为空白的行。
    返回 (groups, skipped)：
      groups: {工号: [行字典, ...]}，行字典含 OUTPUT_FIELDS 的 7 个字段
      skipped: 被过滤的“员工工号为空白”行数
    """
    if openpyxl is None:
        raise RuntimeError("缺少依赖 openpyxl，请先执行: pip install openpyxl")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"文件为空: {os.path.basename(path)}")

    header = [normalize_header(h) for h in rows[0]]
    if "员工工号" not in header:
        raise ValueError(f"未找到“员工工号”列: {os.path.basename(path)}")

    field_cols = {f: header.index(f) for f in OUTPUT_FIELDS if f in header}
    missing = [f for f in OUTPUT_FIELDS if f not in field_cols]
    if missing:
        raise ValueError(f"{os.path.basename(path)} 缺少字段: {', '.join(missing)}")
    emp_idx = field_cols["员工工号"]

    groups = {}
    skipped = 0
    for r in rows[1:]:
        emp = norm_val(r[emp_idx])
        if emp == "":
            skipped += 1          # 过滤“员工工号”为空白的行
            continue
        rec = {f: r[field_cols[f]] for f in OUTPUT_FIELDS}
        groups.setdefault(emp, []).append(rec)
    return groups, skipped


def field_differs(f, v_base, v_compared):
    """判断某字段在输出表(预支人员信息查询明细表)与对照表之间是否不一致。"""
    if f in NUMERIC_FIELDS:
        n1, n2 = to_number(v_base), to_number(v_compared)
        if n1 is None or n2 is None:
            return norm_val(v_base) != norm_val(v_compared)
        return abs(n1 - n2) > 1e-6
    return norm_val(v_base) != norm_val(v_compared)


def rows_equal(base, compared):
    """判断两行 5 个匹配字段是否全部相等，返回 (是否全等, 不一致字段列表)。"""
    diffs = []
    for f in MATCH_FIELDS:
        if field_differs(f, base.get(f), compared.get(f)):
            diffs.append(f)
    return (len(diffs) == 0, diffs)


def process(base_groups, compared_groups):
    """
    核心匹配逻辑（工号单键 + 多对多）。
    base_groups     = 预支人员信息查询明细表（输出数据，逐行写入 Excel）
    compared_groups = 工人预支申请（匹配对照）
    遍历 base 的每一行，在 compared 中找同工号的行比对 5 字段。
    返回 (out_rows, (matched, mismatched, unmatched_base, unmatched_compared))
    """
    out_rows = []
    matched = mismatched = unmatched_base = 0
    unmatched_compared = 0   # 仅存在于申请表：仅汇总提示，不写入 Excel

    for emp in sorted(base_groups):
        b_list = base_groups[emp]
        c_list = compared_groups.get(emp, [])

        if not c_list:
            # 工号在申请表中不存在 -> 未匹配
            for b in b_list:
                out_rows.append([b.get(f, "") for f in OUTPUT_FIELDS]
                                + ["未匹配", "仅存在于预支人员信息查询明细表"])
                unmatched_base += 1
        else:
            # 输出表每一行都与申请表同工号行比对：任一 5 字段全等即 TRUE
            for b in b_list:
                equal_any = False
                first_diff = None
                for c in c_list:
                    eq, diffs = rows_equal(b, c)
                    if eq:
                        equal_any = True
                        break
                    if first_diff is None:
                        first_diff = diffs
                if equal_any:
                    status, remark = "TRUE", ""
                    matched += 1
                else:
                    status = "FALSE"
                    remark = "、".join(f + "不一致" for f in (first_diff or []))
                    mismatched += 1
                out_rows.append([b.get(f, "") for f in OUTPUT_FIELDS]
                                + [status, remark])

    # 申请表独有的工号（仅汇总提示，不写入 Excel）
    for emp in compared_groups:
        if emp not in base_groups:
            unmatched_compared += len(compared_groups[emp])

    stats = (matched, mismatched, unmatched_base, unmatched_compared)
    return out_rows, stats


def write_output(path, out_rows):
    """写出结果 Excel。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "匹配结果"
    header = OUTPUT_FIELDS + ["匹配状态", "备注"]
    ws.append(header)
    for row in out_rows:
        ws.append(row)
    for i in range(1, len(header) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 18
    wb.save(path)


# ------------------------- GUI -------------------------
class MatchApp(_TkBase):
    def __init__(self):
        super().__init__()
        self.title("预支申请字段匹配工具")
        self.geometry("880x360")
        self.minsize(720, 320)
        # 内部保存全路径，界面只显示文件名
        self._base_full = ""        # 输出数据 = 预支人员信息查询明细表
        self._compared_full = ""    # 匹配对照 = 工人预支申请
        self.base_path = tk.StringVar(value="（未选择文件）")
        self.compared_path = tk.StringVar(value="（未选择文件）")
        self.status_var = tk.StringVar(value="● 就绪")
        self._build()

    def _build(self):
        # 列权重：中间列(文件路径)自动伸展
        self.columnconfigure(0, weight=0, minsize=180)
        self.columnconfigure(1, weight=1)
        self.columnconfigure(2, weight=0, minsize=120)

        # —— 标题 ——
        tk.Label(self, text="预支申请字段匹配工具",
                 font=("Microsoft YaHei", 13, "bold"),
                 fg="#1a237e").grid(row=0, column=0, columnspan=3,
                                    sticky="w", padx=18, pady=(14, 0))
        tk.Label(self, text="分别选择两份 Excel，以员工工号为键匹配 5 个字段",
                 fg="#666666").grid(row=1, column=0, columnspan=3,
                                    sticky="w", padx=18, pady=(0, 10))

        # —— 文件选择区（分组）——
        group = tk.LabelFrame(self, text="  文件选择  ",
                              font=("Microsoft YaHei", 9),
                              fg="#1a237e", bd=1, relief="groove",
                              labelanchor="nw")
        group.grid(row=2, column=0, columnspan=3,
                   sticky="ew", padx=18, pady=4)
        group.columnconfigure(1, weight=1)

        def add_row(r, label_text, path_var, choose_cmd):
            tk.Label(group, text=label_text, anchor="w"
                     ).grid(row=r, column=0, sticky="w",
                            padx=(12, 6), pady=10)
            tk.Entry(group, textvariable=path_var, state="readonly",
                     readonlybackground="#fafafa",
                     font=("Consolas", 10)
                     ).grid(row=r, column=1, sticky="ew", padx=4, pady=10)
            tk.Button(group, text="选择文件…", width=12,
                      command=choose_cmd
                      ).grid(row=r, column=2, padx=(4, 12), pady=10)

        add_row(0, "预支人员信息查询明细表（输出数据）",
                self.base_path, self.select_base)
        add_row(1, "工人预支申请（匹配对照）",
                self.compared_path, self.select_compared)

        # —— 操作按钮 ——
        action = tk.Frame(self)
        action.grid(row=3, column=0, columnspan=3,
                    sticky="ew", padx=18, pady=(14, 6))
        action.columnconfigure(0, weight=1)
        tk.Button(action, text="开 始 处 理",
                  font=("Microsoft YaHei", 11, "bold"),
                  bg="#2e7d32", fg="white",
                  activebackground="#1b5e20", activeforeground="white",
                  relief="flat", cursor="hand2",
                  height=2, width=18,
                  command=self.run).grid(row=0, column=0, sticky="e")

        # —— 状态栏 ——
        tk.Label(self, textvariable=self.status_var, anchor="w",
                 fg="#555555", font=("Microsoft YaHei", 9)
                 ).grid(row=4, column=0, columnspan=3,
                        sticky="ew", padx=18, pady=(6, 0))

        tk.Label(self,
                 text="说明：以员工工号为键；比对 5 个字段（还款状态、预支状态、"
                      "预支金额、部分还款金额、剩余未还金额）；工号空白行将被过滤。",
                 anchor="w", fg="#888888"
                 ).grid(row=5, column=0, columnspan=3,
                        sticky="ew", padx=18, pady=(0, 12))

    def select_base(self):
        p = filedialog.askopenfilename(
            title="选择预支人员信息查询明细表（输出数据）",
            filetypes=[("Excel 文件", "*.xlsx *.xlsm")])
        if p:
            self._base_full = p
            self.base_path.set("📄  " + os.path.basename(p))

    def select_compared(self):
        p = filedialog.askopenfilename(
            title="选择工人预支申请（匹配对照）",
            filetypes=[("Excel 文件", "*.xlsx *.xlsm")])
        if p:
            self._compared_full = p
            self.compared_path.set("📄  " + os.path.basename(p))

    def run(self):
        if not self._base_full or not self._compared_full:
            messagebox.showwarning("提示", "请先选择两个 Excel 文件。")
            return

        self.status_var.set("● 处理中…")
        self.update_idletasks()

        try:
            base_groups, b_skip = load_table(self._base_full)
            compared_groups, c_skip = load_table(self._compared_full)
        except Exception as e:
            self.status_var.set("● 读取失败")
            messagebox.showerror("读取失败", str(e))
            return

        out_rows, (matched, mismatched, unmatched_base, unmatched_compared) = process(
            base_groups, compared_groups)

        summary = (
            f"处理完成：\n"
            f"  完全匹配 (TRUE)：{matched} 条\n"
            f"  存在差异 (FALSE)：{mismatched} 条\n"
            f"  仅存在于预支人员信息查询明细表：{unmatched_base} 条\n"
            f"  仅存在于工人预支申请（仅汇总提示）：{unmatched_compared} 条\n"
            f"  过滤“员工工号空白”行：明细表 {b_skip} 条 / 申请表 {c_skip} 条\n"
            f"  写入 Excel 合计：{len(out_rows)} 条"
        )
        self.status_var.set(f"● 已处理：合计 {len(out_rows)} 条")
        messagebox.showinfo("处理结果", summary)

        save_p = filedialog.asksaveasfilename(
            title="选择保存位置",
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx")],
            initialfile="预支匹配结果.xlsx")
        if not save_p:
            return
        try:
            write_output(save_p, out_rows)
        except Exception as e:
            self.status_var.set("● 保存失败")
            messagebox.showerror("保存失败", str(e))
            return
        self.status_var.set(f"● 已保存：{os.path.basename(save_p)}")
        messagebox.showinfo("完成", f"结果已保存至：\n{save_p}")


def main():
    if tk is None:
        sys.stderr.write(
            "缺少 tkinter，GUI 无法启动。请使用标准 Windows Python "
            "(通常自带 tkinter) 运行本程序。\n")
        return
    if openpyxl is None:
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror(
            "缺少依赖",
            "未能导入 openpyxl。请先执行：\n\npip install openpyxl\n\n后重新运行本程序。")
        return
    app = MatchApp()
    app.mainloop()


if __name__ == "__main__":
    main()
