# -- coding: utf-8 --
# @Time : 2025-03-22 10:33
# @Author : 贝特利
# @Email : 1356087739@qq.com
# @File : pl_pizhu.py
# @Software: PyCharm

import os
from openpyxl import load_workbook
from openpyxl.comments import Comment
from tkinter import Tk
from tkinter.filedialog import askopenfilename
from tkinter.messagebox import showinfo

def add_comments_to_excel():
    # 隐藏Tkinter根窗口
    Tk().withdraw()

    # 选择Excel文件
    file_path = askopenfilename(title="选择Excel文件", filetypes=[("Excel files", "*.xlsx")])

    if not file_path:
        print("未选择文件，程序退出。")
        exit()

    try:
        # 加载工作簿
        wb = load_workbook(filename=file_path)
        ws = wb.active

        # 遍历A列所有有内容的单元格
        for row in ws.iter_rows(min_col=1, max_col=1):  # 只遍历A列
            a_cell = row[0]  # A列单元格
            if a_cell.value:  # 如果A列单元格有内容
                # 获取当前行的B列到K列内容
                start_col = 2  # B列
                end_col = 11  # K列
                merged_content = []
                for col in range(start_col, end_col + 1):
                    cell_value = ws.cell(row=a_cell.row, column=col).value
                    if cell_value:  # 如果单元格有内容
                        merged_content.append(str(cell_value))

                # 用换行符连接内容
                comment_text = "\n".join(merged_content)

                # 添加批注
                if a_cell.comment:  # 如果已经有批注，先删除
                    a_cell.comment = None
                if comment_text:  # 如果有内容才添加批注
                    a_cell.comment = Comment(comment_text, "Author")

        # 保存工作簿
        wb.save(file_path)
        print(f"批注已成功添加到 {os.path.basename(file_path)} 的A列有内容的单元格中。")

        # 提示完成
        showinfo("完成", "批量批注完成！")

    except Exception as e:
        print(f"发生错误: {e}")
        showinfo("错误", f"处理过程中发生错误: {e}")

if __name__ == "__main__":
    add_comments_to_excel()