# -- coding: utf-8 --
# @Time : 2025-08-11 15:20
# @Author : 贝特利
# @Email : 1356087739@qq.com
# @File : cw_gys_qzy.py
# @Software: PyCharm

def run_excel_merger_qzy():
    """运行Excel合并工具"""
    try:
        import os
        import tkinter as tk
        from tkinter import filedialog, messagebox, ttk
        import pandas as pd
        from openpyxl import load_workbook
        import threading

        # 创建新的Toplevel窗口
        merger_window = tk.Toplevel()
        merger_window.title("Excel文件合并工具")
        merger_window.geometry("500x350")

        # 变量初始化
        folder_path = tk.StringVar()
        sheet_name = tk.StringVar()
        sheet_names = []

        # 创建控件
        def create_widgets():
            # 文件夹选择
            folder_frame = tk.Frame(merger_window)
            folder_frame.pack(pady=10, padx=10, fill="x")

            tk.Label(folder_frame, text="文件夹路径:").pack(anchor="w")
            path_frame = tk.Frame(folder_frame)
            path_frame.pack(fill="x", pady=5)

            tk.Entry(path_frame, textvariable=folder_path, state="readonly").pack(side="left", fill="x", expand=True)
            tk.Button(path_frame, text="浏览", command=browse_folder).pack(side="right", padx=(5, 0))

            # 工作表选择
            sheet_frame = tk.Frame(merger_window)
            sheet_frame.pack(pady=10, padx=10, fill="x")

            tk.Label(sheet_frame, text="工作表名称:").pack(anchor="w")
            sheet_combo = ttk.Combobox(sheet_frame, textvariable=sheet_name, state="readonly")
            sheet_combo.pack(fill="x", pady=5)

            # 合并按钮
            button_frame = tk.Frame(merger_window)
            button_frame.pack(pady=20)

            tk.Button(button_frame, text="合并Excel文件", command=start_merge,
                     bg="#4CAF50", fg="white", font=("Arial", 12, "bold")).pack(pady=10)

            # 进度条
            progress = ttk.Progressbar(merger_window, mode="indeterminate")
            progress.pack(fill="x", padx=10, pady=5)

            # 状态标签
            status_label = tk.Label(merger_window, text="准备就绪", relief="sunken", anchor="w")
            status_label.pack(fill="x", padx=10, pady=5)

            return sheet_combo, progress, status_label

        def browse_folder():
            folder_selected = filedialog.askdirectory()
            if folder_selected:
                folder_path.set(folder_selected)
                scan_excel_files()

        def scan_excel_files():
            folder = folder_path.get()
            if not folder:
                return

            # 查找所有Excel文件
            excel_files = [f for f in os.listdir(folder) if f.endswith(('.xlsx', '.xls'))]
            if not excel_files:
                messagebox.showwarning("警告", "选择的文件夹中没有Excel文件!")
                return

            # 获取所有工作表的名称
            nonlocal sheet_names
            sheet_names = []
            sample_file = os.path.join(folder, excel_files[0])
            try:
                wb = load_workbook(sample_file, read_only=True)
                sheet_names = wb.sheetnames
                wb.close()
            except Exception as e:
                messagebox.showerror("错误", f"无法读取Excel文件: {str(e)}")
                return

            # 更新下拉框
            sheet_combo['values'] = sheet_names
            if sheet_names:
                sheet_name.set(sheet_names[0])

        def start_merge():
            if not folder_path.get():
                messagebox.showwarning("警告", "请先选择文件夹!")
                return

            if not sheet_name.get():
                messagebox.showwarning("警告", "请选择工作表名称!")
                return

            # 在新线程中执行合并操作
            thread = threading.Thread(target=merge_excel_files)
            thread.daemon = True
            thread.start()

        def merge_excel_files():
            progress.start()
            status_label.config(text="正在合并文件...")

            try:
                folder = folder_path.get()
                sheet = sheet_name.get()

                # 获取所有Excel文件
                excel_files = [f for f in os.listdir(folder) if f.endswith(('.xlsx', '.xls'))]

                # 合并数据
                merged_data = []
                for i, file in enumerate(excel_files):
                    status_label.config(text=f"正在处理 {file} ({i + 1}/{len(excel_files)})")
                    merger_window.update()

                    file_path = os.path.join(folder, file)
                    try:
                        # 读取Excel文件，跳过第1行（从第2行开始读取）
                        df = pd.read_excel(file_path, sheet_name=sheet, header=None, skiprows=1)
                        if not df.empty:
                            # 添加文件名列以便追踪数据来源
                            df['来源文件'] = file
                            merged_data.append(df)
                    except Exception as e:
                        print(f"读取文件 {file} 时出错: {str(e)}")

                if not merged_data:
                    messagebox.showwarning("警告", "没有找到有效数据!")
                    return

                # 合并所有数据（按列字母顺序合并）
                result = pd.concat(merged_data, ignore_index=True)

                # 数据处理
                status_label.config(text="正在进行数据处理...")
                merger_window.update()

                # 1. 删除D列中"供应商-"的内容（假设D列是第4列，索引为3）
                if len(result.columns) > 3:  # 确保有D列
                    result[3] = result[3].astype(str).str.replace('供应商-', '')

                # 2. 删除A列中空白的内容所在的行（假设A列是第1列，索引为0）
                if len(result.columns) > 0:  # 确保有A列
                    result = result[result[0].notna() & (result[0].astype(str).str.strip() != '')]

                # 3. 保留A列第一行的"序号"，删除其余行中的"序号"
                if len(result.columns) > 0:  # 确保有A列
                    # 获取第一行的值
                    first_row_value = str(result.iloc[0, 0]) if len(result) > 0 else ""
                    # 保留第一行，删除其他行中A列包含"序号"的行
                    if "序号" in first_row_value:
                        mask = (result.index == 0) | (~result[0].astype(str).str.contains("序号"))
                        result = result[mask]

                # 保存结果
                save_path = os.path.join(folder, "合并结果.xlsx")
                result.to_excel(save_path, index=False, header=False)

                status_label.config(text=f"合并完成! 结果已保存到: {save_path}")
                messagebox.showinfo("完成", f"Excel文件合并完成!\n结果已保存到: {save_path}")

            except Exception as e:
                messagebox.showerror("错误", f"合并过程中出现错误: {str(e)}")
                status_label.config(text="合并失败")
            finally:
                progress.stop()

        # 创建控件并获取引用
        sheet_combo, progress, status_label = create_widgets()

    except Exception as e:
        messagebox.showerror("错误", f"启动Excel合并工具时出错: {str(e)}")