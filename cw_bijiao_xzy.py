# -- coding: utf-8 --
# @Time : 2026-05-08
# @Author : 贝特利
# @Email : 1356087739@qq.com
# @File : cw_bijiao_xzy.py
# @Software: PyCharm
# @Description: 工人预支申请表 vs 预支明细表 对比工具 - 核心功能模块

import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import os
import warnings

warnings.filterwarnings('ignore')


class WorkerAdvanceComparatorCore:
    """工人预支对比工具 - 核心功能类"""

    def __init__(self):
        # 第一个表格（工人预支申请表）
        self.file1_path = None
        self.sheet1_name = None
        self.df_raw1 = None
        self.df1_processed = None
        self.header_row1 = None

        # 第二个表格（预支明细表）
        self.file2_path = None
        self.sheet2_name = None
        self.df_raw2 = None
        self.df2_processed = None
        self.header_row2 = None

        # 表1日期范围（工人预支申请表）
        self.start_date1 = None
        self.end_date1 = None

        # 表2日期范围（预支明细表）
        self.start_date2 = None
        self.end_date2 = None

        # 通用
        self.final_df = None

        # 预解析的中间数据（供分步处理使用）
        self._table1_parsed_data = None
        self._table2_parsed_data = None

    def set_file1(self, file_path, sheet_name):
        """设置第一个文件路径和工作表"""
        self.file1_path = file_path
        self.sheet1_name = sheet_name

    def set_file2(self, file_path, sheet_name):
        """设置第二个文件路径和工作表"""
        self.file2_path = file_path
        self.sheet2_name = sheet_name

    def set_date_range1(self, start_date, end_date):
        """设置表1（工人预支申请表）的日期范围"""
        self.start_date1 = pd.to_datetime(start_date)
        self.end_date1 = pd.to_datetime(end_date)
        self.end_date1 = self.end_date1 + pd.Timedelta(days=1) - pd.Timedelta(seconds=1)

    def set_date_range2(self, start_date, end_date):
        """设置表2（预支明细表）的日期范围"""
        self.start_date2 = pd.to_datetime(start_date)
        self.end_date2 = pd.to_datetime(end_date)
        self.end_date2 = self.end_date2 + pd.Timedelta(days=1) - pd.Timedelta(seconds=1)

    @staticmethod
    def _parse_date_flexible(series):
        """更灵活地解析日期，兼容多种格式，统一转为日期（不含时间）"""
        def _try_convert(val):
            if pd.isna(val) or str(val).strip() in ('', 'nan', 'NaT', 'None'):
                return pd.NaT
            str_val_orig = str(val).strip()
            try:
                # 尝试直接解析
                result = pd.to_datetime(val, errors='coerce')
                if pd.notna(result):
                    return result.normalize()
            except Exception:
                pass
            try:
                # 尝试 dayfirst
                result = pd.to_datetime(val, errors='coerce', dayfirst=True)
                if pd.notna(result):
                    return result.normalize()
            except Exception:
                pass
            try:
                # 尝试解析 Excel 日期序列号（数字）
                num_val = float(val)
                if 10000 < num_val < 100000:
                    from datetime import datetime as dt_datetime, timedelta
                    return pd.Timestamp(dt_datetime(1899, 12, 30) + timedelta(days=int(num_val))).normalize()
            except (ValueError, TypeError):
                pass
            # 尝试常见中文日期格式（2026年5月8日 → 2026-5-8）
            str_val = str_val_orig.replace('年', '-').replace('月', '-').replace('日', '')
            try:
                result = pd.to_datetime(str_val, errors='coerce')
                if pd.notna(result):
                    return result.normalize()
            except Exception:
                pass
            # 尝试处理 2026/5/8 格式（斜杠分隔，可能是单数字月/日）
            if '/' in str_val_orig:
                try:
                    # 尝试多种格式
                    for fmt in ['%Y/%m/%d %H:%M:%S', '%Y/%m/%d', '%Y/%m/%d %H:%M']:
                        try:
                            from datetime import datetime
                            result = datetime.strptime(str_val_orig, fmt)
                            return pd.Timestamp(result).normalize()
                        except ValueError:
                            continue
                except Exception:
                    pass
            return pd.NaT

        return series.apply(_try_convert)

    @staticmethod
    def find_header_row(df, required_cols, search_rows=10):
        """
        在DataFrame的前几行中查找表头行
        返回：(header_row_index, column_mapping)
        column_mapping: {标准列名: 实际列索引}
        """
        df_str = df.astype(str)
        for row_idx in range(min(search_rows, len(df))):
            row_values = df_str.iloc[row_idx].tolist()
            found_cols = {}
            for std_col, keywords in required_cols.items():
                for col_idx, val in enumerate(row_values):
                    if any(kw in val for kw in keywords):
                        found_cols[std_col] = col_idx
                        break
            if len(found_cols) >= 3:
                return row_idx, found_cols
        return None, None

    def _load_table1_core(self):
        """读取表1原始数据并解析列，返回解析后的DataFrame（不做日期筛选）"""
        if not self.file1_path or not self.sheet1_name:
            raise ValueError("请先设置第一个文件路径和工作表")

        self.df_raw1 = pd.read_excel(
            self.file1_path, sheet_name=self.sheet1_name, header=None
        )

        required_cols = {
            '姓名': ['姓名', '名字', '员工姓名', '工人姓名'],
            '预支金额': ['预支金额', '预支数额', '预支', '预支工资', '金额'],
            '身份证号': ['身份证', '身份证号', '身份证号码', '证件号', '证件号码'],
            '提交时间': ['提交时间', '提交日期', '申请时间', '申请日期', '日期', '时间']
        }

        header_row, col_mapping = self.find_header_row(self.df_raw1, required_cols)

        if header_row is None:
            raise ValueError(
                "在第一个表格（工人预支申请表）中未找到完整的表头信息。\n"
                "需要包含：姓名、预支金额、身份证号、提交时间"
            )

        self.header_row1 = header_row

        missing_cols = [k for k in required_cols.keys() if k not in col_mapping]
        if missing_cols:
            raise ValueError(f"第一个表格中缺少以下列: {', '.join(missing_cols)}")

        data_df = self.df_raw1.iloc[header_row + 1:].copy()
        data_df.columns = list(self.df_raw1.columns)
        data_df = data_df.reset_index(drop=True)

        name_col_idx = col_mapping['姓名']
        amount_col_idx = col_mapping['预支金额']
        id_col_idx = col_mapping['身份证号']
        date_col_idx = col_mapping['提交时间']

        result_df = pd.DataFrame()
        result_df['姓名'] = data_df.iloc[:, name_col_idx]
        result_df['预支金额'] = pd.to_numeric(data_df.iloc[:, amount_col_idx], errors='coerce')
        result_df['身份证号'] = data_df.iloc[:, id_col_idx].astype(str).str.strip()
        result_df['提交时间'] = self._parse_date_flexible(data_df.iloc[:, date_col_idx])

        result_df = result_df.dropna(subset=['姓名', '预支金额'])
        result_df = result_df[result_df['姓名'].astype(str).str.strip() != '']
        result_df = result_df[result_df['姓名'].astype(str).str.strip() != 'nan']
        result_df = result_df.reset_index(drop=True)
        return result_df

    def get_date_range_table1(self):
        """预加载表1，返回可用日期范围 (min_date_str, max_date_str)"""
        result_df = self._load_table1_core()
        self._table1_parsed_data = result_df

        valid_dates = result_df['提交时间'].dropna()
        if valid_dates.empty:
            return "无有效日期", "无有效日期"
        return valid_dates.min().strftime("%Y-%m-%d"), valid_dates.max().strftime("%Y-%m-%d")

    def process_first_table(self):
        """基于已存储的表1数据和日期范围，执行筛选、汇总"""
        if self._table1_parsed_data is None:
            self._table1_parsed_data = self._load_table1_core()

        result_df = self._table1_parsed_data.copy()
        total_rows = len(result_df)
        valid_dates = result_df['提交时间'].notna().sum()
        na_dates = total_rows - valid_dates

        mask = (result_df['提交时间'] >= self.start_date1) & (result_df['提交时间'] <= self.end_date1)
        result_df = result_df[mask].copy()

        if result_df.empty:
            msg = "在第一个表格中，所选时间范围内没有数据。\n"
            if total_rows > 0:
                msg += f"原因诊断：共读取 {total_rows} 行数据，其中 {valid_dates} 行成功识别日期，{na_dates} 行日期解析失败。\n"
                msg += f"筛选范围：{self.start_date1} ~ {self.end_date1}"
            raise ValueError(msg)

        result_df['身份证号'] = result_df['身份证号'].replace('nan', '').replace('NaN', '').replace('None', '')
        result_df['身份证号'] = result_df.groupby('姓名')['身份证号'].transform(
            lambda x: x.replace('', method='ffill').replace('', method='bfill')
        )

        self.df1_processed = result_df.groupby(['姓名', '身份证号'], as_index=False)['预支金额'].sum()
        self.df1_processed = self.df1_processed[['姓名', '预支金额', '身份证号']]
        return self.df1_processed

    def _load_table2_core(self):
        """读取表2原始数据并解析列，返回解析后的DataFrame（不做日期筛选）"""
        if not self.file2_path or not self.sheet2_name:
            raise ValueError("请先设置第二个文件路径和工作表")

        self.df_raw2 = pd.read_excel(
            self.file2_path, sheet_name=self.sheet2_name, header=None
        )

        required_cols = {
            '姓名': ['姓名', '名字', '员工姓名', '工人姓名'],
            '预支数额': ['预支数额', '预支金额', '预支', '预支工资', '数额', '金额'],
            '预支日期': ['预支日期', '预支时间', '日期', '申请日期', '申请时间', '提交日期']
        }

        header_row, col_mapping = self.find_header_row(self.df_raw2, required_cols)

        if header_row is None:
            raise ValueError(
                "在第二个表格（预支明细表）中未找到完整的表头信息。\n"
                "需要包含：姓名、预支数额、预支日期"
            )

        self.header_row2 = header_row

        missing_cols = [k for k in required_cols.keys() if k not in col_mapping]
        if missing_cols:
            raise ValueError(f"第二个表格中缺少以下列: {', '.join(missing_cols)}")

        data_df = self.df_raw2.iloc[header_row + 1:].copy()
        data_df.columns = list(self.df_raw2.columns)
        data_df = data_df.reset_index(drop=True)

        name_col_idx = col_mapping['姓名']
        amount_col_idx = col_mapping['预支数额']
        date_col_idx = col_mapping['预支日期']

        result_df = pd.DataFrame()
        result_df['姓名'] = data_df.iloc[:, name_col_idx]
        result_df['预支数额'] = pd.to_numeric(data_df.iloc[:, amount_col_idx], errors='coerce')
        result_df['预支日期'] = self._parse_date_flexible(data_df.iloc[:, date_col_idx])

        result_df = result_df.dropna(subset=['姓名', '预支数额'])
        result_df = result_df[result_df['姓名'].astype(str).str.strip() != '']
        result_df = result_df[result_df['姓名'].astype(str).str.strip() != 'nan']
        result_df = result_df.reset_index(drop=True)
        return result_df

    def get_date_range_table2(self):
        """预加载表2，返回可用日期范围 (min_date_str, max_date_str)"""
        result_df = self._load_table2_core()
        self._table2_parsed_data = result_df

        valid_dates = result_df['预支日期'].dropna()
        if valid_dates.empty:
            return "无有效日期", "无有效日期"
        return valid_dates.min().strftime("%Y-%m-%d"), valid_dates.max().strftime("%Y-%m-%d")

    def process_second_table(self):
        """基于已存储的表2数据和日期范围，执行筛选、汇总"""
        if self._table2_parsed_data is None:
            self._table2_parsed_data = self._load_table2_core()

        result_df = self._table2_parsed_data.copy()
        total_rows = len(result_df)
        valid_dates = result_df['预支日期'].notna().sum()
        na_dates = total_rows - valid_dates

        mask = (result_df['预支日期'] >= self.start_date2) & (result_df['预支日期'] <= self.end_date2)
        result_df = result_df[mask].copy()

        if result_df.empty:
            msg = "在第二个表格中，所选时间范围内没有数据。\n"
            if total_rows > 0:
                msg += f"原因诊断：共读取 {total_rows} 行数据，其中 {valid_dates} 行成功识别日期，{na_dates} 行日期解析失败。\n"
                msg += f"筛选范围：{self.start_date2} ~ {self.end_date2}"
            raise ValueError(msg)

        self.df2_processed = result_df.groupby(['姓名'], as_index=False)['预支数额'].sum()
        return self.df2_processed

    def merge_results(self):
        """合并两个表格的结果并进行对比，返回最终的DataFrame"""
        if self.df1_processed is None or self.df2_processed is None:
            raise ValueError("请先处理两个表格的数据")

        result_rows = []
        right_map = dict(zip(self.df2_processed['姓名'], self.df2_processed['预支数额']))

        for _, left_row in self.df1_processed.iterrows():
            left_name = left_row['姓名']
            left_amount = left_row['预支金额']
            left_id = left_row['身份证号']

            if left_name in right_map:
                right_amount = right_map[left_name]
                if abs(left_amount - right_amount) < 0.01:
                    compare_result = "TRUE"
                    diff = 0.0
                else:
                    compare_result = "FALSE"
                    diff = left_amount - right_amount

                result_rows.append({
                    '姓名': left_name,
                    '预支金额': left_amount,
                    '身份证号': left_id,
                    '比较结果': compare_result,
                    '差额': round(diff, 2),
                    '姓名(右)': left_name,
                    '预支数额(右)': right_amount
                })
            else:
                result_rows.append({
                    '姓名': left_name,
                    '预支金额': left_amount,
                    '身份证号': left_id,
                    '比较结果': 'FALSE',
                    '差额': round(left_amount, 2),
                    '姓名(右)': '',
                    '预支数额(右)': 0
                })

        left_names = set(self.df1_processed['姓名'])
        for _, right_row in self.df2_processed.iterrows():
            right_name = right_row['姓名']
            if right_name not in left_names:
                result_rows.append({
                    '姓名': right_name,
                    '预支金额': 0,
                    '身份证号': '',
                    '比较结果': 'FALSE',
                    '差额': round(-right_row['预支数额'], 2),
                    '姓名(右)': right_name,
                    '预支数额(右)': right_row['预支数额']
                })

        self.final_df = pd.DataFrame(result_rows)

        if not self.final_df.empty:
            self.final_df = self.final_df.sort_values(
                by=['姓名', '预支金额'], ascending=[True, False]
            ).reset_index(drop=True)

        return self.final_df

    @staticmethod
    def format_number(value):
        """格式化数字"""
        try:
            num = float(value)
            if num == int(num):
                return str(int(num))
            return f"{num:.2f}"
        except (ValueError, TypeError):
            return str(value)

    @staticmethod
    def save_to_excel(final_df, save_path):
        """将结果保存到Excel文件"""
        if final_df is None or final_df.empty:
            raise ValueError("没有数据可以保存")

        wb = Workbook()
        ws = wb.active
        ws.title = "对比结果"

        # 样式定义
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        true_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        false_fill = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")

        data_font = Font(name="微软雅黑", size=10)
        data_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        # 表头
        headers = ['姓名', '预支金额', '身份证号', '比较结果', '差额', '姓名(右)', '预支数额(右)']
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border

        # 数据行
        for row_idx, (_, row_data) in enumerate(final_df.iterrows(), 2):
            values = [
                str(row_data['姓名']),
                float(row_data['预支金额']),
                str(row_data['身份证号']),
                str(row_data['比较结果']),
                float(row_data['差额']),
                str(row_data['姓名(右)']),
                float(row_data['预支数额(右)'])
            ]

            is_true = row_data['比较结果'] == 'TRUE'

            for col_idx, value in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = data_font
                cell.alignment = data_alignment
                cell.border = thin_border
                cell.fill = true_fill if is_true else false_fill

        # 列宽
        col_widths_map = {1: 14, 2: 14, 3: 22, 4: 12, 5: 12, 6: 14, 7: 14}
        for col_idx, width in col_widths_map.items():
            ws.column_dimensions[chr(64 + col_idx)].width = width

        # 冻结首行
        ws.freeze_panes = "A2"

        # 统计信息
        stat_row = len(final_df) + 3
        ws.cell(row=stat_row, column=1, value="统计信息").font = Font(name="微软雅黑", size=10, bold=True)

        true_count = len(final_df[final_df['比较结果'] == 'TRUE'])
        false_count = len(final_df[final_df['比较结果'] == 'FALSE'])

        ws.cell(row=stat_row + 1, column=1, value="相符（TRUE）").font = data_font
        ws.cell(row=stat_row + 1, column=2, value=true_count).font = data_font
        ws.cell(row=stat_row + 2, column=1, value="不符（FALSE）").font = data_font
        ws.cell(row=stat_row + 2, column=2, value=false_count).font = data_font
        ws.cell(row=stat_row + 3, column=1, value="总计").font = data_font
        ws.cell(row=stat_row + 3, column=2, value=len(final_df)).font = data_font

        wb.save(save_path)
        wb.close()


def get_sheet_names(file_path):
    """获取Excel文件的所有工作表名称"""
    from openpyxl import load_workbook
    wb = load_workbook(file_path, read_only=True)
    sheet_names = wb.sheetnames
    wb.close()
    return sheet_names
