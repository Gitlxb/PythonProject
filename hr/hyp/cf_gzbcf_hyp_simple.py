import os
import re
import shutil
import subprocess
import sys
import tkinter as tk
from copy import copy
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook


class SplitGroupDialog(tk.Toplevel):
    """选择拆分组合弹窗"""

    def __init__(self, parent, all_values):
        super().__init__(parent)
        self.title("选择拆分组合")
        self.geometry("680x420")
        self.transient(parent)
        self.grab_set()
        self.resizable(False, False)

        self.all_values = list(all_values)
        self.groups = []  # 已创建的组合，每个元素是成员列表
        self.result = None

        self._build_ui()

    def _build_ui(self):
        tk.Label(
            self,
            text="在左侧多选值，点击『添加到新组合』将其归为一组（输出一个文件）\n"
                  "未选择的值将各自单独拆分。",
            font=("Microsoft YaHei", 9), fg="#555555", justify="left"
        ).pack(anchor="w", padx=15, pady=(12, 8))

        mid_frame = tk.Frame(self)
        mid_frame.pack(fill="both", expand=True, padx=15)

        # ---- 左侧：可选值 ----
        left = tk.Frame(mid_frame)
        left.pack(side="left", fill="both", expand=True)

        tk.Label(left, text="可选值（多选）:", font=("Microsoft YaHei", 10, "bold")).pack(anchor="w")
        self.list_available = tk.Listbox(
            left, selectmode="multiple", width=22, height=15,
            font=("Microsoft YaHei", 10), exportselection=False
        )
        self.list_available.pack(fill="both", expand=True, pady=(4, 6))
        self._refresh_available()

        btn_add = tk.Button(
            left, text="添加到新组合 >>", font=("Microsoft YaHei", 10),
            bg="#27ae60", fg="white", relief="flat", cursor="hand2",
            command=self._add_group
        )
        btn_add.pack(fill="x", pady=(0, 4))

        # ---- 右侧：已创建组合 ----
        right = tk.Frame(mid_frame)
        right.pack(side="right", fill="both", expand=True, padx=(12, 0))

        tk.Label(right, text="已创建的组合:", font=("Microsoft YaHei", 10, "bold")).pack(anchor="w")
        self.list_groups = tk.Listbox(
            right, width=48, height=15,
            font=("Microsoft YaHei", 10), exportselection=False
        )
        self.list_groups.pack(fill="both", expand=True, pady=(4, 6))

        btn_del = tk.Button(
            right, text="<< 删除选中组合", font=("Microsoft YaHei", 10),
            bg="#e74c3c", fg="white", relief="flat", cursor="hand2",
            command=self._remove_group
        )
        btn_del.pack(fill="x", pady=(0, 4))

        # ---- 底部按钮 ----
        bottom = tk.Frame(self)
        bottom.pack(fill="x", padx=15, pady=(6, 12))

        tk.Label(
            bottom, text="提示：未分组的值将各自单独拆分",
            font=("Microsoft YaHei", 9), fg="#888888"
        ).pack(side="left")

        tk.Button(
            bottom, text="确认", width=10, font=("Microsoft YaHei", 10),
            bg="#2980b9", fg="white", relief="flat", cursor="hand2",
            command=self._confirm
        ).pack(side="right", padx=(6, 0))
        tk.Button(
            bottom, text="取消", width=10, font=("Microsoft YaHei", 10),
            bg="#95a5a6", fg="white", relief="flat", cursor="hand2",
            command=self.destroy
        ).pack(side="right")

    def _get_available(self):
        """返回当前未分组的剩余值"""
        grouped = set()
        for g in self.groups:
            for v in g:
                grouped.add(v)
        return [v for v in self.all_values if v not in grouped]

    def _refresh_available(self):
        self.list_available.delete(0, "end")
        for v in self._get_available():
            self.list_available.insert("end", v)

    def _refresh_groups(self):
        self.list_groups.delete(0, "end")
        for i, g in enumerate(self.groups, 1):
            label = f"组合{i}: {', '.join(g)}"
            self.list_groups.insert("end", label)

    def _add_group(self):
        sel_indices = self.list_available.curselection()
        if not sel_indices:
            messagebox.showwarning("提示", "请先在左侧多选要组合的值。", parent=self)
            return
        available = self._get_available()
        members = [available[i] for i in sel_indices]
        self.groups.append(members)
        self._refresh_available()
        self._refresh_groups()

    def _remove_group(self):
        sel = self.list_groups.curselection()
        if not sel:
            messagebox.showwarning("提示", "请先在右侧选中要删除的组合。", parent=self)
            return
        idx = sel[0]
        if 0 <= idx < len(self.groups):
            del self.groups[idx]
            self._refresh_available()
            self._refresh_groups()

    def _confirm(self):
        self.result = self.groups
        self.destroy()


def _sanitize_filename(name: str) -> str:
    """清理文件名/Sheet名中的非法字符"""
    name = name.replace('\n', ' ').replace('\r', ' ').replace('\t', ' ').strip()
    invalid_chars = '\\/:*?"<>|'
    for ch in invalid_chars:
        name = name.replace(ch, '_')
    return name


class SplitExcelApp:
    """
    简化版 Excel 表格拆分工具。
    支持嵌入模式：传入 parent 参数时，UI 构建在 parent 上；
    不传 parent 时，自动创建 tk.Tk 顶级窗口（兼容直接运行）。
    """
    def __init__(self, parent=None):
        if parent is None:
            self.root = tk.Tk()
            self._is_top_level = True
        else:
            self.root = parent.winfo_toplevel()
            self._is_top_level = False

        self.parent = parent

        # 业务数据
        self.filepath = ""
        self.wb = None
        self.ws = None
        self.sheet_names = []
        self.headers = []
        self.split_col_idx = None
        self.footer_start = None
        self.merge_map = {}
        self.dept_rows = {}
        self.dept_order = []
        self.selected_groups = None  # 用户通过弹窗设置的组合
        self.output_dir = ""

        self._build_ui()

        if self._is_top_level:
            self.root.title("Excel表格拆分工具（简化版）")
            self.root.geometry("600x400")
            self.root.resizable(False, False)

    def _build_ui(self):
        """构建界面，所有组件挂载在 self.parent 或 self.root 上"""
        padx = 10
        pady = 5

        # 确定 UI 挂载的父容器
        container = self.parent if self.parent is not None else self.root
        # 嵌入模式下也需要配置容器权重，使日志区域自动扩展
        container.columnconfigure(1, weight=1)
        container.rowconfigure(4, weight=1)

        tk.Label(container, text="Excel文件:").grid(row=0, column=0, sticky="e", padx=padx, pady=pady)
        self.entry_file = tk.Entry(container, width=45)
        self.entry_file.grid(row=0, column=1, sticky="ew", padx=padx, pady=pady)
        tk.Button(container, text="浏览", command=self._choose_file).grid(row=0, column=2, padx=padx, pady=pady)

        tk.Label(container, text="工作表:").grid(row=1, column=0, sticky="e", padx=padx, pady=pady)
        self.combo_sheet = ttk.Combobox(container, width=42, state="readonly")
        self.combo_sheet.grid(row=1, column=1, sticky="ew", padx=padx, pady=pady)
        self.combo_sheet.bind("<<ComboboxSelected>>", self._on_sheet_selected)

        tk.Label(container, text="拆分依据:").grid(row=2, column=0, sticky="e", padx=padx, pady=pady)
        self.combo_col = ttk.Combobox(container, width=42, state="readonly")
        self.combo_col.grid(row=2, column=1, sticky="ew", padx=padx, pady=pady)

        # 操作按钮
        btn_frame = tk.Frame(container)
        btn_frame.grid(row=3, column=0, columnspan=3, pady=15)
        tk.Button(btn_frame, text="选择拆分", command=self._open_group_dialog, width=15).pack(side=tk.LEFT, padx=5)
        tk.Button(btn_frame, text="开始拆分", command=self._start_split, width=15).pack(side=tk.LEFT, padx=5)
        tk.Button(btn_frame, text="打开输出目录", command=self._open_output_dir, width=15).pack(side=tk.LEFT, padx=5)

        # 状态/日志
        tk.Label(container, text="日志:").grid(row=4, column=0, sticky="ne", padx=padx, pady=pady)
        self.text_log = tk.Text(container, width=65, height=12, state="disabled")
        self.text_log.grid(row=4, column=1, columnspan=2, padx=padx, pady=pady, sticky="nsew")

    def _log(self, msg: str):
        self.text_log.config(state="normal")
        self.text_log.insert(tk.END, msg + "\n")
        self.text_log.see(tk.END)
        self.text_log.config(state="disabled")
        if self.parent is not None:
            self.parent.update_idletasks()
        else:
            self.root.update_idletasks()

    def _open_group_dialog(self):
        """打开选择拆分组合弹窗"""
        if not self.wb or not self.ws:
            messagebox.showwarning("提示", "请先选择Excel文件和工作表，并选择拆分依据列。")
            return
        col_text = self.combo_col.get()
        if not col_text:
            messagebox.showwarning("提示", "请先选择拆分依据列。")
            return
        self.split_col_idx = int(col_text.split(".")[0])

        # 收集拆分列的所有不重复值
        self.merge_map = self._build_merge_map(self.ws, self.split_col_idx)
        seen = set()
        all_values = []
        for r in range(3, self.ws.max_row + 1):
            val = self.merge_map.get(r)
            if val is None:
                val = self.ws.cell(r, self.split_col_idx).value
            if val is not None:
                dept = str(val).strip()
                if dept and dept not in seen:
                    seen.add(dept)
                    all_values.append(dept)

        if not all_values:
            messagebox.showwarning("提示", "未找到可拆分的值。")
            return

        dialog = SplitGroupDialog(self.root, all_values)
        self.root.wait_window(dialog)

        if dialog.result is not None:
            self.selected_groups = dialog.result
            group_desc = []
            for g in self.selected_groups:
                group_desc.append("+".join(g))
            self._log(f"已设置组合: {', '.join(group_desc)}")
            messagebox.showinfo("选择拆分", f"已创建 {len(self.selected_groups)} 个组合，未选择的值将单独拆分。")

    def _open_output_dir(self):
        """打开拆分输出目录"""
        if not self.output_dir or not os.path.exists(self.output_dir):
            messagebox.showwarning("提示", "暂无输出目录，请先执行拆分操作。")
            return
        try:
            if os.name == 'nt':
                os.startfile(self.output_dir)
            elif sys.platform == 'darwin':
                subprocess.run(['open', self.output_dir])
            else:
                subprocess.run(['xdg-open', self.output_dir])
        except Exception as e:
            messagebox.showerror("错误", f"无法打开输出目录:\n{e}")

    def _choose_file(self):
        path = filedialog.askopenfilename(filetypes=[("Excel文件", "*.xlsx")])
        if not path:
            if self.parent is None:
                self.root.lift()
                self.root.focus_force()
            return
        self.filepath = path
        self.entry_file.delete(0, tk.END)
        self.entry_file.insert(0, path)
        try:
            self.wb = load_workbook(path, data_only=False)
            self.sheet_names = self.wb.sheetnames
            self.combo_sheet["values"] = self.sheet_names
            if self.sheet_names:
                self.combo_sheet.current(0)
                self._on_sheet_selected(None)
            self._log(f"已加载文件: {os.path.basename(path)}")
            # 重置组合选择
            self.selected_groups = None
        except Exception as e:
            messagebox.showerror("错误", f"无法加载文件:\n{e}")
        finally:
            if self.parent is None:
                self.root.lift()
                self.root.focus_force()

    def _on_sheet_selected(self, event):
        if not self.wb:
            return
        sheet_name = self.combo_sheet.get()
        if not sheet_name:
            return
        self.ws = self.wb[sheet_name]
        headers = []
        row2 = self.ws[2]
        for idx, cell in enumerate(row2, 1):
            val = cell.value if cell.value is not None else ""
            headers.append((idx, str(val).strip()))
        self.headers = headers
        self.combo_col["values"] = [f"{col_idx}. {name}" for col_idx, name in headers]
        if headers:
            self.combo_col.current(0)
        self._log(f"已选择工作表: {sheet_name}, 第2行共 {len(headers)} 个表头")
        # 重置组合选择
        self.selected_groups = None

    def _build_merge_map(self, ws, col_idx: int):
        merge_map = {}
        for merged_range in ws.merged_cells.ranges:
            min_col, min_row, max_col, max_row = merged_range.bounds
            if min_col <= col_idx <= max_col:
                val = ws.cell(min_row, min_col).value
                for r in range(min_row, max_row + 1):
                    merge_map[r] = val
        return merge_map

    def _detect_footer(self, ws, col_idx: int):
        merge_map = self._build_merge_map(ws, col_idx)
        max_row = ws.max_row
        last_data_row = 2
        for r in range(3, max_row + 1):
            val = merge_map.get(r)
            if val is None:
                val = ws.cell(r, col_idx).value
            if val is not None and str(val).strip():
                txt = str(val).strip()
                if txt not in ("合计", "总计", "汇总", "小计", "平均"):
                    last_data_row = r
        return last_data_row + 1

    def _copy_cell_style(self, src, dst):
        if not src.has_style:
            return
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.number_format = copy(src.number_format)
        dst.protection = copy(src.protection)
        dst.alignment = copy(src.alignment)

    def _adjust_formula(self, formula: str, row_map: dict, old_last_data_row: int, new_last_data_row: int) -> str:
        cross_pattern = re.compile(
            r"('[^']+'|[A-Za-z0-9_\u4e00-\u9fa5]+)!"
            r"(\$?[A-Za-z]{1,3}\$?\d+"
            r"(:\$?[A-Za-z]{1,3}\$?\d+)?)"
        )
        cross_ph = []

        def replace_cross(m):
            cross_ph.append(m.group(0))
            return f"__CROSS{len(cross_ph) - 1}__"

        temp = cross_pattern.sub(replace_cross, formula)

        pattern = re.compile(r'(\$?[A-Za-z]{1,3}\$?)(\d+)')

        def repl(m):
            col_part = m.group(1)
            row_num = int(m.group(2))
            if row_num in row_map:
                new_num = row_map[row_num]
            elif row_num == old_last_data_row:
                new_num = new_last_data_row
            else:
                new_num = row_num
            return f"{col_part}{new_num}"

        temp = pattern.sub(repl, temp)

        for i, ph in enumerate(cross_ph):
            temp = temp.replace(f"__CROSS{i}__", ph)

        return temp

    def _get_safe_sheet_name(self, wb, base_name: str) -> str:
        name = base_name
        idx = 1
        while name in wb.sheetnames:
            name = f"{base_name}_{idx}"
            idx += 1
        return name

    def _start_split(self):
        if not self.filepath or not self.ws:
            messagebox.showwarning("提示", "请先选择Excel文件和工作表")
            return

        col_text = self.combo_col.get()
        if not col_text:
            messagebox.showwarning("提示", "请选择拆分依据列")
            return

        self.split_col_idx = int(col_text.split(".")[0])

        suggested_footer = self._detect_footer(self.ws, self.split_col_idx)
        max_row = self.ws.max_row

        if suggested_footer > max_row:
            messagebox.showinfo("提示", "未检测到表尾，所有行将被视为数据行")
            self.footer_start = max_row + 1
        else:
            confirm = messagebox.askyesno(
                "确认表尾",
                f"检测到数据区域: 第3行 ~ 第{suggested_footer - 1}行\n"
                f"表尾起始行: 第{suggested_footer}行 (共 {max_row - suggested_footer + 1} 行)\n\n"
                f"是否确认并开始拆分？"
            )
            if not confirm:
                return
            self.footer_start = suggested_footer

        # 构建部门映射
        self.merge_map = self._build_merge_map(self.ws, self.split_col_idx)
        self.dept_rows = {}
        self.dept_order = []
        seen = set()

        for r in range(3, self.footer_start):
            val = self.merge_map.get(r)
            if val is None:
                val = self.ws.cell(r, self.split_col_idx).value
            if val is not None:
                dept = str(val).strip()
                if dept and dept not in seen:
                    seen.add(dept)
                    self.dept_order.append(dept)
                    self.dept_rows[dept] = []
                if dept in self.dept_rows:
                    self.dept_rows[dept].append(r)

        if not self.dept_order:
            messagebox.showwarning("提示", "未找到有效的拆分值")
            return

        # 根据用户选择构建分组
        if self.selected_groups:
            grouped_depts = set()
            split_groups = []
            for g in self.selected_groups:
                split_groups.append(g)
                for d in g:
                    grouped_depts.add(d)
            for d in self.dept_order:
                if d not in grouped_depts:
                    split_groups.append([d])
        else:
            split_groups = [[d] for d in self.dept_order]

        self._log(f"拆分分组: {split_groups}")

        # 预先计算实际有内容的表尾行
        actual_footer_rows = []
        for r in range(self.footer_start, max_row + 1):
            has_content = False
            for c in range(1, self.ws.max_column + 1):
                cell = self.ws.cell(r, c)
                if cell.value is not None and str(cell.value).strip():
                    has_content = True
                    break
            if has_content:
                actual_footer_rows.append(r)

        # 执行拆分
        base_dir = os.path.dirname(self.filepath)
        base_name = os.path.splitext(os.path.basename(self.filepath))[0]
        original_sheet_name = self.ws.title

        # 输出目录命名：若已存在则自动加数字后缀
        base_output_dir = os.path.join(base_dir, f"{base_name}_拆分结果")
        self.output_dir = base_output_dir
        counter = 1
        while os.path.exists(self.output_dir):
            self.output_dir = f"{base_output_dir}{counter}"
            counter += 1
        os.makedirs(self.output_dir)
        self._log(f"已创建输出目录: {os.path.basename(self.output_dir)}")

        old_last_data_row = self.footer_start - 1
        max_col = self.ws.max_column

        for group in split_groups:
            group_label = "_".join(_sanitize_filename(d) for d in group)
            rows = []
            for d in group:
                rows.extend(self.dept_rows.get(d, []))

            if not rows:
                continue

            new_filename = f"{group_label}.xlsx"
            new_filepath = os.path.join(self.output_dir, new_filename)

            if os.path.exists(new_filepath):
                os.remove(new_filepath)
            shutil.copy(self.filepath, new_filepath)

            wb_new = load_workbook(new_filepath, data_only=False)
            ws_old = wb_new[original_sheet_name]

            # 过滤幽灵行
            filtered_rows = []
            for old_r in rows:
                has_content = False
                for c in range(1, max_col + 1):
                    if c == self.split_col_idx:
                        continue
                    cell = ws_old.cell(old_r, c)
                    if cell.value is not None and str(cell.value).strip():
                        has_content = True
                        break
                if has_content:
                    filtered_rows.append(old_r)

            if not filtered_rows:
                self._log(f"  跳过 {group_label}: 过滤后无有效数据行")
                wb_new.close()
                os.remove(new_filepath)
                continue

            rows = filtered_rows

            keep_rows = [1, 2] + rows + actual_footer_rows
            keep_set = set(keep_rows)
            row_map = {}
            new_r = 1
            for old_r in keep_rows:
                row_map[old_r] = new_r
                new_r += 1

            new_data_last_row = row_map[rows[-1]]

            # 先删除原始Sheet
            if original_sheet_name in wb_new.sheetnames:
                del wb_new[original_sheet_name]

            # 创建新Sheet
            safe_sheet_name = _sanitize_filename(original_sheet_name)
            ws_new = wb_new.create_sheet(title=safe_sheet_name, index=0)

            for new_r, old_r in enumerate(keep_rows, 1):
                for c in range(1, max_col + 1):
                    old_cell = ws_old.cell(old_r, c)
                    new_cell = ws_new.cell(new_r, c)
                    if isinstance(old_cell.value, str) and old_cell.value.startswith("="):
                        new_cell.value = old_cell.value
                    else:
                        new_cell.value = old_cell.value
                    self._copy_cell_style(old_cell, new_cell)

                if old_r in ws_old.row_dimensions:
                    ws_new.row_dimensions[new_r].height = ws_old.row_dimensions[old_r].height

            for col_letter in ws_old.column_dimensions:
                ws_new.column_dimensions[col_letter].width = ws_old.column_dimensions[col_letter].width

            # 重建合并单元格
            for merged_range in ws_old.merged_cells.ranges:
                mc_min_col, mc_min_row, mc_max_col, mc_max_row = merged_range.bounds
                rows_in_range = set(range(mc_min_row, mc_max_row + 1))
                if rows_in_range.issubset(keep_set):
                    new_min_row = row_map[mc_min_row]
                    new_max_row = row_map[mc_max_row]
                    ws_new.merge_cells(start_row=new_min_row, start_column=mc_min_col,
                                       end_row=new_max_row, end_column=mc_max_col)

            # 为"个税"列填充固定公式
            col_idx_map = {}
            for c in range(1, max_col + 1):
                val = ws_new.cell(2, c).value
                if val is not None:
                    col_idx_map[str(val).strip()] = c

            if "个税" in col_idx_map:
                geshui_col = col_idx_map["个税"]
                for r in range(3, new_data_last_row + 1):
                    ws_new.cell(r, geshui_col).value = (
                        f"=-ROUND(MAX((AK{r}-5000-AJ{r}-AT{r}+AM{r}+AN{r})"
                        f"*{{3,10,20,25,30,35,45}}%"
                        f"-5*{{0,42,282,532,882,1432,3032}},),2)"
                    )
                self._log(f"  已填充'个税'列公式（列{geshui_col}）")

            # 调整公式
            new_max_row = ws_new.max_row
            for r in range(1, new_max_row + 1):
                for c in range(1, max_col + 1):
                    cell = ws_new.cell(r, c)
                    if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                        old_formula = cell.value
                        new_formula = self._adjust_formula(
                            old_formula, row_map, old_last_data_row, new_data_last_row
                        )
                        cell.value = new_formula

            sheets_to_remove = ["各公司支出", "未发工资的社保扣除人员"]
            for sheet_name in sheets_to_remove:
                if sheet_name in wb_new.sheetnames:
                    wb_new.remove(wb_new[sheet_name])

            wb_new.save(new_filepath)
            wb_new.close()
            self._log(f"已生成: {new_filename} (数据行 {len(rows)} 条)")

        self._log("拆分完成!")
        messagebox.showinfo("完成", f"拆分完成，共生成 {len(split_groups)} 个文件。")


def main():
    app = SplitExcelApp()
    app.root.mainloop()


if __name__ == "__main__":
    main()
