#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
数据去重工具 - 核心逻辑模块
功能：按姓名去重，保留日期最新的一整行数据，写入新Sheet并保留格式
"""

import os
import re
from datetime import datetime
from typing import Optional, List, Dict, Any, Tuple

import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import (
    Alignment, Font, PatternFill, Border, Side, Protection, GradientFill
)
from openpyxl.styles.numbers import FORMAT_GENERAL
from openpyxl.worksheet.worksheet import Worksheet


def get_sheet_names(file_path: str) -> List[str]:
    """获取Excel文件中所有工作表名称"""
    wb = load_workbook(file_path, read_only=True, data_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def get_headers(file_path: str, sheet_name: str, header_row: int = 1) -> List[str]:
    """
    读取指定sheet的表头行
    :param file_path: Excel文件路径
    :param sheet_name: 工作表名称
    :param header_row: 表头所在行（默认第1行）
    :return: 表头列表
    """
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    headers = []
    for cell in ws[header_row]:
        value = cell.value
        headers.append(str(value).strip() if value is not None else "")
    wb.close()
    return headers


def _parse_chinese_date(date_str: str) -> Optional[datetime]:
    """
    解析中文日期字符串，如 '1月21日'、'12月3日'
    返回 datetime 对象（年份使用当前年份）
    """
    if not date_str or not isinstance(date_str, str):
        return None
    match = re.match(r"(\d{1,2})月(\d{1,2})日?", date_str.strip())
    if match:
        month, day = int(match.group(1)), int(match.group(2))
        try:
            return datetime(datetime.now().year, month, day)
        except ValueError:
            return None
    return None


def _parse_excel_date(cell_value: Any) -> Optional[datetime]:
    """
    尝试解析单元格中的日期值
    支持：datetime对象、Excel日期序列号（int/float）、中文日期字符串
    """
    # 情况1：datetime 对象（data_only=True 且单元格是日期格式时返回）
    if isinstance(cell_value, datetime):
        return cell_value
    # 情况2：Excel日期序列号（int/float），如 46045.0
    # 注意：排除 bool 类型（bool 是 int 的子类）
    if isinstance(cell_value, (int, float)) and not isinstance(cell_value, bool) and cell_value > 0:
        try:
            from openpyxl.utils.datetime import from_excel
            return from_excel(cell_value)
        except Exception:
            return None
    # 情况3：字符串，尝试解析中文日期
    if isinstance(cell_value, str):
        return _parse_chinese_date(cell_value)
    return None


def find_column_index(headers: List[str], keyword: str) -> Optional[int]:
    """
    在表头中查找包含关键字的列索引（从0开始）
    """
    for idx, h in enumerate(headers):
        if keyword in h:
            return idx
    return None


def process_deduplication(
    file_path: str,
    sheet_name: str,
    name_col_keyword: str = "姓名",
    date_col_keyword: str = "日期",
    id_col_keyword: str = "身份证",
    header_row: int = 1,
) -> Tuple[List[str], List[List[Any]], int, int]:
    """
    执行去重处理（按姓名+身份证联合去重）

    :param file_path: Excel文件路径
    :param sheet_name: 源工作表名称
    :param name_col_keyword: 姓名列的表头关键字
    :param date_col_keyword: 日期列的表头关键字
    :param id_col_keyword: 身份证列的表头关键字
    :param header_row: 表头所在行
    :return: (表头列表, 去重后的数据行, 总数据行数, 去重后行数)
    """
    # data_only=True 获取公式计算后的值（用于去重比较和预览显示）
    wb = load_workbook(file_path, data_only=True)
    ws = wb[sheet_name]

    # 读取表头
    headers = []
    for cell in ws[header_row]:
        value = cell.value
        headers.append(str(value).strip() if value is not None else "")

    # 定位姓名列、日期列、身份证列
    name_idx = find_column_index(headers, name_col_keyword)
    date_idx = find_column_index(headers, date_col_keyword)
    id_idx = find_column_index(headers, id_col_keyword)

    if name_idx is None:
        wb.close()
        raise ValueError(f"未找到包含'{name_col_keyword}'的表头列")

    # 读取所有数据行（从表头行的下一行开始）
    rows_data = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1), start=header_row + 1):
        # 边界检查：防止某行列数不足导致越界
        if name_idx >= len(row):
            continue
        name_val = row[name_idx].value
        if name_val is None or str(name_val).strip() == "":
            continue
        name_val = str(name_val).strip()

        # 解析身份证（用于联合分组）
        id_val = ""
        if id_idx is not None and id_idx < len(row):
            raw_id = row[id_idx].value
            if raw_id is not None:
                id_val = str(raw_id).strip()

        # 解析日期
        date_val = None
        if date_idx is not None and date_idx < len(row):
            date_val = _parse_excel_date(row[date_idx].value)

        # 联合key：姓名|身份证（身份证为空时退化为只按姓名）
        group_key = f"{name_val}|{id_val}" if id_val else name_val

        rows_data.append({
            "row_index": row_idx,
            "name": name_val,
            "id": id_val,
            "group_key": group_key,
            "date": date_val,
            "date_raw": row[date_idx].value if date_idx is not None and date_idx < len(row) else None,
        })

    total_count = len(rows_data)

    # 按联合key分组，取日期最新的一行
    # 规则：
    #   1) 日期不同 → 取日期最新的
    #   2) 日期相同 → 取最后出现的那条（后出现的覆盖先出现的）
    #   3) 无日期 → 取最后出现的
    group_map: Dict[str, Dict] = {}
    for item in rows_data:
        key = item["group_key"]
        if key not in group_map:
            group_map[key] = item
        else:
            existing = group_map[key]
            item_date = item["date"]
            exist_date = existing["date"]

            # 情况A：item日期 > existing日期 → 替换
            if item_date is not None and exist_date is not None:
                if item_date > exist_date:
                    group_map[key] = item
                elif item_date == exist_date:
                    # 日期相同 → 取最后出现的（直接替换）
                    group_map[key] = item
            # 情况B：item有日期，existing无日期 → 替换
            elif item_date is not None and exist_date is None:
                group_map[key] = item
            # 情况C：日期都相同（都为None）→ 取最后出现的
            elif item_date is None and exist_date is None:
                group_map[key] = item

    # 按原顺序排序去重后的行索引
    selected_rows = sorted(group_map.values(), key=lambda x: x["row_index"])

    deduplicated_count = len(selected_rows)

    # 收集去重后的整行数据（用于预览）
    preview_data = []
    max_col = len(headers)
    for item in selected_rows:
        row_idx = item["row_index"]
        row_data = []
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            row_data.append(val)
        preview_data.append(row_data)

    wb.close()

    return headers, preview_data, total_count, deduplicated_count


def write_deduplication_result(
    file_path: str,
    sheet_name: str,
    new_sheet_name: str,
    header_row: int = 1,
    name_col_keyword: str = "姓名",
    date_col_keyword: str = "日期",
    id_col_keyword: str = "身份证",
) -> str:
    """
    将去重结果写入新Sheet（保留格式、公式等）
    分两步执行：
      1) data_only=True 读取用于去重比较（获取公式计算值）
      2) data_only=False 读取用于复制整行（保留公式和格式）

    :param file_path: Excel文件路径
    :param sheet_name: 源工作表名称
    :param new_sheet_name: 新工作表名称
    :param header_row: 表头所在行
    :param name_col_keyword: 姓名列的表头关键字
    :param date_col_keyword: 日期列的表头关键字
    :param id_col_keyword: 身份证列的表头关键字
    :return: 实际创建的新Sheet名称
    """
    # ========== 第1步：data_only=True 去重比较，获取需要保留的源行号 ==========
    wb_compare = load_workbook(file_path, data_only=True)
    ws_compare = wb_compare[sheet_name]

    # 读取表头
    headers = []
    for cell in ws_compare[header_row]:
        value = cell.value
        headers.append(str(value).strip() if value is not None else "")

    # 定位姓名列、日期列、身份证列
    name_idx = find_column_index(headers, name_col_keyword)
    date_idx = find_column_index(headers, date_col_keyword)
    id_idx = find_column_index(headers, id_col_keyword)

    if name_idx is None:
        wb_compare.close()
        raise ValueError(f"未找到包含'{name_col_keyword}'的表头列")

    # 读取所有数据行进行去重比较
    rows_data = []
    for row_idx, row in enumerate(ws_compare.iter_rows(min_row=header_row + 1), start=header_row + 1):
        if name_idx >= len(row):
            continue
        name_val = row[name_idx].value
        if name_val is None or str(name_val).strip() == "":
            continue
        name_val = str(name_val).strip()

        # 解析身份证（用于联合分组）
        id_val = ""
        if id_idx is not None and id_idx < len(row):
            raw_id = row[id_idx].value
            if raw_id is not None:
                id_val = str(raw_id).strip()

        date_val = None
        if date_idx is not None and date_idx < len(row):
            date_val = _parse_excel_date(row[date_idx].value)

        # 联合key：姓名|身份证（身份证为空时退化为只按姓名）
        group_key = f"{name_val}|{id_val}" if id_val else name_val

        rows_data.append({
            "row_index": row_idx,
            "name": name_val,
            "id": id_val,
            "group_key": group_key,
            "date": date_val,
        })

    # 按联合key分组，取日期最新的一行
    # 规则：
    #   1) 日期不同 → 取日期最新的
    #   2) 日期相同 → 取最后出现的那条（后出现的覆盖先出现的）
    #   3) 无日期 → 取最后出现的
    group_map: Dict[str, Dict] = {}
    for item in rows_data:
        key = item["group_key"]
        if key not in group_map:
            group_map[key] = item
        else:
            existing = group_map[key]
            item_date = item["date"]
            exist_date = existing["date"]

            # 情况A：item日期 > existing日期 → 替换
            if item_date is not None and exist_date is not None:
                if item_date > exist_date:
                    group_map[key] = item
                elif item_date == exist_date:
                    # 日期相同 → 取最后出现的（直接替换）
                    group_map[key] = item
            # 情况B：item有日期，existing无日期 → 替换
            elif item_date is not None and exist_date is None:
                group_map[key] = item
            # 情况C：日期都相同（都为None）→ 取最后出现的
            elif item_date is None and exist_date is None:
                group_map[key] = item

    # 按原顺序排序，得到需要保留的源行号列表
    selected_items = sorted(group_map.values(), key=lambda x: x["row_index"])
    selected_row_indices = [item["row_index"] for item in selected_items]

    wb_compare.close()

    # ========== 第2步：data_only=False 复制整行（保留公式和格式） ==========
    wb = load_workbook(file_path, data_only=False)

    # 处理新Sheet名称冲突
    actual_new_name = new_sheet_name
    if actual_new_name in wb.sheetnames:
        counter = 1
        while f"{new_sheet_name}_{counter}" in wb.sheetnames:
            counter += 1
        actual_new_name = f"{new_sheet_name}_{counter}"

    # 创建新Sheet
    new_ws = wb.create_sheet(title=actual_new_name)

    # 获取源Sheet
    src_ws = wb[sheet_name]

    # 复制表头行（保留格式）
    _copy_row(src_ws, new_ws, header_row, 1, len(headers), date_idx=date_idx)

    # 复制选中的行到新Sheet
    for new_row_idx, src_row_idx in enumerate(selected_row_indices, start=2):
        _copy_row(src_ws, new_ws, src_row_idx, new_row_idx, len(headers), date_idx=date_idx)

    # 自动调整列宽
    _auto_adjust_column_width(new_ws, src_ws, len(headers))

    # 保存文件
    wb.save(file_path)
    wb.close()

    return actual_new_name


def _copy_cell_style(src_cell, dst_cell):
    """复制单元格样式"""
    if src_cell.has_style:
        dst_cell.font = src_cell.font.copy()
        dst_cell.fill = src_cell.fill.copy()
        dst_cell.border = src_cell.border.copy()
        dst_cell.alignment = src_cell.alignment.copy()
        dst_cell.number_format = src_cell.number_format
        dst_cell.protection = src_cell.protection.copy()
        if src_cell.style and src_cell.style != "Normal":
            dst_cell.style = src_cell.style


def _copy_row(
    src_ws: Worksheet,
    dst_ws: Worksheet,
    src_row_idx: int,
    dst_row_idx: int,
    max_col: int,
    date_idx: Optional[int] = None,
):
    """复制整行数据（包括值、公式、格式）
    :param date_idx: 日期列的0-based索引。若指定，该列的datetime值会被转为中文日期文本（如"1月21日"），避免Excel底层存储为datetime导致编辑栏显示日期序列值。
    """
    # 日期列的1-based列号（openpyxl使用1-based）
    date_col = date_idx + 1 if date_idx is not None else None

    # 使用 cell(row=, column=) 方式访问，避免新Sheet行单元格不足导致索引越界
    for col_idx in range(1, max_col + 1):
        src_cell = src_ws.cell(row=src_row_idx, column=col_idx)
        dst_cell = dst_ws.cell(row=dst_row_idx, column=col_idx)

        # 复制值（如果是公式，保留公式）
        if src_cell.value is not None:
            # 日期列特殊处理：datetime 或 Excel日期序列号 转为中文日期字符串文本
            if col_idx == date_col:
                val = src_cell.value
                converted = False
                # 情况1：datetime 对象
                if isinstance(val, datetime):
                    dst_cell.value = f"{val.month}月{val.day}日"
                    converted = True
                # 情况2：整数/浮点数（Excel日期序列号）
                elif isinstance(val, (int, float)) and not isinstance(val, bool) and val > 0:
                    try:
                        from openpyxl.utils.datetime import from_excel
                        dt = from_excel(val)
                        dst_cell.value = f"{dt.month}月{dt.day}日"
                        converted = True
                    except Exception:
                        dst_cell.value = val  # 不是有效日期序列号，保持原值
                else:
                    # 文本等其他类型，保持原样
                    dst_cell.value = val

                if converted:
                    dst_cell.number_format = "@"  # 强制设为文本格式

            elif src_cell.data_type == 'f':  # 公式
                dst_cell.value = src_cell.value
            else:
                dst_cell.value = src_cell.value

        # 复制样式（日期列已手动设置number_format，跳过完整样式复制中的number_format覆盖）
        if col_idx == date_col and src_cell.value is not None:
            val = src_cell.value
            is_date = isinstance(val, datetime) or (
                isinstance(val, (int, float)) and not isinstance(val, bool) and val > 0
            )
            if is_date:
                # 保留除number_format外的其他样式
                if src_cell.has_style:
                    dst_cell.font = src_cell.font.copy()
                    dst_cell.fill = src_cell.fill.copy()
                    dst_cell.border = src_cell.border.copy()
                    dst_cell.alignment = src_cell.alignment.copy()
                    dst_cell.protection = src_cell.protection.copy()
                    if src_cell.style and src_cell.style != "Normal":
                        dst_cell.style = src_cell.style
            else:
                _copy_cell_style(src_cell, dst_cell)
        else:
            _copy_cell_style(src_cell, dst_cell)


def _auto_adjust_column_width(dst_ws: Worksheet, src_ws: Worksheet, max_col: int):
    """自动调整列宽（基于源Sheet的列宽）"""
    for col_idx in range(1, max_col + 1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        if col_letter in src_ws.column_dimensions:
            dst_ws.column_dimensions[col_letter].width = src_ws.column_dimensions[col_letter].width
