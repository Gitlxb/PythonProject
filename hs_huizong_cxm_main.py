"""
Excel汇总处理工具 - GUI启动界面
"""
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
import pandas as pd
import os
import re
from openpyxl.utils import get_column_letter

from hs_huizong_cxm import SummaryProcessor, DataChecker, copy_sheet_to_workbook
from openpyxl import load_workbook


class SummaryProcessorApp:
    """汇总处理GUI应用程序"""

    def __init__(self, root: tk.Tk):
        """
        初始化应用程序

        Args:
            root: Tkinter根窗口
        """
        self.root = root
        self.root.title("Excel汇总处理工具")
        self.root.geometry("1000x800")
        self.root.minsize(900, 600)

        # 配置窗口可调整大小
        self.root.columnconfigure(0, weight=1)
        self.root.rowconfigure(0, weight=1)

        self.processor = SummaryProcessor()
        self.current_sheet = tk.StringVar()
        self.selected_months_for_save = None  # 保存用户选择的月份列表，用于生成工作表标题
        self.last_saved_path = None  # 记录上次保存的文件路径，用于打开输出目录
        self.reference_path = None  # 基准表（汇总表）路径，专用于核对功能
        self.check_result_path = None  # 核对结果保存路径，专用于合并到原表

        # 配置Treeview样式，减小行高
        self._setup_treeview_style()

        self._create_ui()

    def _setup_treeview_style(self):
        """配置Treeview样式，优化显示效果"""
        style = ttk.Style()

        # 设置Treeview的行高
        style.configure("Treeview",
                       rowheight=25,
                       font=('Microsoft YaHei UI', 9))

        # 设置标题样式
        style.configure("Treeview.Heading",
                       font=('Microsoft YaHei UI', 10, 'bold'))

        # 设置选中项样式
        style.map("Treeview",
                 background=[('selected', '#0078d7')],
                 foreground=[('selected', 'white')])

    def _create_ui(self):
        """创建用户界面"""
        # 主框架
        main_frame = ttk.Frame(self.root, padding="15")
        main_frame.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))

        # 配置主框架的网格权重
        main_frame.columnconfigure(0, weight=1)

        # ========== 文件选择区域 ==========
        file_frame = ttk.LabelFrame(main_frame, text="文件选择", padding="12")
        file_frame.grid(row=0, column=0, sticky=(tk.W, tk.E), pady=(0, 12))
        file_frame.columnconfigure(0, weight=1)

        self.file_path_var = tk.StringVar()
        file_entry = ttk.Entry(file_frame, textvariable=self.file_path_var, font=('Microsoft YaHei UI', 9))
        file_entry.grid(row=0, column=0, padx=(0, 10), sticky=(tk.W, tk.E))

        browse_btn = ttk.Button(file_frame, text="浏览...", command=self._browse_file, width=12)
        browse_btn.grid(row=0, column=1, padx=(0, 15))

        ttk.Label(file_frame, text="工作表:", font=('Microsoft YaHei UI', 9)).grid(row=0, column=2, padx=(0, 5))
        self.sheet_combo = ttk.Combobox(file_frame, textvariable=self.current_sheet,
                                       state='readonly', width=25, font=('Microsoft YaHei UI', 9))
        self.sheet_combo.grid(row=0, column=3)

        # ========== 操作按钮区域 ==========
        button_frame = ttk.Frame(main_frame)
        button_frame.grid(row=1, column=0, sticky=(tk.W, tk.E), pady=(0, 12))

        process_btn = ttk.Button(button_frame, text="📊数据处理", command=self._process_data, width=14)
        process_btn.grid(row=0, column=0, padx=(0, 8))

        save_btn = ttk.Button(button_frame, text="💾文件保存", command=self._save_file, width=14)
        save_btn.grid(row=0, column=1, padx=(0, 8))

        check_btn = ttk.Button(button_frame, text="✅核对", command=self._check_data, width=14)
        check_btn.grid(row=0, column=2, padx=(0, 8))

        merge_btn = ttk.Button(button_frame, text="🔗合并到原表", command=self._merge_to_original, width=14)
        merge_btn.grid(row=0, column=3, padx=(0, 8))

        open_dir_btn = ttk.Button(button_frame, text="📂打开输出目录", command=self._open_output_dir, width=14)
        open_dir_btn.grid(row=0, column=4, padx=(0, 8))

        clear_btn = ttk.Button(button_frame, text="🗑️清空", command=self._clear_all, width=14)
        clear_btn.grid(row=0, column=5)

        # ========== 数据预览区域 ==========
        preview_frame = ttk.LabelFrame(main_frame, text="数据预览", padding="12")
        preview_frame.grid(row=2, column=0, sticky=(tk.W, tk.E, tk.N, tk.S), pady=(0, 12))
        preview_frame.columnconfigure(0, weight=1)
        preview_frame.rowconfigure(0, weight=1)

        # 创建表格
        self.tree = ttk.Treeview(preview_frame, selectmode='browse', show='headings')
        self.tree.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))

        # 滚动条
        vsb = ttk.Scrollbar(preview_frame, orient="vertical", command=self.tree.yview)
        vsb.grid(row=0, column=1, sticky=(tk.N, tk.S))
        hsb = ttk.Scrollbar(preview_frame, orient="horizontal", command=self.tree.xview)
        hsb.grid(row=1, column=0, sticky=(tk.W, tk.E))

        # 配置Treeview的滚动命令
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        # 初始化滚动位置到最左边
        self.tree.xview_moveto(0)

        # 绑定鼠标滚轮事件
        self.tree.bind("<MouseWheel>", self._on_mousewheel)
        self.tree.bind("<Button-4>", self._on_mousewheel)
        self.tree.bind("<Button-5>", self._on_mousewheel)

        # ========== 状态信息栏 ==========
        status_frame = ttk.LabelFrame(main_frame, text="状态信息", padding="12")
        status_frame.grid(row=3, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        status_frame.columnconfigure(0, weight=1)
        status_frame.rowconfigure(0, weight=1)

        self.status_text = scrolledtext.ScrolledText(status_frame, wrap=tk.WORD,
                                                   font=('Consolas', 9))
        self.status_text.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))

        # 配置网格权重（让数据预览和状态信息自动调整大小）
        main_frame.rowconfigure(2, weight=3)  # 数据预览占3份
        main_frame.rowconfigure(3, weight=1)  # 状态信息占1份

        # 配置文件选择区域的网格权重
        file_frame.columnconfigure(0, weight=1)

        self._log_status("欢迎使用Excel汇总处理工具！")
        self._log_status("请选择一个Excel文件开始处理。")

    def _on_mousewheel(self, event):
        """
        处理鼠标滚轮事件，支持表格滚动

        Args:
            event: 滚轮事件
        """
        if event.num == 4 or event.delta > 0:
            self.tree.yview_scroll(-1, "units")
        elif event.num == 5 or event.delta < 0:
            self.tree.yview_scroll(1, "units")

    def _log_status(self, message: str):
        """
        记录状态信息

        Args:
            message: 状态消息
        """
        timestamp = pd.Timestamp.now().strftime("%H:%M:%S")
        self.status_text.insert(tk.END, f"[{timestamp}] {message}\n")
        self.status_text.see(tk.END)
        self.root.update()

    def _browse_file(self):
        """浏览并选择Excel文件"""
        file_path = filedialog.askopenfilename(
            title="选择Excel文件",
            filetypes=[("Excel文件", "*.xlsx *.xls"), ("所有文件", "*.*")]
        )

        if file_path:
            self._log_status(f"正在读取文件：{os.path.basename(file_path)}...")
            success, message, sheets = self.processor.load_file(file_path)

            if success:
                self.file_path_var.set(file_path)
                self.sheet_combo['values'] = sheets
                if sheets:
                    self.sheet_combo.current(0)
                self._log_status(message)
                self._log_status(f"包含{len(sheets)}个工作表")
            else:
                messagebox.showerror("错误", message)
                self._log_status(f"错误：{message}")

    def _process_data(self):
        """执行数据处理"""
        file_path = self.file_path_var.get()
        sheet_name = self.current_sheet.get()

        if not file_path:
            messagebox.showwarning("警告", "请先选择Excel文件")
            return

        if not sheet_name:
            messagebox.showwarning("警告", "请选择工作表")
            return

        self._log_status("=" * 50)
        self._log_status("开始数据处理...")

        # 加载工作表
        self._log_status(f"1. 加载工作表：{sheet_name}")
        success, message = self.processor.load_sheet(sheet_name)

        if not success:
            messagebox.showerror("错误", message)
            self._log_status(f"错误：{message}")
            return

        self._log_status(message)

        # 处理数据
        self._log_status("2. 检测列...")
        if "备注2" not in self.processor.data_frame.columns:
            messagebox.showerror("错误", "未找到'备注2'列")
            self._log_status("错误：未找到'备注2'列")
            return

        self._log_status("   ✓ 找到'备注2'列")

        if "供应商名称" not in self.processor.data_frame.columns:
            messagebox.showerror("错误", "未找到'供应商名称'列")
            self._log_status("错误：未找到'供应商名称'列")
            return

        self._log_status("   ✓ 找到'供应商名称'列")

        if "收款姓名" not in self.processor.data_frame.columns:
            messagebox.showerror("错误", "未找到'收款姓名'列")
            self._log_status("错误：未找到'收款姓名'列")
            return

        self._log_status("   ✓ 找到'收款姓名'列")

        # 检测打款类型、月份和打款金额列
        self._log_status("2.5 检测可选数据...")
        payment_types = self.processor.detect_payment_types()
        months = self.processor.detect_months()
        amount_columns = self.processor.detect_amount_columns()

        selected_types = None
        selected_months = None
        selected_amount_col_idx = None

        has_multi_amount = len(amount_columns) >= 2
        needs_dialog = payment_types or months or has_multi_amount

        if needs_dialog:
            self._log_status(f"   检测到打款类型：{', '.join(payment_types) if payment_types else '无'}")
            self._log_status(f"   检测到月份：{', '.join(months) if months else '无'}")
            if amount_columns:
                cols_desc = ', '.join([f'{col_letter}列' for _, col_letter in amount_columns])
                self._log_status(f"   检测到打款金额列：{cols_desc}")

            selected_types, selected_months, selected_amount_col_idx = self._show_selection_dialog(
                payment_types, months, amount_columns
            )
            if selected_types is None:  # 用户取消
                self._log_status("用户取消了选择")
                return

            self.selected_months_for_save = selected_months
            self._log_status(f"   已选择打款类型：{', '.join(selected_types) if selected_types else '无'}")
            self._log_status(f"   已选择月份：{', '.join(selected_months) if selected_months else '无'}")
            if selected_amount_col_idx is not None:
                self._log_status(f"   已选择打款金额列：{get_column_letter(selected_amount_col_idx + 1)}列")
        else:
            self._log_status("   未检测到可选项，将处理全部数据")

        self._log_status("3. 应用缩写规则...")
        success, message = self.processor.process_data(selected_types, selected_months, selected_amount_col_idx)

        if not success:
            messagebox.showerror("错误", message)
            self._log_status(f"错误：{message}")
            return

        self._log_status(message)

        # 显示数据预览
        self._log_status("4. 生成汇总表...")
        self._show_preview()

        self._log_status("5. 数据处理完成！")
        self._log_status("=" * 50)

        messagebox.showinfo("成功", "数据处理完成！\n可以点击'文件保存'按钮保存结果。")

    def _show_preview(self):
        """显示数据预览"""
        # 清空现有数据
        for item in self.tree.get_children():
            self.tree.delete(item)

        if self.processor.output_data is None:
            return

        df = self.processor.output_data

        # 设置列
        columns = list(df.columns)
        self.tree['columns'] = columns

        # 设置列标题和列宽
        for col in columns:
            self.tree.heading(col, text=col)

            # 根据列名动态设置列宽（优化宽度以减少横向滚动）
            if col == '供应商名称':
                self.tree.column(col, width=180, anchor=tk.W)
            elif col == '公司':
                self.tree.column(col, width=250, anchor=tk.W)
            elif col == '总计':
                self.tree.column(col, width=100, anchor=tk.E)
            elif col == '':
                self.tree.column(col, width=30, anchor=tk.CENTER)
            else:
                # 公司列
                self.tree.column(col, width=90, anchor=tk.E)

        # 显示数据（最多显示50行）
        display_rows = min(50, len(df))
        for row in df.head(display_rows).itertuples(index=False):
            # 处理空值和零值
            processed_row = tuple(
                '' if (isinstance(v, (int, float)) and v == 0)
                else v if pd.notna(v)
                else ''
                for v in row
            )
            self.tree.insert('', tk.END, values=processed_row)

        # 强制更新界面
        self.tree.update_idletasks()

        # 滚动到最左边
        self.tree.xview_moveto(0)

        self._log_status(f"   ✓ 预览显示前{display_rows}行数据")
        self._log_status(f"   ✓ 共{len(df)}行，{len(columns)}列")

    def _save_file(self):
        """保存文件"""
        if self.processor.output_data is None:
            messagebox.showwarning("警告", "没有数据可保存，请先进行数据处理")
            return

        file_path = filedialog.asksaveasfilename(
            title="保存文件",
            defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")],
            initialfile="供应商拆分汇总"
        )

        if file_path:
            self._log_status("正在保存文件...")
            # 只取单月份（不含"-"）用于工作表标题，忽略"1-2月"等范围格式
            month_for_title = None
            if self.selected_months_for_save:
                single_months = [m for m in self.selected_months_for_save if '-' not in m]
                if single_months:
                    month_for_title = single_months[0]
            success, message = self.processor.save_to_excel(file_path, month_for_title)

            if success:
                self.last_saved_path = file_path
                self.reference_path = file_path  # 基准表路径同步更新
                self._log_status(message)
                messagebox.showinfo("成功", message)
            else:
                messagebox.showerror("错误", message)
                self._log_status(f"错误：{message}")

    def _clear_all(self):
        """清空所有数据"""
        self.file_path_var.set("")
        self.current_sheet.set("")
        self.sheet_combo['values'] = []
        self.selected_months_for_save = None
        self.last_saved_path = None
        self.reference_path = None
        self.check_result_path = None

        # 清空预览
        for item in self.tree.get_children():
            self.tree.delete(item)

        # 重置处理器
        self.processor = SummaryProcessor()

        self._log_status("已清空所有数据")

    def _open_output_dir(self):
        """打开上次保存文件的目录"""
        if not self.last_saved_path:
            messagebox.showwarning("警告", "尚未保存文件，请先进行文件保存操作")
            return
        dir_path = os.path.dirname(self.last_saved_path)
        if os.path.exists(dir_path):
            os.startfile(dir_path)
            self._log_status(f"已打开输出目录：{dir_path}")
        else:
            messagebox.showerror("错误", f"目录不存在：{dir_path}")

    def _check_data(self):
        """核对数据：对比待核对表与基准表的金额"""
        # 1. 检查基准表（使用专门的reference_path，避免被核对结果覆盖）
        if not self.reference_path or not os.path.exists(self.reference_path):
            messagebox.showwarning("警告", "未找到基准表，请先完成数据处理并保存文件", parent=self.root)
            return

        self._log_status("=" * 50)
        self._log_status("开始核对数据...")

        # 2. 选择待核对文件
        check_file = filedialog.askopenfilename(
            title="选择待核对Excel文件",
            filetypes=[("Excel文件", "*.xlsx *.xls"), ("所有文件", "*.*")]
        )
        if not check_file:
            self._log_status("用户取消了选择文件")
            return

        self._log_status(f"选择待核对文件：{os.path.basename(check_file)}")

        # 3. 加载获取工作表
        try:
            wb = load_workbook(check_file, read_only=True)
            sheets = wb.sheetnames
            wb.close()
            self._log_status(f"读取到{len(sheets)}个工作表")
        except Exception as e:
            messagebox.showerror("错误", f"读取文件失败：{str(e)}", parent=self.root)
            self._log_status(f"错误：读取文件失败 {str(e)}")
            return

        if len(sheets) == 1:
            selected_sheet = sheets[0]
            self._log_status(f"自动选择工作表：{selected_sheet}")
        else:
            selected_sheet = self._show_sheet_selection_dialog(sheets)
            if not selected_sheet:
                self._log_status("用户取消了选择工作表")
                return
            self._log_status(f"选择工作表：{selected_sheet}")

        # 4. 使用核心模块快速加载（只读需要的三列，自动检测表头行）
        self._log_status("正在加载工作表数据（快速模式）...")
        df, month_col, person_col, amount_col, msg = DataChecker.load_check_file(check_file, selected_sheet)
        if df is None:
            messagebox.showerror("错误", msg, parent=self.root)
            self._log_status(f"错误：{msg}")
            return
        self._log_status(f"加载成功，{msg}")
        self._log_status(f"识别到列：月份={month_col}, 介绍人={person_col}, 金额={amount_col}")

        # 5. 加载基准表
        checker = DataChecker(self.reference_path)
        self._log_status("正在加载基准表...")
        success, msg = checker.load_reference()
        if not success:
            messagebox.showerror("错误", msg, parent=self.root)
            self._log_status(f"错误：{msg}")
            return
        self._log_status(msg)

        # 6. 获取唯一值
        self._log_status("正在检测月份和介绍人...")
        months = checker.get_unique_months(df, month_col)
        persons = checker.get_unique_persons(df, person_col)

        self._log_status(f"检测到月份：{', '.join(months) if months else '无'}")
        self._log_status(f"检测到介绍人：{len(persons)}个")

        if not months:
            messagebox.showwarning("警告", "未检测到任何月份数据，请检查'月份'列内容格式是否为'X月'或'X-Y月'", parent=self.root)
            return
        if not persons:
            messagebox.showwarning("警告", "未检测到任何介绍人数据", parent=self.root)
            return

        # 6.5 计算每个月份对应的介绍人映射（用于弹窗联动，一次遍历优化）
        month_person_map = {m: set() for m in months}
        for val, person in zip(df[month_col], df[person_col]):
            if pd.isna(val) or pd.isna(person):
                continue
            val_str = str(val).strip()
            person_str = str(person).strip()
            if not person_str:
                continue
            for m in months:
                if re.search(rf'(?<!\d){re.escape(m)}(?!\d)', val_str):
                    month_person_map[m].add(person_str)
        month_person_map = {m: sorted(list(ps)) for m, ps in month_person_map.items()}

        # 7. 弹窗选择月份和介绍人（传入月份-介绍人映射，实现联动）
        self._log_status("打开条件选择弹窗...")
        selected_month, selected_persons = self._show_check_condition_dialog(months, month_person_map)
        if selected_month is None:
            self._log_status("用户取消了条件选择")
            return

        self._log_status(f"已选择月份：{selected_month}")
        self._log_status(f"已选择介绍人：{', '.join(selected_persons)}")

        # 8. 执行核对
        self._log_status("执行核对...")
        success, msg, result_df = checker.check(
            df, selected_month, selected_persons, month_col, person_col, amount_col
        )
        if not success:
            messagebox.showerror("错误", msg, parent=self.root)
            self._log_status(f"错误：{msg}")
            return

        self._log_status(msg)

        # 9. 显示结果摘要
        checked = len(result_df[result_df['状态'] == '已核对'])
        pending = len(result_df[result_df['状态'] == '待核对'])
        unmatched = len(result_df[result_df['状态'] == '未匹配'])
        self._log_status(f"核对结果：已核对{checked}条，待核对{pending}条，未匹配{unmatched}条")

        # 10. 保存结果
        self._save_check_result(result_df)

    def _show_sheet_selection_dialog(self, sheets):
        """显示工作表选择弹窗"""
        dialog = tk.Toplevel(self.root)
        dialog.title("选择工作表")
        dialog.transient(self.root)
        dialog.grab_set()
        dialog.resizable(False, False)

        frame = tk.Frame(dialog)
        frame.pack(padx=20, pady=15)

        tk.Label(frame, text="请选择工作表：", font=('Microsoft YaHei UI', 10)).pack(anchor=tk.W, pady=(0, 8))

        sheet_var = tk.StringVar(value=sheets[0] if sheets else "")
        combo = ttk.Combobox(frame, textvariable=sheet_var, values=sheets, state='readonly', width=30)
        combo.pack()

        btn_frame = tk.Frame(dialog)
        btn_frame.pack(pady=(0, 15))

        result = [None]

        def on_confirm():
            result[0] = sheet_var.get()
            dialog.destroy()

        def on_cancel():
            dialog.destroy()

        ttk.Button(btn_frame, text="确定", command=on_confirm, width=10).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="取消", command=on_cancel, width=10).pack(side=tk.LEFT, padx=5)

        dialog.update_idletasks()
        self.center_dialog(dialog, dialog.winfo_reqwidth(), dialog.winfo_reqheight())
        self.root.wait_window(dialog)
        return result[0]

    def _merge_to_original(self):
        """将已生成的汇总表和核对结果追加到指定原始文件中（保留格式）"""
        has_summary = self.reference_path and os.path.exists(self.reference_path)
        has_check = self.check_result_path and os.path.exists(self.check_result_path)

        if not has_summary and not has_check:
            messagebox.showwarning("警告", "没有可合并的表格，请先完成数据处理或核对", parent=self.root)
            return

        target_path = filedialog.askopenfilename(
            title="选择要追加到的原始文件",
            filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")]
        )
        if not target_path:
            self._log_status("用户取消了合并操作")
            return

        try:
            target_wb = load_workbook(target_path)
            existing_sheets = set(target_wb.sheetnames)

            def get_unique_name(base_name):
                if base_name not in existing_sheets:
                    return base_name
                i = 1
                while f"{base_name}_{i}" in existing_sheets:
                    i += 1
                return f"{base_name}_{i}"

            merged = []
            if has_summary:
                name = get_unique_name('供应商拆分汇总')
                copy_sheet_to_workbook(self.reference_path, target_wb, name)
                existing_sheets.add(name)
                merged.append(name)

            if has_check:
                name = get_unique_name('核对结果')
                copy_sheet_to_workbook(self.check_result_path, target_wb, name)
                existing_sheets.add(name)
                merged.append(name)

            target_wb.save(target_path)
            target_wb.close()

            self._log_status(f"已合并到原始文件：{target_path}，新增工作表：{', '.join(merged)}")
            messagebox.showinfo("成功", f"已成功追加以下工作表到原始文件：\n" + "\n".join(merged), parent=self.root)

        except Exception as e:
            messagebox.showerror("错误", f"合并失败：{str(e)}", parent=self.root)
            self._log_status(f"错误：合并失败 {str(e)}")

    def _show_check_condition_dialog(self, months, month_person_map):
        """显示月份（单选）和介绍人（多选）选择弹窗，月份与介绍人联动"""
        dialog = tk.Toplevel(self.root)
        dialog.title("核对条件")
        dialog.grab_set()
        dialog.resizable(True, True)
        dialog.geometry("900x650")

        # 主容器
        main_container = tk.Frame(dialog)
        main_container.pack(fill=tk.BOTH, expand=True, padx=15, pady=15)
        main_container.columnconfigure(0, weight=0)
        main_container.columnconfigure(1, weight=1)
        main_container.rowconfigure(0, weight=1)

        # ===== 左侧：月份（单选） =====
        left_frame = ttk.LabelFrame(main_container, text="月份（单选）")
        left_frame.grid(row=0, column=0, sticky=(tk.N, tk.S, tk.W), padx=(0, 10))

        month_var = tk.StringVar(value=months[0] if months else "")

        # ===== 右侧：介绍人（多选，多列网格+滚动+查询+全选/全不选） =====
        all_persons = sorted(set(p for plist in month_person_map.values() for p in plist))
        right_frame = ttk.LabelFrame(main_container, text="介绍人（多选）")
        right_frame.grid(row=0, column=1, sticky=(tk.N, tk.S, tk.W, tk.E))
        right_frame.rowconfigure(1, weight=1)
        right_frame.columnconfigure(0, weight=1)

        # 工具栏：查询 + 全选 + 全不选
        toolbar = tk.Frame(right_frame)
        toolbar.grid(row=0, column=0, sticky=(tk.W, tk.E), pady=(0, 5))

        search_var = tk.StringVar()
        search_entry = ttk.Entry(toolbar, textvariable=search_var, width=18)
        search_entry.pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(toolbar, text="查询", width=6).pack(side=tk.LEFT, padx=(0, 10))
        ttk.Button(toolbar, text="全选", width=6).pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(toolbar, text="全不选", width=6).pack(side=tk.LEFT)

        # Canvas + 滚动条
        canvas = tk.Canvas(right_frame, highlightthickness=0)
        v_scroll = ttk.Scrollbar(right_frame, orient="vertical", command=canvas.yview)
        h_scroll = ttk.Scrollbar(right_frame, orient="horizontal", command=canvas.xview)
        canvas.configure(yscrollcommand=v_scroll.set, xscrollcommand=h_scroll.set)

        v_scroll.grid(row=1, column=1, sticky=(tk.N, tk.S))
        h_scroll.grid(row=2, column=0, sticky=(tk.W, tk.E))
        canvas.grid(row=1, column=0, sticky=(tk.N, tk.S, tk.W, tk.E))

        # Canvas 内 Frame
        inner_frame = tk.Frame(canvas)
        canvas.create_window((0, 0), window=inner_frame, anchor=tk.NW)

        # 介绍人：每列20个，横向排列（创建所有，但初始只显示当前月份的）
        person_vars = {}
        checkbuttons = {}
        items_per_col = 20
        for idx, p in enumerate(all_persons):
            var = tk.BooleanVar(value=False)
            person_vars[p] = var
            col = idx // items_per_col
            row = idx % items_per_col
            cb = ttk.Checkbutton(inner_frame, text=p, variable=var)
            cb.grid(row=row, column=col, sticky=tk.W, pady=2, padx=8)
            checkbuttons[p] = cb

        # 更新 Canvas 滚动区域
        inner_frame.update_idletasks()
        canvas.config(scrollregion=canvas.bbox("all"))

        # 联动过滤：根据当前月份 + 搜索条件，显示/隐藏并重新排列介绍人
        def filter_persons():
            current_month = month_var.get()
            valid_persons = set(month_person_map.get(current_month, []))
            query = search_var.get().strip().lower()
            visible_idx = 0
            for p, cb in checkbuttons.items():
                if p not in valid_persons:
                    cb.grid_remove()
                    continue
                if not query or query in p.lower():
                    cb.grid()
                    col = visible_idx // items_per_col
                    row = visible_idx % items_per_col
                    cb.grid_configure(row=row, column=col)
                    visible_idx += 1
                else:
                    cb.grid_remove()
            inner_frame.update_idletasks()
            canvas.config(scrollregion=canvas.bbox("all"))
            right_frame.configure(text=f"介绍人（多选，{current_month}，显示{visible_idx}/{len(valid_persons)}个）")

        # 月份切换回调
        def on_month_change():
            filter_persons()

        # 创建月份单选按钮（绑定联动回调）
        for m in months:
            ttk.Radiobutton(left_frame, text=m, variable=month_var, value=m, command=on_month_change).pack(anchor=tk.W, pady=3, padx=10)

        # 全选 / 全不选（只影响当前月份下可见的介绍人）
        def select_all():
            current_month = month_var.get()
            valid_persons = set(month_person_map.get(current_month, []))
            query = search_var.get().strip().lower()
            for p, var in person_vars.items():
                if p in valid_persons and (not query or query in p.lower()):
                    var.set(True)

        def deselect_all():
            current_month = month_var.get()
            valid_persons = set(month_person_map.get(current_month, []))
            query = search_var.get().strip().lower()
            for p, var in person_vars.items():
                if p in valid_persons and (not query or query in p.lower()):
                    var.set(False)

        # 配置工具栏按钮命令
        for widget in toolbar.winfo_children():
            if isinstance(widget, ttk.Button):
                text = widget.cget("text")
                if text == "查询":
                    widget.configure(command=filter_persons)
                elif text == "全选":
                    widget.configure(command=select_all)
                elif text == "全不选":
                    widget.configure(command=deselect_all)

        search_entry.bind("<Return>", lambda e: filter_persons())

        # 鼠标滚轮绑定（进入 Canvas 区域时启用）
        def _on_mousewheel(event):
            canvas.yview_scroll(-1 if event.delta > 0 else 1, "units")

        def _bind_wheel(event=None):
            canvas.bind_all("<MouseWheel>", _on_mousewheel)

        def _unbind_wheel(event=None):
            canvas.unbind_all("<MouseWheel>")

        canvas.bind("<Enter>", _bind_wheel)
        canvas.bind("<Leave>", _unbind_wheel)

        # 初始化：显示默认月份的介绍人
        if months:
            on_month_change()

        # ===== 底部按钮（居中） =====
        btn_frame = tk.Frame(dialog)
        btn_frame.pack(pady=(10, 15))

        result = {"month": None, "persons": None}

        def on_confirm():
            selected_month = month_var.get()
            valid_persons = set(month_person_map.get(selected_month, []))
            selected_persons = [p for p, v in person_vars.items() if v.get() and p in valid_persons]

            if not selected_month:
                messagebox.showwarning("警告", "请选择一个月份", parent=dialog)
                return
            if not selected_persons:
                messagebox.showwarning("警告", "请至少选择一个介绍人", parent=dialog)
                return

            result["month"] = selected_month
            result["persons"] = selected_persons
            _unbind_wheel()
            dialog.destroy()

        def on_cancel():
            _unbind_wheel()
            dialog.destroy()

        ttk.Button(btn_frame, text="确定", command=on_confirm, width=12).pack(side=tk.LEFT, padx=6)
        ttk.Button(btn_frame, text="取消", command=on_cancel, width=12).pack(side=tk.LEFT, padx=6)

        # 居中显示
        dialog.update_idletasks()
        self.center_dialog(dialog, dialog.winfo_reqwidth(), dialog.winfo_reqheight())
        self.root.wait_window(dialog)

        return result["month"], result["persons"]

    def _save_check_result(self, result_df):
        """保存核对结果到Excel（调用核心模块的save_result）"""
        file_path = filedialog.asksaveasfilename(
            title="保存核对结果",
            defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")],
            initialfile="核对结果"
        )

        if not file_path:
            self._log_status("用户取消了保存核对结果")
            return

        success, msg = DataChecker.save_result(result_df, file_path)
        if success:
            self.last_saved_path = file_path
            self.check_result_path = file_path
            self._log_status(msg)
            messagebox.showinfo("成功", msg, parent=self.root)
        else:
            messagebox.showerror("错误", msg, parent=self.root)
            self._log_status(f"错误：{msg}")

    def _show_selection_dialog(self, payment_types, months, amount_columns):
        """显示打款类型、月份和打款金额列选择对话框"""
        dialog = tk.Toplevel(self.root)
        dialog.title("选择数据")
        dialog.transient(self.root)
        dialog.grab_set()
        dialog.resizable(False, False)

        # 主容器
        main_container = tk.Frame(dialog)
        main_container.pack(fill=tk.BOTH, expand=True, padx=15, pady=15)

        # 上方区域：左侧打款类型 + 右侧月份
        top_frame = tk.Frame(main_container)
        top_frame.pack(fill=tk.BOTH, expand=True)

        # 左侧：打款类型（多选，默认不选中）
        if payment_types:
            left_frame = ttk.LabelFrame(top_frame, text="打款类型")
            left_frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 8))

            type_vars = {}
            for pt in payment_types:
                var = tk.BooleanVar(value=False)  # 默认不选中
                type_vars[pt] = var
                ttk.Checkbutton(left_frame, text=pt, variable=var).pack(anchor=tk.W, pady=3, padx=8)
        else:
            type_vars = {}

        # 右侧：月份（多选，默认不选中）
        if months:
            right_frame = ttk.LabelFrame(top_frame, text="月份")
            right_frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(8, 0))

            month_vars = {}
            for m in months:
                var = tk.BooleanVar(value=False)  # 默认不选
                month_vars[m] = var
                ttk.Checkbutton(right_frame, text=m, variable=var).pack(anchor=tk.W, pady=3, padx=8)
        else:
            month_vars = {}

        # 中间区域：打款金额列选择（有多列时才显示）
        amount_var = tk.IntVar(value=-1)  # -1 表示未选择
        if len(amount_columns) >= 2:
            amount_frame = ttk.LabelFrame(main_container, text="打款金额列")
            amount_frame.pack(fill=tk.X, pady=(10, 0))

            for idx, col_letter in amount_columns:
                ttk.Radiobutton(
                    amount_frame,
                    text=f"打款金额（{col_letter}列）",
                    variable=amount_var,
                    value=idx
                ).pack(anchor=tk.W, pady=2, padx=8)

        # 按钮区域 - 底部居中
        btn_frame = tk.Frame(dialog)
        btn_frame.pack(pady=(10, 15))

        result = {"types": None, "months": None, "amount_col_idx": None}

        def on_confirm():
            selected_types = [t for t, v in type_vars.items() if v.get()]
            selected_months = [m for m, v in month_vars.items() if v.get()]

            if payment_types and not selected_types:
                messagebox.showwarning("警告", "请至少选择一个打款类型", parent=dialog)
                return
            if months and not selected_months:
                messagebox.showwarning("警告", "请至少选择一个月份", parent=dialog)
                return
            if len(amount_columns) >= 2 and amount_var.get() == -1:
                messagebox.showwarning("警告", "请选择一个打款金额列", parent=dialog)
                return

            result["types"] = selected_types if payment_types else None
            result["months"] = selected_months if months else None
            if len(amount_columns) >= 2:
                result["amount_col_idx"] = amount_var.get()
            elif len(amount_columns) == 1:
                result["amount_col_idx"] = amount_columns[0][0]
            dialog.destroy()

        def on_cancel():
            dialog.destroy()

        ttk.Button(btn_frame, text="确定", command=on_confirm, width=10).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="取消", command=on_cancel, width=10).pack(side=tk.LEFT, padx=5)

        # 自适应大小并居中
        dialog.update_idletasks()
        width = dialog.winfo_reqwidth()
        height = dialog.winfo_reqheight()
        self.center_dialog(dialog, width, height)

        self.root.wait_window(dialog)
        return result["types"], result["months"], result["amount_col_idx"]

    def center_dialog(self, dialog, width, height):
        """将对话框居中显示在主窗口上"""
        self.root.update_idletasks()
        root_x = self.root.winfo_x()
        root_y = self.root.winfo_y()
        root_width = self.root.winfo_width()
        root_height = self.root.winfo_height()

        x = root_x + (root_width - width) // 2
        y = root_y + (root_height - height) // 2

        dialog.geometry(f"{width}x{height}+{x}+{y}")


def main():
    """主函数"""
    root = tk.Tk()
    app = SummaryProcessorApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
