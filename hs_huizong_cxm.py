"""
Excel汇总处理工具
根据备注2列内容缩写，汇总打款金额，生成供应商汇总表
"""
import pandas as pd
import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import re
import traceback
from typing import Dict, List, Set, Tuple, Any, Optional
import warnings
warnings.filterwarnings('ignore')


# ==================== 公共工具函数 ====================

def apply_excel_styles(ws: openpyxl.worksheet.worksheet.Worksheet, df: pd.DataFrame) -> None:
    """应用统一的Excel样式：表头蓝色背景、数据行居中边框、自动列宽"""
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col in range(1, df.shape[1] + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border

    for row_idx in range(2, df.shape[0] + 2):
        for col_idx in range(1, df.shape[1] + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border

    for col in range(1, df.shape[1] + 1):
        max_length = 0
        column_letter = get_column_letter(col)
        for row_idx in range(1, df.shape[0] + 2):
            cell = ws.cell(row=row_idx, column=col)
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        adjusted_width = min(max_length + 2, 30)
        ws.column_dimensions[column_letter].width = max(adjusted_width, 10)


def match_month_in_text(month: str, text: str) -> bool:
    """精确匹配月份字符串，避免'1月'匹配到'12月'"""
    if pd.isna(text):
        return False
    pattern = rf'(?<!\d){re.escape(month)}(?!\d)'
    return bool(re.search(pattern, str(text).strip()))


def copy_sheet_to_workbook(source_path: str, target_wb, target_name: str) -> None:
    """跨工作簿复制工作表，保留值、格式、列宽、行高、合并单元格"""
    from copy import copy as copy_obj
    source_wb = load_workbook(source_path)
    source_ws = source_wb.active
    target_ws = target_wb.create_sheet(target_name)

    for row in source_ws.iter_rows():
        for cell in row:
            new_cell = target_ws.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                new_cell.font = copy_obj(cell.font)
                new_cell.border = copy_obj(cell.border)
                new_cell.fill = copy_obj(cell.fill)
                new_cell.number_format = copy_obj(cell.number_format)
                new_cell.protection = copy_obj(cell.protection)
                new_cell.alignment = copy_obj(cell.alignment)

    for col_letter, col_dim in source_ws.column_dimensions.items():
        if col_dim.width:
            target_ws.column_dimensions[col_letter].width = col_dim.width

    for row_num, row_dim in source_ws.row_dimensions.items():
        if row_dim.height:
            target_ws.row_dimensions[row_num].height = row_dim.height

    for merged_range in source_ws.merged_cells.ranges:
        target_ws.merge_cells(str(merged_range))

    source_wb.close()


class SummaryProcessor:
    """
    汇总处理核心类
    功能包括：
    - 读取Excel文件和工作表
    - 检测"备注2"列并应用缩写规则
    - 按供应商和公司汇总打款金额
    - 生成汇总报表
    """
    
    # 空白字符正则表达式（包括零宽空格、不间断空格等）
    BLANK_PATTERN = r'[\s\u200b\u200c\u200d\u00a0\u3000]+'
    
    def __init__(self):
        """初始化处理器"""
        self.current_file_path: Optional[str] = None
        self.workbook: Optional[openpyxl.Workbook] = None
        self.sheet_names: List[str] = []
        self.current_sheet: Optional[str] = None
        self.data_frame: Optional[pd.DataFrame] = None
        self.output_data: Optional[pd.DataFrame] = None
        
    def load_file(self, file_path: str) -> Tuple[bool, str, List[str]]:
        """
        加载Excel文件
        
        Args:
            file_path: Excel文件路径
            
        Returns:
            (success, message, sheet_names): 是否成功、消息、工作表列表
        """
        try:
            self.current_file_path = file_path
            wb = load_workbook(file_path, data_only=True)
            self.sheet_names = wb.sheetnames
            wb.close()  # 及时释放资源，后续用pd.read_excel读取数据
            return True, f"成功读取文件：{os.path.basename(file_path)}", self.sheet_names
        except Exception as e:
            return False, f"读取文件失败：{str(e)}", []
    
    def load_sheet(self, sheet_name: str) -> Tuple[bool, str]:
        """
        加载指定工作表
        
        Args:
            sheet_name: 工作表名称
            
        Returns:
            (success, message): 是否成功、消息
        """
        try:
            self.current_sheet = sheet_name
            self.data_frame = pd.read_excel(self.current_file_path, sheet_name=sheet_name)
            return True, f"成功加载工作表：{sheet_name}"
        except Exception as e:
            return False, f"加载工作表失败：{str(e)}"
    
    def _clean_blank_chars(self, text: str) -> str:
        """移除所有空白字符（包括不可见字符）"""
        return re.sub(self.BLANK_PATTERN, '', str(text))
    
    def _build_unique_key(self, row, has_payee_name: bool):
        """构建供应商唯一键"""
        supplier = row['供应商名称']
        if pd.isna(supplier) or supplier == '':
            return None
        supplier_str = str(supplier).strip()
        
        if has_payee_name and pd.notna(row['收款姓名']) and str(row['收款姓名']).strip() != '':
            payee_name = str(row['收款姓名']).strip()
            return (supplier_str, payee_name)
        return supplier_str
    
    def _match_supplier(self, row_supplier, target_supplier: str) -> bool:
        """匹配供应商名称（支持空白字符容错）"""
        if pd.isna(row_supplier):
            return False
        row_clean = str(row_supplier).strip()
        if row_clean == target_supplier:
            return True
        return self._clean_blank_chars(row_clean) == target_supplier
    
    def _match_payee(self, row_payee, target_payee: str, has_payee_name: bool) -> bool:
        """匹配收款姓名（支持空白字符容错）"""
        if not has_payee_name:
            return True
        if pd.notna(row_payee):
            row_clean = str(row_payee).strip()
            if row_clean == target_payee:
                return True
            return self._clean_blank_chars(row_clean) == target_payee
        return target_payee == ''
    
    def _abbreviate_company_name(self, company_name: str) -> Optional[str]:
        """
        根据缩写规则处理公司名称
        
        Args:
            company_name: 原始公司名称
            
        Returns:
            缩写后的公司名称
        """
        # 空白或空值返回"（空白）"
        if pd.isna(company_name) or company_name == '':
            return "（空白）"
        
        # 使用正则表达式移除所有空白字符（包括不可见字符）
        company_str = str(company_name)
        cleaned_str = re.sub(self.BLANK_PATTERN, '', company_str)
        
        if cleaned_str == '':
            return "（空白）"
        
        company_str = company_str.strip()
        
        # 规则1: "刘先锋现金卡支付" → "刘先锋现金卡"
        if company_str == "刘先锋现金卡支付":
            return "刘先锋现金卡"
        
        # 规则2: "XXX对公打款"或"XXX代发打款" → "XXX"
        for suffix in ["对公打款", "代发打款"]:
            if company_str.endswith(suffix):
                return company_str.replace(suffix, "")
        
        # 不符合规则，返回原值
        return company_str
    
    def _extract_payment_type(self, remark2: str) -> Optional[str]:
        """从备注2值中提取打款类型"""
        if pd.isna(remark2):
            return None
        val_str = str(remark2).strip()
        
        if "刘先锋现金卡" in val_str:
            return "刘先锋现金卡"
        if val_str.endswith("对公打款"):
            return "对公打款"
        if val_str.endswith("代发打款"):
            return "代发打款"
        return None
    
    def detect_payment_types(self) -> List[str]:
        """检测备注2列中存在的打款类型"""
        if self.data_frame is None:
            return []
        if "备注2" not in self.data_frame.columns:
            return []
        
        types_found = set()
        for val in self.data_frame["备注2"].dropna():
            pt = self._extract_payment_type(val)
            if pt:
                types_found.add(pt)
        
        # 按固定顺序排序
        order = {"刘先锋现金卡": 0, "对公打款": 1, "代发打款": 2}
        return sorted(list(types_found), key=lambda x: order.get(x, 99))
    
    def detect_months(self) -> List[str]:
        """检测备注1列中存在的月份（支持X月和X-Y月格式，如1月、1-2月）"""
        if self.data_frame is None:
            return []
        if "备注1" not in self.data_frame.columns:
            return []
        
        months_found = set()
        for val in self.data_frame["备注1"].dropna():
            val_str = str(val).strip()
            # 精确提取单月份，避免"12月"被拆成"1月"和"2月"
            found = re.findall(r'(?<!\d)(\d{1,2}月)(?!\d)', val_str)
            for m in found:
                try:
                    num = int(m.replace('月', ''))
                    if 1 <= num <= 12:
                        months_found.add(m)
                except ValueError:
                    pass
            
            # 提取范围月份，如"1-2月"、"3-5月"
            range_found = re.findall(r'(?<!\d)(\d{1,2}-\d{1,2}月)(?!\d)', val_str)
            for m in range_found:
                try:
                    start, end = m.replace('月', '').split('-')
                    start_num = int(start)
                    end_num = int(end)
                    if 1 <= start_num <= 12 and 1 <= end_num <= 12:
                        months_found.add(m)
                except ValueError:
                    pass
        
        # 排序：单月份在前，范围月份在后，均按起始数字排序
        def sort_key(x):
            if '-' in x:
                start = int(x.split('-')[0])
                return (1, start)
            return (0, int(x.replace('月', '')))
        
        return sorted(list(months_found), key=sort_key)
    
    def detect_amount_columns(self) -> List[Tuple[int, str]]:
        """检测所有名为'打款金额'的列，返回列索引和列字母标识"""
        if self.data_frame is None:
            return []
        
        amount_cols = [(i, get_column_letter(i + 1)) for i, col in enumerate(self.data_frame.columns) if '打款金额' in str(col)]
        return amount_cols
    
    def process_data(self, selected_types: Optional[List[str]] = None,
                     selected_months: Optional[List[str]] = None,
                     selected_amount_col_idx: Optional[int] = None) -> Tuple[bool, str]:
        """
        处理数据：缩写、汇总、生成报表
        
        Returns:
            (success, message): 是否成功、消息
        """
        try:
            if self.data_frame is None:
                return False, "请先选择工作表"
            
            df = self.data_frame
            
            # 检测"备注2"列
            if "备注2" not in df.columns:
                return False, "未找到'备注2'列"
            
            # 使用用户选择的打款金额列
            if selected_amount_col_idx is not None:
                amount_col_idx = selected_amount_col_idx
            else:
                return False, "未选择或检测到'打款金额'列"
            
            # 应用缩写规则
            df['公司缩写'] = df['备注2'].apply(self._abbreviate_company_name)
            
            # 根据用户选择的打款类型筛选
            if selected_types:
                def match_payment_type(remark2):
                    if pd.isna(remark2):
                        return False
                    val_str = str(remark2).strip()
                    for pt in selected_types:
                        if pt == "刘先锋现金卡" and "刘先锋现金卡" in val_str:
                            return True
                        if pt == "对公打款" and val_str.endswith("对公打款"):
                            return True
                        if pt == "代发打款" and val_str.endswith("代发打款"):
                            return True
                    return False
                
                df = df[df['备注2'].apply(match_payment_type)]
                
                if df.empty:
                    return False, "没有符合所选打款类型的数据"
            
            # 根据用户选择的月份筛选（精确字符串匹配，不展开范围）
            if selected_months:
                if "备注1" not in df.columns:
                    return False, "未找到'备注1'列，无法进行月份筛选"
                
                df = df[df['备注1'].apply(
                    lambda val: any(match_month_in_text(m, val) for m in selected_months)
                )]
                
                if df.empty:
                    return False, f"没有符合所选月份的数据"
            
            # 获取所有唯一的供应商名称
            if "供应商名称" not in df.columns:
                return False, "未找到'供应商名称'列"
            
            # 获取所有唯一的缩写公司名称
            unique_companies = df['公司缩写'].unique().tolist()
            unique_companies = [c for c in unique_companies if c is not None and c != '']
            
            if not unique_companies:
                return False, "没有有效的公司缩写数据"
            
            # 确保"（空白）"列在最后（在"总计"列之前）
            if "（空白）" in unique_companies:
                unique_companies.remove("（空白）")
                unique_companies.append("（空白）")
            
            # 检查是否有"收款姓名"列
            has_payee_name = "收款姓名" in df.columns
            
            # 第一步：按供应商+收款姓名分组收集空白行和非空白行
            blank_groups = {}  # {key: [rows]}
            non_blank_groups = {}  # {key: {company: total_amount}}
            
            for _, row in df.iterrows():
                unique_key = self._build_unique_key(row, has_payee_name)
                if unique_key is None:
                    continue
                
                remark2 = row['备注2']
                is_blank = pd.isna(remark2) or self._clean_blank_chars(remark2) == ''
                
                if is_blank:
                    # 收集空白行
                    if unique_key not in blank_groups:
                        blank_groups[unique_key] = []
                    blank_groups[unique_key].append(row)
                else:
                    # 汇总非空白行的金额
                    company_abbrev = row['公司缩写']
                    if unique_key not in non_blank_groups:
                        non_blank_groups[unique_key] = {}
                    
                    try:
                        if len(row) > amount_col_idx:
                            amount = row.iloc[amount_col_idx]
                            if pd.notna(amount):
                                amount_value = float(amount)
                                if company_abbrev in non_blank_groups[unique_key]:
                                    non_blank_groups[unique_key][company_abbrev] += amount_value
                                else:
                                    non_blank_groups[unique_key][company_abbrev] = amount_value
                    except (ValueError, TypeError, IndexError):
                        pass
            
            # 第二步：初始化汇总数据字典
            all_keys = list(set(list(blank_groups.keys()) + list(non_blank_groups.keys())))
            summary_data = {}
            for key in all_keys:
                if isinstance(key, tuple):
                    supplier_str, payee_name = key
                    summary_data[key] = {
                        'supplier': supplier_str,
                        'payee': payee_name,
                        'amounts': {company: None for company in unique_companies}
                    }
                else:
                    summary_data[key] = {
                        'supplier': key,
                        'payee': '',
                        'amounts': {company: None for company in unique_companies}
                    }
            
            # 第三步：填充非空白行的汇总数据
            for key, company_amounts in non_blank_groups.items():
                if key not in summary_data:
                    continue
                for company, amount in company_amounts.items():
                    if company != "（空白）" and company in summary_data[key]['amounts']:
                        summary_data[key]['amounts'][company] = amount
            
            # 第四步：处理空白行，汇总到"（空白）"列
            for key, rows in blank_groups.items():
                if key not in summary_data:
                    continue
                
                supplier_str = summary_data[key]['supplier']
                payee_name = summary_data[key]['payee']
                blank_total = 0.0
                
                for row in rows:
                    # 匹配供应商
                    if not self._match_supplier(row['供应商名称'], supplier_str):
                        continue
                    
                    # 匹配收款姓名
                    if not self._match_payee(row.get('收款姓名'), payee_name, has_payee_name):
                        continue
                    
                    # 累加金额
                    try:
                        if len(row) > amount_col_idx:
                            amount = row.iloc[amount_col_idx]
                            if pd.notna(amount):
                                amount_value = float(amount)
                                blank_total += amount_value
                    except (ValueError, TypeError, IndexError):
                        pass
                
                # 设置空白列总额
                if blank_total > 0 and "（空白）" in summary_data[key]['amounts']:
                    summary_data[key]['amounts']["（空白）"] = blank_total
            
            # 创建输出DataFrame
            output_columns = ['供应商名称'] + unique_companies
            output_rows = []
            
            for key in summary_data:
                row_data = {'供应商名称': summary_data[key]['supplier']}
                for company in unique_companies:
                    row_data[company] = summary_data[key]['amounts'][company]
                # 保存收款姓名，用于后续处理"公司"列
                row_data['_payee_name'] = summary_data[key]['payee']
                output_rows.append(row_data)
            
            self.output_data = pd.DataFrame(output_rows, columns=output_columns + ['_payee_name'])
            
            # 添加"总计"列
            company_cols = [col for col in output_columns if col != '供应商名称']
            self.output_data['总计'] = self.output_data[company_cols].sum(axis=1)
            
            # 添加"公司"列数据
            self._add_company_column()
            
            # 重新排列列顺序
            # 只在"总计"列和"公司"列之间保留一个空列
            final_columns = ['供应商名称'] + unique_companies + ['总计', '', '公司']
            self.output_data = self.output_data.reindex(columns=final_columns)
            
            # 获取H列名称
            #h_col_name = df.columns[amount_col_idx] if amount_col_idx < len(df.columns) else f'第{amount_col_idx + 1}列'
            
            return True, f"处理完成！共{len(self.output_data)}个供应商，{len(unique_companies)}个公司"
            
        except Exception as e:
            import traceback
            error_msg = f"数据处理失败：{str(e)}\n{traceback.format_exc()}"
            return False, error_msg
    
    def _add_company_column(self) -> None:
        """
        添加"公司"列数据
        
        从"收款姓名"列提取符合条件的值：
        - 包含"刘先锋现金卡支付"
        - 包含"个体工商户"
        - 包含"有限公司"
        - 包含"分公司"
        
        注意：现在每个供应商+收款姓名的组合是一行，所以"公司"列显示对应的收款姓名
        """
        if self.output_data is None or '_payee_name' not in self.output_data.columns:
            return
        
        # 向量化处理公司列
        keywords = ["刘先锋现金卡支付", "个体工商户", "有限公司", "分公司"]
        
        def extract_company(payee_name):
            if payee_name and any(keyword in payee_name for keyword in keywords):
                return payee_name
            return ''
        
        self.output_data['公司'] = self.output_data['_payee_name'].apply(extract_company)
        
        # 删除临时列
        if '_payee_name' in self.output_data.columns:
            del self.output_data['_payee_name']
    

    
    def _extract_date_from_sheet_name(self, sheet_name: str, include_timestamp: bool = False) -> str:
        """
        从工作表名称中提取日期部分

        Args:
            sheet_name: 工作表名称
            include_timestamp: 是否在无法提取日期时附加时间戳

        Returns:
            提取的日期字符串或清理后的名称
        """
        if not sheet_name:
            if include_timestamp:
                return f"供应商汇总_{pd.Timestamp.now().strftime('%Y%m%d_%H%M%S')}"
            return "供应商汇总"

        pattern = r'(\d{4}年\d{1,2}月)'
        match = re.search(pattern, sheet_name)

        if match:
            return match.group(1)
        else:
            clean_name = re.sub(r'[<>:"/\\|?*]', '', sheet_name)
            if clean_name:
                return clean_name
            if include_timestamp:
                return f"供应商汇总_{pd.Timestamp.now().strftime('%Y%m%d_%H%M%S')}"
            return "供应商汇总"
    
    def save_to_excel(self, save_path: str, selected_month: Optional[str] = None) -> Tuple[bool, str]:
        """
        保存到Excel文件
        
        Args:
            save_path: 保存路径
            selected_month: 用户选择的月份，用于生成工作表标题
            
        Returns:
            (success, message): 是否成功、消息
        """
        try:
            if self.output_data is None:
                return False, "没有数据可保存"
            
            # 使用openpyxl创建新工作簿
            wb = openpyxl.Workbook()
            ws = wb.active
            
            # 生成工作表标题：年份（从工作表名提取）+ 选择的月份
            if selected_month:
                year_match = re.search(r'(\d{4}年)', str(self.current_sheet))
                year = year_match.group(1) if year_match else ""
                sheet_title = f"{year}{selected_month}" if year else selected_month
            else:
                sheet_title = self._extract_date_from_sheet_name(self.current_sheet)
            ws.title = sheet_title
            
            # 写入表头（第1行）
            columns = list(self.output_data.columns)
            for col_idx, col_name in enumerate(columns, start=1):
                ws.cell(row=1, column=col_idx, value=col_name)
            
            # 写入数据（从第2行开始）
            for row_idx, row in enumerate(self.output_data.itertuples(index=False), start=2):
                for col_idx, value in enumerate(row, start=1):
                    # 如果数值为0或空字符串，则显示为空白
                    if isinstance(value, (int, float)) and value == 0:
                        ws.cell(row=row_idx, column=col_idx, value=None)
                    elif pd.isna(value):
                        ws.cell(row=row_idx, column=col_idx, value=None)
                    elif value == '':
                        ws.cell(row=row_idx, column=col_idx, value=None)
                    else:
                        ws.cell(row=row_idx, column=col_idx, value=value)
            
            # 设置样式（调用公共样式函数）
            apply_excel_styles(ws, self.output_data)
            
            # 保存
            wb.save(save_path)
            return True, f"文件已保存到：{save_path}"
            
        except Exception as e:
            import traceback
            error_msg = f"保存失败：{str(e)}\n{traceback.format_exc()}"
            return False, error_msg
    
    def _apply_excel_styles(self, ws: openpyxl.worksheet.worksheet.Worksheet, df: pd.DataFrame) -> None:
        """
        应用Excel样式（已提取为公共函数，此处保留向后兼容）
        
        Args:
            ws: 工作表对象
            df: 数据框
        """
        apply_excel_styles(ws, df)


class DataChecker:
    """数据核对器：对比待核对表与基准表的金额"""

    def __init__(self, reference_path: str):
        self.reference_path = reference_path
        self.reference_data: Dict[str, float] = {}  # 供应商名称 -> 总计

    def load_reference(self) -> Tuple[bool, str]:
        """加载基准表数据（供应商名称 -> 总计），使用向量化操作替代iterrows"""
        try:
            df = pd.read_excel(self.reference_path)
            if '供应商名称' not in df.columns:
                return False, f"基准表未找到'供应商名称'列"
            if '总计' not in df.columns:
                return False, f"基准表未找到'总计'列"

            self.reference_df = df  # 保存原始DataFrame，用于后续提取供应商名称

            # 向量化处理：比iterrows快10-100倍
            valid_df = df.dropna(subset=['供应商名称']).copy()
            suppliers = valid_df['供应商名称'].astype(str).str.strip()
            totals = pd.to_numeric(valid_df['总计'], errors='coerce').fillna(0.0)
            self.reference_data = dict(zip(suppliers, totals))

            return True, f"加载基准表成功，共{len(self.reference_data)}条数据"
        except Exception as e:
            return False, f"加载基准表失败：{str(e)}"

    def detect_columns(self, df: pd.DataFrame) -> Tuple[Optional[str], Optional[str], Optional[str]]:
        """识别月份、介绍人、金额三列，返回(月份列, 介绍人列, 金额列)"""
        month_col = None
        person_col = None
        amount_col = None

        for col in df.columns:
            col_str = str(col).strip()
            if month_col is None and '月份' in col_str:
                month_col = col
            if person_col is None and '介绍人' in col_str:
                person_col = col
            if amount_col is None and '金额' in col_str:
                amount_col = col

        return month_col, person_col, amount_col

    def get_unique_months(self, df: pd.DataFrame, month_col: str) -> List[str]:
        """获取月份列的唯一值（支持X月和X-Y月格式）"""
        if month_col not in df.columns:
            return []
        months_found = set()
        for val in df[month_col].dropna():
            val_str = str(val).strip()
            # 单月份
            found = re.findall(r'(?<!\d)(\d{1,2}月)(?!\d)', val_str)
            for m in found:
                try:
                    num = int(m.replace('月', ''))
                    if 1 <= num <= 12:
                        months_found.add(m)
                except ValueError:
                    pass
            # 范围月份
            range_found = re.findall(r'(?<!\d)(\d{1,2}-\d{1,2}月)(?!\d)', val_str)
            for m in range_found:
                try:
                    start, end = m.replace('月', '').split('-')
                    if 1 <= int(start) <= 12 and 1 <= int(end) <= 12:
                        months_found.add(m)
                except ValueError:
                    pass

        def sort_key(x):
            if '-' in x:
                return (1, int(x.split('-')[0]))
            return (0, int(x.replace('月', '')))
        return sorted(list(months_found), key=sort_key)

    def get_unique_persons(self, df: pd.DataFrame, person_col: str) -> List[str]:
        """获取介绍人列的非空唯一值"""
        if person_col not in df.columns:
            return []
        values = df[person_col].dropna().astype(str).str.strip()
        values = values[values != '']
        valid_values = []
        for v in values.unique():
            if len(v) > 0 and not v.startswith('#') and v.lower() not in ['nan', 'none', 'null']:
                valid_values.append(v)
        return sorted(valid_values)

    def check(self, df: pd.DataFrame, selected_month: str, selected_persons: List[str],
              month_col: str, person_col: str, amount_col: str) -> Tuple[bool, str, Optional[pd.DataFrame]]:
        """执行核对：按介绍人分组累加金额，与基准表总计对比"""
        try:
            # 筛选月份（使用公共月份匹配函数）
            df_filtered = df[df[month_col].apply(lambda val: match_month_in_text(selected_month, val))]

            if df_filtered.empty:
                return False, f"月份'{selected_month}'没有匹配到数据", None

            results = []
            for person in selected_persons:
                person_df = df_filtered[df_filtered[person_col].astype(str).str.strip() == person]
                check_amount = person_df[amount_col].sum()

                base_total = self.reference_data.get(person)

                if base_total is not None:
                    if abs(float(check_amount) - base_total) < 0.01:
                        status = "已核对"
                    else:
                        status = "待核对"
                else:
                    status = "未匹配"

                # 从基准表原始DataFrame中提取对应的供应商名称
                supplier_name = ''
                if hasattr(self, 'reference_df') and self.reference_df is not None:
                    matched = self.reference_df[self.reference_df['供应商名称'].astype(str).str.strip() == person]
                    if not matched.empty:
                        supplier_name = str(matched.iloc[0]['供应商名称']).strip()

                results.append({
                    '介绍人': person,
                    '待核对金额总计': round(float(check_amount), 2),
                    '供应商名称': supplier_name,
                    '对公打款总计': round(base_total, 2) if base_total is not None else '',
                    '状态': status
                })

            result_df = pd.DataFrame(results)
            return True, f"核对完成，共{len(results)}条记录", result_df

        except Exception as e:
            import traceback
            return False, f"核对失败：{str(e)}\n{traceback.format_exc()}", None

    @staticmethod
    def load_check_file(file_path: str, sheet_name: str) -> Tuple[Optional[pd.DataFrame], Optional[str], Optional[str], Optional[str], str]:
        """使用openpyxl快速加载待核对数据，自动检测表头行，只读取需要的三列"""
        try:
            wb = load_workbook(file_path, read_only=True, data_only=True)
            ws = wb[sheet_name]

            # 自动检测表头行（前5行内查找包含"月份"、"介绍人"、"金额"的行）
            header_row = None
            headers = []
            for row_idx in range(1, min(6, ws.max_row + 1)):
                row_headers = [str(cell.value).strip() if cell.value else '' for cell in ws[row_idx]]
                has_month = any('月份' in h for h in row_headers)
                has_person = any('介绍人' in h for h in row_headers)
                has_amount = any('金额' in h for h in row_headers)
                if has_month and has_person and has_amount:
                    header_row = row_idx
                    headers = row_headers
                    break

            if header_row is None:
                wb.close()
                return None, None, None, None, "未在表头中找到'月份'、'介绍人'、'金额'列，请确认表头在前5行内"

            # 找到列索引
            month_col = person_col = amount_col = None
            month_idx = person_idx = amount_idx = None
            for i, h in enumerate(headers):
                if month_col is None and '月份' in h:
                    month_col = h
                    month_idx = i
                if person_col is None and '介绍人' in h:
                    person_col = h
                    person_idx = i
                if amount_col is None and '金额' in h:
                    amount_col = h
                    amount_idx = i

            if month_idx is None or person_idx is None or amount_idx is None:
                wb.close()
                return None, None, None, None, f"识别列失败，检测到的表头：{headers}"

            # 只读取三列数据
            data = []
            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                data.append({
                    month_col: row[month_idx] if month_idx < len(row) else None,
                    person_col: row[person_idx] if person_idx < len(row) else None,
                    amount_col: row[amount_idx] if amount_idx < len(row) else None,
                })

            df = pd.DataFrame(data)
            wb.close()
            return df, month_col, person_col, amount_col, f"共{len(df)}行，表头在第{header_row}行"

        except Exception as e:
            import traceback
            return None, None, None, None, f"加载失败：{str(e)}\n{traceback.format_exc()}"

    @staticmethod
    def save_result(result_df: pd.DataFrame, file_path: str) -> Tuple[bool, str]:
        """保存核对结果到Excel，应用统一样式"""
        try:
            with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
                result_df.to_excel(writer, sheet_name='核对结果', index=False)

                wb = writer.book
                ws = writer.sheets['核对结果']
                apply_excel_styles(ws, result_df)

            return True, f"核对结果已保存到：{file_path}"
        except Exception as e:
            import traceback
            return False, f"保存失败：{str(e)}\n{traceback.format_exc()}"
