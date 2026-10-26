# -*- coding: utf-8 -*-
"""hbgs_core.py - 合并个税核心逻辑 (VBA "个税汇总工具_项目提取_改版" Python 重写)

功能对应:
  VBA ListSubFolders/ListFileNames/汇总  ->  scan_excel_files + run_merge_collect
  VBA LoadFeeRules                       ->  load_fee_rules
  VBA ComputeFee_V2                      ->  compute_fee_v2
  VBA ClassifyRuleExact                  ->  classify_rule_exact
  VBA IsBankExempt                       ->  is_bank_exempt
  VBA ResolveAmount                      ->  resolve_amount
  VBA HandleDualUnitRule                 ->  handle_dual_unit_rule
  VBA CleanProjectName/RemoveSpaces/ExtractMonthFromFileName/NormalizeName -> 同名函数

依赖: openpyxl (必须)
"""

import os
import re
import threading

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


# ============================================================
# 文本工具函数
# ============================================================

def remove_spaces(s):
    """去除半角空格与全角空格(与 VBA RemoveSpaces 等价, 仅去空格不含其他字符)"""
    if s is None:
        return ""
    return re.sub(r"[\u3000 ]", "", str(s))


def normalize_name(s):
    """姓名归一化: 去空格, 全角数字/字母转半角 (与 VBA NormalizeName 等价)"""
    if s is None:
        return ""
    s = str(s).strip()
    out = []
    for ch in s:
        if ch == " " or ch == "\u3000":
            continue
        o = ord(ch)
        if 0xFF10 <= o <= 0xFF19:          # 全角数字
            out.append(chr(o - 0xFF10 + 0x30))
        elif 0xFF21 <= o <= 0xFF3A:        # 全角大写字母
            out.append(chr(o - 0xFF21 + 0x41))
        elif 0xFF41 <= o <= 0xFF5A:        # 全角小写字母
            out.append(chr(o - 0xFF41 + 0x61))
        else:
            out.append(ch)
    return "".join(out)


def clean_project_name(raw):
    """清洗"项目"长文本: 去年月/工资表/单位元等 (与 VBA CleanProjectName 等价)"""
    if raw is None:
        return ""
    s = str(raw).strip()
    if not s:
        return ""
    s = re.sub(r"^\d{2,4}年\d{1,2}月份?", "", s)
    s = re.sub(r"工资表", "", s)
    s = re.sub(r"单位：元", "", s)
    s = re.sub(r"\s*（", "（", s)
    s = re.sub(r"\s*\(", "(", s)
    s = re.sub(r"\s+", " ", s)
    return s.strip()


def extract_month_from_filename(fname):
    """从文件名提取月份数字 (与 VBA ExtractMonthFromFileName 等价)"""
    m = re.search(r"(\d+)月", fname)
    if not m:
        return ""
    return m.group(1)


# ============================================================
# 输出列布局 (2026-10-26 对齐 VBA「按列名映射重写版」/ Java HbgsService)
#   固定 19 列白名单 + 白名单外按邻列插入 + 忽略名单 + 尾部压缩
# ============================================================

_STD_COLUMNS = [
    "文件名", "项目", "姓名", "身份证", "入职时间", "离职时间", "应发工资", "个税",
    "预支扣款", "需扣回", "实发工资", "备注", "发薪单位", "联系电话", "社保", "公积金",
    "手续费", "个税申报所属期", "采用/不采用",
]

# 忽略名单: 源表出现这些列名时整列跳过(不输出、不占列位)
#   开户行   —— 手续费引擎从「员工信息-X月」表取
#   发放公司 —— 不输出, 但 M 类双单位规则仍用它判定发薪单位
#   电话号码 —— 与「联系电话」重复
_IGNORE_COLUMNS = {"开户行", "发放公司", "电话号码"}

# 尾部压缩时必须保留的列(即使整列为空)
_KEEP_COLUMNS = {"文件名", "项目", "姓名", "身份证", "手续费"}

_ALIAS = {
    "个税扣款": "个税",
    "入职日期": "入职时间",
    "离职日期": "离职时间",
    "发薪公司": "发薪单位",
    "手机号": "联系电话",
    "预支": "预支扣款",
    "预支扣回": "预支扣款",
}


def alias_name(h):
    """源表头 -> 标准列名: 别名归并 (与 VBA AliasName 等价)

    注意: 「发放公司」「电话号码」【不归并】—— 它们要命中忽略名单整列跳过；
         若归并成「发薪单位」「联系电话」, 跳过就失效了。
    """
    if h is None:
        return ""
    h = str(h)
    if h in _ALIAS:
        return _ALIAS[h]
    if h.startswith("个税申报所属期"):
        return "个税申报所属期"
    return h


def local_name(hdr_val):
    """源表某列的表头 -> 本地列名; 返回 "" 表示无表头(该列跳过)"""
    if hdr_val is None:
        return ""
    h = remove_spaces(str(hdr_val).strip())
    if h == "":
        return ""
    h = alias_name(h)
    # 与生成列同名的源列 -> 加「(源)」后缀另立一列
    if h in ("文件名", "项目", "手续费", "采用/不采用"):
        h = h + "(源)"
    return h


def _index_of(names):
    """列名表 -> {列名: 1-based 列号}"""
    return {n: i + 1 for i, n in enumerate(names)}


def _column_has_data(rows, col0):
    """该列(0-based)在数据行中是否有非空值"""
    for r in rows:
        if col0 < len(r):
            v = r[col0]
            if v is not None and remove_spaces(str(v)) != "":
                return True
    return False


def _id_text(v):
    """身份证文本化: 空保持空; 数值转全位数字符串(防 18 位变科学计数法)"""
    if v is None:
        return None
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else str(v)
    s = str(v)
    return s if s != "" else None


def _resolve_insert_pos(loc, s, src_cols, u_dict, u_names):
    """新列插入位置 (与 VBA SyncUnion 等价): 左右邻相邻→插中间; 否则以左邻为准"""
    p_name = ""
    n_name = ""
    for j in range(s - 1, 1, -1):
        if loc[j] and loc[j] in u_dict:
            p_name = loc[j]
            break
    for j in range(s + 1, src_cols + 1):
        if loc[j] and loc[j] in u_dict:
            n_name = loc[j]
            break
    if p_name and n_name:
        p_idx, n_idx = u_dict[p_name], u_dict[n_name]
        if n_idx == p_idx + 1:
            return p_idx + 1
        return p_idx + 1 if n_idx > p_idx else n_idx
    if p_name:
        return u_dict[p_name] + 1
    if n_name:
        return u_dict[n_name]
    return len(u_names) if u_names[-1] == "采用/不采用" else len(u_names) + 1


def _compact_tail(u_names, rows):
    """列尾压缩 (与 VBA CompactTailColumns 等价)

    删除「最后一个有效列」右侧的空列, 使「采用/不采用」紧跟数据。
    只裁尾部 —— 中间空列保留, 保证跨月列序稳定。
    返回删除的列数。
    """
    if "采用/不采用" not in u_names:
        return 0
    tag_idx = u_names.index("采用/不采用") + 1          # 1-based
    if tag_idx <= 1 or tag_idx != len(u_names):
        return 0
    last_data = 0
    for c in range(tag_idx - 1, 0, -1):
        nm = u_names[c - 1]
        if nm in _KEEP_COLUMNS:
            last_data = c
            break
        if _column_has_data(rows, c - 1):
            last_data = c
            break
    if last_data == 0 or last_data >= tag_idx - 1:
        return 0
    removed = tag_idx - 1 - last_data
    del u_names[last_data:tag_idx - 1]
    for r in rows:
        del r[last_data:tag_idx - 1]
    return removed


# ============================================================
# 手续费规则分类与计算
# ============================================================

# A~N 各类精确匹配串 (必须与 VBA 原值逐字一致, 含全角标点)
_RULE_SETS = {
    "A": {
        "只有代工费", "只有管理费，没有工资", "只有代工费，没有工资", "不扣手续费",
        "日结", "联动代发不收手续费", "不发工资，没有手续费", "不扣手续费，锦途发放",
        "不收手续费", "都不扣",
    },
    "B": {"除招商都扣", "招商", "招商不收", "除招商外都扣"},
    "C": {
        "除招商、建设外都扣，离职员工不扣，前15个员工不扣", "除招商和建设外都扣",
        "招商、建设不收", "除招商、建设外都扣", "招商、建设不扣，其他银行都扣",
    },
    "D": {"除平安外都扣"},
    "E": {"除宁波外都扣"},
    "F": {"除工商都扣"},
    "G": {"除温州银行外都扣手续费", "除温州银行外都收手续费"},
    "H": {"邮政卡不收手续费（除邮政都收）", "除了邮政之外的手续费员工承担"},
    "I": {"除邮政银行，其他银行都扣"},
    "J": {"招商和邮政不扣手续费", "除招商、邮政外都扣", "招商、邮储不扣，其他银行都扣"},
    "K": {"除邮政，中信外都扣"},
    "L": {"中信不扣，其他都扣"},
    "M": {"签浙江锦途：招商、建设不扣；签温州驰锦：民生不扣"},
    # N-全扣: 所有银行都扣手续费 (2026-10-26 补齐, 此前缺失导致该规则的项目手续费留空)
    "N": {"所有银行都扣手续费"},
}


def classify_rule_exact(rule_text):
    """将规则文本精确匹配到类别 A~N; 未命中返回空串 (与 VBA ClassifyRuleExact 等价)"""
    s = remove_spaces(rule_text)
    for cat, sset in _RULE_SETS.items():
        if s in sset:
            return cat
    return ""


def is_bank_exempt(category, bank_name):
    """按类别 + 开户行(包含匹配)判断是否免扣 (与 VBA IsBankExempt 等价)"""
    b = remove_spaces(bank_name)
    if category == "A":
        return True
    if category == "B":
        return "招商银行" in b
    if category == "C":
        return ("招商银行" in b) or ("建设银行" in b)
    if category == "D":
        return "平安银行" in b
    if category == "E":
        return "宁波银行" in b
    if category == "F":
        return "工商" in b
    if category == "G":
        return "温州银行" in b
    if category in ("H", "I"):
        return ("邮政" in b) or ("邮储" in b)
    if category == "J":
        return ("招商银行" in b) or ("邮政" in b) or ("邮储" in b)
    if category == "K":
        return ("邮政" in b) or ("邮储" in b) or ("中信银行" in b)
    if category == "L":
        return "中信银行" in b
    # N-全扣: 所有银行都不免除 (对齐 VBA IsBankExempt 的 Case "N")
    if category == "N":
        return False
    return False


def resolve_amount(table_amount, category):
    """解析最终金额: 表格有效正数优先, 否则按类别默认 (A=0, I=3, 其余=5)"""
    if table_amount is not None and table_amount != "":
        try:
            num = float(table_amount)
            if num > 0:
                return int(num)
        except (ValueError, TypeError):
            pass
    if category == "A":
        return 0
    if category == "I":
        return 3
    return 5


def handle_dual_unit_rule(pay_unit, bank, table_amount):
    """M 类双单位特殊规则 (与 VBA HandleDualUnitRule 等价)"""
    pu = remove_spaces(pay_unit)
    b = remove_spaces(bank)
    amt = 5
    if table_amount is not None and table_amount != "":
        try:
            tn = float(table_amount)
            if tn > 0:
                amt = int(tn)
        except (ValueError, TypeError):
            pass
    if pu == "浙江锦途人力资源有限公司":
        return 0 if ("招商银行" in b or "建设银行" in b) else amt
    if pu == "温州驰锦企业服务有限公司":
        return 0 if "民生" in b else amt
    # 兜底: 招商/建行不扣
    return 0 if ("招商银行" in b or "建设银行" in b) else amt


def _is_invalid_name_id(v):
    """姓名/身份证无效判定: 空 / "0" / 错误值"""
    if v is None:
        return True
    s = str(v).strip()
    if s == "" or s == "0":
        return True
    if s.startswith("#"):
        return True
    return False


def compute_fee_v2(name_val, id_val, pay_unit_val, info_lookup, fee_rules, missing_ops=None):
    """逐行计算手续费 (与 VBA ComputeFee_V2 等价)

    info_lookup: dict{ normalized_name: [(ops, bank), ...] }
    fee_rules:   dict{ 运营项目: [扣费银行规则文本, 扣费金额] }
    missing_ops: 可选 set, 用于收集"源文件有但规则表缺"的运营项目(对应 VBA gMissingOps)
    返回: 0 / 3 / 5 / 其他金额, 或 None(留空)
    """
    if _is_invalid_name_id(name_val) or _is_invalid_name_id(id_val):
        return None
    if not fee_rules:
        return None
    if not info_lookup:
        return None
    key = normalize_name(name_val)
    entries = info_lookup.get(key)
    if not entries:
        return None
    for ops, bank in entries:
        if ops is None:
            continue
        base_ops = str(ops).strip()
        if base_ops == "":
            continue
        # "代发Xxx" 去掉前缀后按基础运营项目查规则
        if base_ops.startswith("代发"):
            base_ops = base_ops[2:]
        if base_ops not in fee_rules:
            # 对应 VBA gMissingOps: 源文件有此运营项目但规则表无该行
            if missing_ops is not None:
                missing_ops.add(base_ops)
            continue
        rule_text, table_amount = fee_rules[base_ops]
        category = classify_rule_exact(rule_text)
        if category == "":
            # 与最新 VBA 一致: 规则文本无法识别时直接留空, 不静默扣5元, 也不再收集
            continue
        if category == "M":
            return handle_dual_unit_rule(pay_unit_val, bank, table_amount)
        if is_bank_exempt(category, bank):
            return 0
        return resolve_amount(table_amount, category)
    return None


# ============================================================
# 规则表加载
# ============================================================

def load_fee_rules(path, on_log=None):
    """从规则 Excel 加载规则字典 (与 VBA LoadFeeRules 等价)

    自动选择名字含"手续费"的工作表, 否则取第一个;
    动态定位 客户状态/运营项目/扣费银行（多填）/扣费金额 列;
    跳过"客户状态=终止合作"; 重复运营项目以最后一条为准。
    返回: dict{ 运营项目: [扣费银行规则文本, 扣费金额] }
    """
    if on_log is None:
        on_log = lambda m: None
    wb = load_workbook(path, data_only=True, read_only=True)
    sheet_name = None
    for sn in wb.sheetnames:
        if "手续费" in sn:
            sheet_name = sn
            break
    if sheet_name is None:
        sheet_name = wb.sheetnames[0]
    on_log(f"规则表使用工作表: {sheet_name}")
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return {}
    header = rows[0]
    col_status = col_ops = col_bank = col_amount = 0
    for i, h in enumerate(header, start=1):
        ht = remove_spaces(h) if h is not None else ""
        if ht == "客户状态":
            col_status = i
        elif ht == "运营项目":
            col_ops = i
        elif ht in ("扣费银行（多填）", "扣费银行(多填)"):
            col_bank = i
        elif ht == "扣费金额":
            col_amount = i
    if col_ops == 0 or col_bank == 0:
        on_log("⚠ 规则表缺少「运营项目」或「扣费银行（多填）」列, 无法加载")
        return {}
    rules = {}
    for r in rows[1:]:
        if col_status > 0:
            st = r[col_status - 1] if len(r) >= col_status else None
            if st is not None and remove_spaces(str(st)) == "终止合作":
                continue
        ops_val = r[col_ops - 1] if len(r) >= col_ops else None
        if ops_val is None or str(ops_val).strip() == "":
            continue
        ops_key = str(ops_val).strip()
        bank_val = r[col_bank - 1] if len(r) >= col_bank else None
        amount_val = r[col_amount - 1] if (col_amount > 0 and len(r) >= col_amount) else None
        rules[ops_key] = [bank_val, amount_val]
    on_log(f"已加载手续费规则 {len(rules)} 条")
    return rules


# ============================================================
# 文件扫描
# ============================================================

def scan_excel_files(root, recursive=True, exclude_paths=None):
    """扫描 Excel 文件列表

    recursive=True 递归子目录 (默认); 否则仅根目录一层。
    排除: ~$ 临时文件, 非 .xlsx/.xls, exclude_paths 中的路径。
    """
    exclude = set()
    if exclude_paths:
        exclude.update(os.path.normcase(os.path.abspath(p)) for p in exclude_paths)

    result = []

    def _should_skip(full):
        fn = os.path.basename(full)
        if fn.startswith("~$"):
            return True
        if not fn.lower().endswith((".xlsx", ".xls")):
            return True
        if os.path.normcase(os.path.abspath(full)) in exclude:
            return True
        return False

    if recursive:
        for dirpath, _dirnames, filenames in os.walk(root):
            for fn in filenames:
                full = os.path.join(dirpath, fn)
                if not _should_skip(full):
                    result.append(full)
    else:
        for fn in os.listdir(root):
            full = os.path.join(root, fn)
            if not os.path.isfile(full):
                continue
            if not _should_skip(full):
                result.append(full)
    return result


# ============================================================
# 读取单个个税工作簿
# ============================================================

def _read_payroll_file(fpath):
    """读取单个工作簿的个税统计表 + 员工信息表 (纯 openpyxl, data_only)

    返回: (src_headers, src_rows, src_cols, col_name, col_id, col_gjj, col_pay,
           info_lookup, info_found, warns)
    若无"个税统计"表返回 None; 若缺关键列抛 ValueError。
    """
    wb = None
    try:
        wb = load_workbook(fpath, data_only=True, read_only=True)

        # ---- 找"个税统计"表 ----
        target = None
        for sn in wb.sheetnames:
            if remove_spaces(sn) == "个税统计":
                target = sn
                break
        if target is None:
            wb.close()
            return None
        ws = wb[target]
        rows = list(ws.iter_rows(values_only=True))

        # ---- 找"员工信息-X月"表 ----
        info_sheet = None
        month = extract_month_from_filename(os.path.basename(fpath))
        for sn in wb.sheetnames:
            if month and remove_spaces(sn) == ("员工信息-" + month + "月"):
                info_sheet = sn
                break
        if info_sheet is None:
            for sn in wb.sheetnames:
                if remove_spaces(sn).startswith("员工信息-"):
                    info_sheet = sn
                    break

        info_lookup = {}
        info_found = info_sheet is not None
        warns = []
        if info_sheet is not None:
            iws = wb[info_sheet]
            irows = list(iws.iter_rows(values_only=True))
            if irows:
                iheader = irows[0]
                ci = co = cb = 0
                for i, h in enumerate(iheader, start=1):
                    ht = remove_spaces(h) if h is not None else ""
                    if ht == "姓名":
                        ci = i
                    elif ht == "运营项目":
                        co = i
                    elif ht == "开户行":
                        cb = i
                if ci > 0 and co > 0 and cb > 0:
                    for r in irows[1:]:
                        nm = r[ci - 1] if len(r) >= ci else None
                        if nm is None or str(nm).strip() == "":
                            continue
                        ops = r[co - 1] if len(r) >= co else None
                        bank = r[cb - 1] if len(r) >= cb else None
                        info_lookup.setdefault(normalize_name(nm), []).append((ops, bank))
                else:
                    warns.append("⚠ 员工信息表缺少 姓名/运营项目/开户行 列")

        if not rows:
            wb.close()
            return None
        header = list(rows[0])

        # ---- 实际列数(最后有数据的列) ----
        src_cols = 0
        for r in rows:
            for j in range(len(r) - 1, -1, -1):
                if r[j] is not None and str(r[j]).strip() != "":
                    if j + 1 > src_cols:
                        src_cols = j + 1
                    break
        if src_cols == 0:
            src_cols = len(header)

        # ---- 定位关键列 ----
        col_name = col_id = col_gjj = col_pay = 0
        for i, h in enumerate(header, start=1):
            if i > src_cols:
                break
            ht = remove_spaces(h) if h is not None else ""
            if ht == "":
                continue
            if ht == "姓名":
                col_name = i
            elif ht == "身份证":
                col_id = i
            elif ht == "公积金":
                col_gjj = i
            # 发薪单位: 先别名归并(发薪公司→发薪单位), 再兼容「发放公司」写法。
            # 「发放公司」不进汇总表(忽略名单), 但 M 类双单位规则仍需要它判定发薪单位。
            elif col_pay == 0 and (alias_name(ht) == "发薪单位" or ht == "发放公司"):
                col_pay = i

        # 2026-10-26: 公积金列不再强制要求 (对齐 VBA —— 缺列留空, 不整簿跳过)
        if col_name == 0 or col_id == 0:
            wb.close()
            raise ValueError("表头缺少「姓名」或「身份证」列")

        src_headers = [(header[i - 1] if i - 1 < len(header) else None)
                       for i in range(1, src_cols + 1)]

        src_rows = []
        last_non_blank = -1
        for r in rows[1:]:
            rr = list(r)
            if len(rr) < src_cols:
                rr = rr + [None] * (src_cols - len(rr))
            else:
                rr = rr[:src_cols]
            src_rows.append(rr)
            if any(v is not None and remove_spaces(str(v)) != "" for v in rr):
                last_non_blank = len(src_rows) - 1
        # 裁掉尾部「仅格式无值」的行 (对应 VBA src_rows = SheetRow - 1),
        # 否则这些行会以空白数据进入汇总并被标成「不采用」
        if last_non_blank + 1 < len(src_rows):
            src_rows = src_rows[:last_non_blank + 1]

        wb.close()
        return (src_headers, src_rows, src_cols, col_name, col_id, col_gjj, col_pay,
                info_lookup, info_found, warns)
    except Exception:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass
        raise


def write_merge_output(output_path, header, all_rows, target_id_col):
    """写出汇总表, 身份证列文本化"""
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for out in all_rows:
        row_copy = list(out)
        if target_id_col - 1 < len(row_copy):
            v = row_copy[target_id_col - 1]
            row_copy[target_id_col - 1] = "" if v is None else str(v)
        ws.append(row_copy)
    if 0 < target_id_col <= len(header):
        col_letter = get_column_letter(target_id_col)
        for cell in ws[col_letter]:
            cell.number_format = "@"
    wb.save(output_path)


# ============================================================
# 主流程
# ============================================================

def run_merge_collect(scan_root, rules_path, cancel_event=None,
                      on_log=None, on_progress=None, on_done=None):
    """合并个税 - 阶段一: 扫描 + 读取 + 内存汇总(不写文件)

    参数:
      scan_root    扫描根目录(始终递归子目录)
      rules_path   手续费规则表路径
      cancel_event threading.Event, 置位则中止
      on_log/on_progress/on_done 回调
    结果经 on_done 返回 dict:
      header        输出表头(list，19 列白名单 + 白名单外新列，尾部空列已压缩)
      rows          汇总数据(list[list])
      target_id_col 身份证列在 header 中的 1-based 位置
      total_files / total_rows / skipped / rule_count
      errors / empty_keyword_ops / missing_ops / skipped_columns
      output_path   固定为 None (本阶段不写文件, 由 GUI 弹出保存对话框后再调 write_merge_output)
    """
    if on_log is None:
        on_log = lambda m: None
    if on_progress is None:
        on_progress = lambda c, t, m: None
    if on_done is None:
        on_done = lambda r: None
    if cancel_event is None:
        cancel_event = threading.Event()

    result = {"header": None, "rows": [], "target_id_col": None,
              "total_files": 0, "total_rows": 0, "skipped": 0, "rule_count": 0,
              "errors": [], "empty_keyword_ops": set(), "missing_ops": set(),
              "skipped_columns": set(), "output_path": None}

    # ---- 加载规则 ----
    try:
        fee_rules = load_fee_rules(rules_path, on_log)
    except Exception as e:
        on_log(f"⚠ 读取规则表失败: {e}")
        on_done(result)
        return
    result["rule_count"] = len(fee_rules)
    if not fee_rules:
        on_log("⚠ 未加载到任何手续费规则, 手续费列将全部留空")

    # ---- 校验：规则表中"扣费银行（多填）"列是否为空 (对应 VBA gEmptyKeywordOps) ----
    empty_keyword_ops = result["empty_keyword_ops"]
    for ops_key, entry in fee_rules.items():
        rule_text = entry[0]
        if rule_text is None or remove_spaces(str(rule_text)) == "":
            empty_keyword_ops.add(ops_key)

    # ---- 扫描文件(始终递归) ----
    file_list = scan_excel_files(scan_root, recursive=True,
                                  exclude_paths=[rules_path])
    if not file_list:
        on_log("未找到任何 Excel 文件")
        on_done(result)
        return

    result["total_files"] = len(file_list)
    missing_ops = result["missing_ops"]

    # ============================================================
    # 输出列布局：固定 19 列白名单（对齐 VBA「按列名映射重写版」/ Java HbgsService）
    #   · 源簿缺列 → 该列留空，不影响其他文件
    #   · 白名单外的列 → 按「左邻/右邻」插入输出表并写表头
    #   · 忽略名单（开户行/发放公司/电话号码）→ 整列跳过
    #   · 末尾做「尾部压缩」，「采用/不采用」跟随最后一个有效列
    # ============================================================
    u_names = list(_STD_COLUMNS)
    u_dict = _index_of(u_names)
    all_rows = []
    skipped_columns = result["skipped_columns"]
    u_count = len(u_names)

    for idx, fpath in enumerate(file_list, start=1):
        if cancel_event.is_set():
            on_log("⚠ 用户取消, 停止处理")
            break
        on_progress(idx, len(file_list), f"读取中: {os.path.basename(fpath)}")
        try:
            data = _read_payroll_file(fpath)
        except Exception as e:
            msg = f"✗ 无法读取 {os.path.basename(fpath)}: {e}"
            on_log(msg)
            result["errors"].append(msg)
            result["skipped"] += 1
            continue
        if data is None:
            msg = f"✗ {os.path.basename(fpath)} 无「个税统计」表, 已跳过"
            on_log(msg)
            result["errors"].append(msg)
            result["skipped"] += 1
            continue

        (src_headers, src_rows, src_cols, col_name, col_id, _col_gjj, col_pay,
         info_lookup, info_found, warns) = data
        for w in warns:
            on_log(w)
        if not info_found:
            on_log(f"⚠ {os.path.basename(fpath)} 无员工信息表, 手续费将留空")

        # ---- ① 表头 → 本地列名：别名归并 + 生成列加「(源)」+ 同簿重名 #2 ----
        loc = [""] * (src_cols + 1)
        used_names = set()
        for s in range(2, src_cols + 1):
            nm = local_name(src_headers[s - 1] if s - 1 < len(src_headers) else None)
            if nm:
                if nm in used_names:
                    c = 2
                    while f"{nm}#{c}" in used_names:
                        c += 1
                    nm = f"{nm}#{c}"
                used_names.add(nm)
            loc[s] = nm

        # ---- ② 忽略名单 / 无表头告警 / 白名单外插列 ----
        for s in range(2, src_cols + 1):
            nm = loc[s]
            if not nm:
                # 表头空但列下有值 → 该列数据会丢，明确告警
                if _column_has_data(src_rows, s - 1):
                    on_log(f"⚠ {os.path.basename(fpath)}：第 {s} 列无表头，该列数据未汇总")
                continue
            if nm in _IGNORE_COLUMNS:
                skipped_columns.add(nm)
                continue
            if nm in u_dict:
                continue
            pos = _resolve_insert_pos(loc, s, src_cols, u_dict, u_names)
            u_names.insert(pos - 1, nm)
            u_dict = _index_of(u_names)
            # 等价 VBA 的 gS2.Columns(pos).Insert：已汇总的行整体右移一格
            for r in all_rows:
                r.insert(pos - 1, None)
            on_log(f"⚠ {os.path.basename(fpath)}：新增列「{nm}」插入第 {pos} 列")

        # ---- ③ 源列号 → 输出列号 ----
        u_count = len(u_names)
        g_idx = [0] * (src_cols + 1)
        for s in range(2, src_cols + 1):
            if loc[s]:
                g_idx[s] = u_dict.get(loc[s], 0)
        f_idx = u_dict["文件名"]
        p_idx = u_dict["项目"]
        id_idx = u_dict["身份证"]
        fee_idx = u_dict["手续费"]
        fname = os.path.splitext(os.path.basename(fpath))[0]

        # ---- ④ 逐行落位 ----
        for row in src_rows:
            out = [None] * u_count
            out[f_idx - 1] = fname
            out[p_idx - 1] = clean_project_name(row[0] if len(row) > 0 else None)
            for s in range(2, min(src_cols, len(row)) + 1):
                if g_idx[s]:
                    out[g_idx[s] - 1] = row[s - 1]
            # 身份证文本化（防 18 位变科学计数法）
            out[id_idx - 1] = _id_text(out[id_idx - 1])
            # 手续费：引擎按「源表列号」取 姓名 / 身份证 / 发薪单位
            name_v = row[col_name - 1] if (0 < col_name <= len(row)) else None
            id_v = row[col_id - 1] if (0 < col_id <= len(row)) else None
            pay_v = row[col_pay - 1] if (0 < col_pay <= len(row)) else None
            out[fee_idx - 1] = compute_fee_v2(name_v, id_v, pay_v, info_lookup,
                                              fee_rules, missing_ops)
            all_rows.append(out)

    if not all_rows:
        on_log("没有有效数据可汇总！")
        on_done(result)
        return

    # ---- ⑤ 尾部压缩：「采用/不采用」跟随最后一个有效列 ----
    removed = _compact_tail(u_names, all_rows)
    u_dict = _index_of(u_names)
    u_count = len(u_names)

    # ---- ⑥ 写「采用/不采用」----
    tag_idx = u_dict["采用/不采用"]
    name_out_idx = u_dict["姓名"]
    id_out_idx = u_dict["身份证"]
    for r in all_rows:
        nv = r[name_out_idx - 1]
        iv = r[id_out_idx - 1]
        r[tag_idx - 1] = "不采用" if (_is_invalid_name_id(nv) or _is_invalid_name_id(iv)) else "采用"

    result["header"] = u_names
    result["rows"] = all_rows
    result["target_id_col"] = u_dict["身份证"]
    result["total_rows"] = len(all_rows)
    on_log(f"✔ 内存汇总完成: 共 {len(all_rows)} 条数据, 输出 {u_count} 列, 请在弹出的对话框中选择保存位置")
    if removed:
        on_log(f"列尾压缩：删除尾部 {removed} 个空列，「采用/不采用」位于第 {tag_idx} 列")

    # ---- 提示汇总①：规则表"扣费银行（多填）"列为空 (对应 VBA gEmptyKeywordOps) ----
    if empty_keyword_ops:
        e_list = " ".join("[" + str(o) + "]" for o in sorted(empty_keyword_ops))
        on_log('【规则表提示】以下运营项目的"扣费银行（多填）"列为空，需要填写关键词：' + e_list)

    # ---- 提示汇总②：源文件有但规则表缺的运营项目 (对应 VBA gMissingOps) ----
    if missing_ops:
        m_list = " ".join("[" + str(o) + "]" for o in sorted(missing_ops))
        on_log('【项目提示】以下运营项目在规则表中找不到，请在"收取银行手续费信息"工作表中填写运营项目及文本关键词：' + m_list)

    # ---- 提示汇总③：被忽略名单整列跳过的源列 (对应 VBA gSkippedCols) ----
    if skipped_columns:
        s_list = " ".join("[" + str(o) + "]" for o in sorted(skipped_columns))
        on_log("【已跳过列】以下源表列按忽略名单整列跳过，未汇总到结果表：" + s_list)

    on_done(result)
