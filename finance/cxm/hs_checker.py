# -*- coding: utf-8 -*-
"""
数据核对模块
包含：DataChecker 类、save_combined_excel 函数
"""

import re
import pandas as pd
from openpyxl import load_workbook
from typing import Dict, List, Tuple, Optional

from finance.cxm.hs_utils import match_month_in_text


class DataChecker:
    """数据核对器：对比待核对表与基准表的金额"""

    def __init__(self, reference_path: str):
        self.reference_path = reference_path
        self.reference_data: Dict[str, float] = {}  # 供应商名称 -> 总计
        self.reference_df: Optional[pd.DataFrame] = None  # 保存原始DataFrame

    def load_reference(self) -> Tuple[bool, str]:
        """加载基准表数据（供应商名称 -> 总计），使用向量化操作替代iterrows"""
        try:
            df = pd.read_excel(self.reference_path)
            if '供应商名称' not in df.columns:
                return False, f"基准表未找到'供应商名称'列"
            if '总计' not in df.columns:
                return False, f"基准表未找到'总计'列"

            self.reference_df = df  # 保存原始DataFrame，用于后续提取供应商名称

            # 向量化处理：比iterrows快10-100倍
            valid_df = df.dropna(subset=['供应商名称']).copy()
            suppliers = valid_df['供应商名称'].astype(str).str.strip()
            totals = pd.to_numeric(valid_df['总计'], errors='coerce').fillna(0.0)
            self.reference_data = dict(zip(suppliers, totals))

            return True, f"加载基准表成功，共{len(self.reference_data)}条数据"
        except Exception as e:
            return False, f"加载基准表失败：{str(e)}"

    def load_reference_from_dataframe(self, df: pd.DataFrame) -> Tuple[bool, str]:
        """从内存中的DataFrame加载基准表数据（供应商名称 -> 总计）"""
        try:
            if '供应商名称' not in df.columns:
                return False, f"基准表未找到'供应商名称'列"
            if '总计' not in df.columns:
                return False, f"基准表未找到'总计'列"

            self.reference_df = df

            valid_df = df.dropna(subset=['供应商名称']).copy()
            suppliers = valid_df['供应商名称'].astype(str).str.strip()
            totals = pd.to_numeric(valid_df['总计'], errors='coerce').fillna(0.0)
            self.reference_data = dict(zip(suppliers, totals))

            return True, f"从内存加载基准表成功，共{len(self.reference_data)}条数据"
        except Exception as e:
            return False, f"从内存加载基准表失败：{str(e)}"

    def detect_columns(self, df: pd.DataFrame) -> Tuple[Optional[str], Optional[str], Optional[str]]:
        """识别月份、介绍人、金额三列，返回(月份列, 介绍人列, 金额列)"""
        month_col = None
        person_col = None
        amount_col = None

        for col in df.columns:
            col_str = str(col).strip()
            if month_col is None and '月份' in col_str:
                month_col = col
            if person_col is None and '介绍人' in col_str:
                person_col = col
            if amount_col is None and '金额' in col_str:
                amount_col = col

        return month_col, person_col, amount_col

    def get_unique_months(self, df: pd.DataFrame, month_col: str) -> List[str]:
        """获取月份列的唯一值（支持X月和X-Y月格式）"""
        if month_col not in df.columns:
            return []
        months_found = set()
        for val in df[month_col].dropna():
            val_str = str(val).strip()
            # 单月份
            found = re.findall(r'(?<!\d)(\d{1,2}月)(?!\d)', val_str)
            for m in found:
                try:
                    num = int(m.replace('月', ''))
                    if 1 <= num <= 12:
                        months_found.add(m)
                except ValueError:
                    pass
            # 范围月份
            range_found = re.findall(r'(?<!\d)(\d{1,2}-\d{1,2}月)(?!\d)', val_str)
            for m in range_found:
                try:
                    start, end = m.replace('月', '').split('-')
                    if 1 <= int(start) <= 12 and 1 <= int(end) <= 12:
                        months_found.add(m)
                except ValueError:
                    pass

        def sort_key(x):
            if '-' in x:
                return (1, int(x.split('-')[0]))
            return (0, int(x.replace('月', '')))

        return sorted(list(months_found), key=sort_key)

    def get_unique_persons(self, df: pd.DataFrame, person_col: str) -> List[str]:
        """获取介绍人列的非空唯一值"""
        if person_col not in df.columns:
            return []
        values = df[person_col].dropna().astype(str).str.strip()
        values = values[values != '']
        valid_values = []
        for v in values.unique():
            if len(v) > 0 and not v.startswith('#') and v.lower() not in ['nan', 'none', 'null']:
                valid_values.append(v)
        return sorted(valid_values)

    def check(self, df: pd.DataFrame, selected_month: str, selected_persons: List[str],
              month_col: str, person_col: str, amount_col: str) -> Tuple[bool, str, Optional[pd.DataFrame]]:
        """执行核对：按介绍人分组累加金额，与基准表总计对比

        供应商名称匹配规则：去除末尾数字后作为基础名称进行匹配，
        例如"供应商-张三"、"供应商-张三1"、"供应商-张三2"均视为"供应商-张三"
        """
        try:

            def get_base_name(name: str) -> str:
                """去除名称末尾的数字，获取基础名称（如"供应商-张三1" → "供应商-张三"）"""
                return re.sub(r'\d+$', '', str(name).strip())

            # 筛选月份（使用公共月份匹配函数）
            df_filtered = df[df[month_col].apply(lambda val: match_month_in_text(selected_month, val))]

            if df_filtered.empty:
                return False, f"月份'{selected_month}'没有匹配到数据", None

            results = []
            for person in selected_persons:
                person_base = get_base_name(person)

                # 从基准表中查找所有基础名称匹配的供应商，累加金额
                matched_suppliers = []
                base_total = 0.0
                for supplier, total in self.reference_data.items():
                    if get_base_name(supplier) == person_base:
                        matched_suppliers.append(supplier)
                        base_total += total

                # 计算待核对金额
                person_df = df_filtered[df_filtered[person_col].astype(str).str.strip() == person]
                check_amount = person_df[amount_col].sum()

                if matched_suppliers:
                    if abs(float(check_amount) - base_total) < 0.01:
                        status = "已核对"
                    else:
                        status = "待核对"
                    # 供应商名称显示为基础名称（去除数字后缀）
                    supplier_name = person_base
                else:
                    status = "未匹配"
                    supplier_name = ''
                    base_total = None

                results.append({
                    '介绍人': person,
                    '待核对金额总计': round(float(check_amount), 2),
                    '供应商名称': supplier_name,
                    '对公打款总计': round(base_total, 2) if base_total is not None else '',
                    '状态': status
                })

            result_df = pd.DataFrame(results)
            return True, f"核对完成，共{len(results)}条记录", result_df

        except Exception as e:
            import traceback
            return False, f"核对失败：{str(e)}\n{traceback.format_exc()}", None

    @staticmethod
    def load_check_file(file_path: str, sheet_name: str) -> Tuple[Optional[pd.DataFrame], Optional[str], Optional[str], Optional[str], str]:
        """使用openpyxl快速加载待核对数据，自动检测表头行，只读取需要的三列"""
        try:
            wb = load_workbook(file_path, read_only=True, data_only=True)
            ws = wb[sheet_name]

            # 自动检测表头行（前5行内查找包含"月份"、"介绍人"、"金额"的行）
            header_row = None
            headers = []
            for row_idx in range(1, min(6, ws.max_row + 1)):
                row_headers = [str(cell.value).strip() if cell.value else '' for cell in ws[row_idx]]
                has_month = any('月份' in h for h in row_headers)
                has_person = any('介绍人' in h for h in row_headers)
                has_amount = any('金额' in h for h in row_headers)
                if has_month and has_person and has_amount:
                    header_row = row_idx
                    headers = row_headers
                    break

            if header_row is None:
                wb.close()
                return None, None, None, None, "未在表头中找到'月份'、'介绍人'、'金额'列，请确认表头在前5行内"

            # 找到列索引
            month_col = person_col = amount_col = None
            month_idx = person_idx = amount_idx = None
            for i, h in enumerate(headers):
                if month_col is None and '月份' in h:
                    month_col = h
                    month_idx = i
                if person_col is None and '介绍人' in h:
                    person_col = h
                    person_idx = i
                if amount_col is None and '金额' in h:
                    amount_col = h
                    amount_idx = i

            if month_idx is None or person_idx is None or amount_idx is None:
                wb.close()
                return None, None, None, None, f"识别列失败，检测到的表头：{headers}"

            # 只读取三列数据
            data = []
            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                data.append({
                    month_col: row[month_idx] if month_idx < len(row) else None,
                    person_col: row[person_idx] if person_idx < len(row) else None,
                    amount_col: row[amount_idx] if amount_idx < len(row) else None,
                })

            df = pd.DataFrame(data)
            wb.close()
            return df, month_col, person_col, amount_col, f"共{len(df)}行，表头在第{header_row}行"

        except Exception as e:
            import traceback
            return None, None, None, None, f"加载失败：{str(e)}\n{traceback.format_exc()}"





