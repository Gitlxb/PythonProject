# -*- coding: utf-8 -*-
"""
Tab 2「核对」UI Mixin
包含：发放记录表区、供应商对账单区、核对按钮栏、汇总动作、自动保存、对比输出

通过多重继承混入 SummaryProcessorApp：
  class SummaryProcessorApp(GuiActionsMixin, SplitTabMixin, CheckTabMixin):
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import pandas as pd
import re
import os
from openpyxl import load_workbook
from finance.cxm.hs_check_processor import CheckProcessor
from finance.cxm.hs_utils import apply_excel_styles
from openpyxl.styles import PatternFill


class CheckTabMixin:
    """Tab 2「核对」的 UI 创建和事件处理"""

    # ==================== Tab 2 内容组装 ====================

    def _create_tab2_content(self, parent):
        """构建 Tab 2（核对）的全部 UI

        布局 grid 行分配：
          row=0  A区 发放记录表
          row=1  B区 供应商对账单
          row=2  操作按钮栏
          row=3  状态信息（共用 ScrolledText）
        """
        parent.columnconfigure(0, weight=1)
        parent.rowconfigure(3, weight=1)

        self._create_check_disbursement_frame(parent)   # row=0
        self._create_check_statement_frame(parent)     # row=1
        self._create_check_button_frame(parent)        # row=2
        self._create_status_frame(parent, row=3)       # row=3

    # ────────────────────────────────────────────────
    #  A. 发放记录表区域
    # ────────────────────────────────────────────────

    def _create_check_disbursement_frame(self, parent):
        f = ttk.LabelFrame(parent, text="A. 发放记录表（汇总条件）", padding="12")
        f.grid(row=0, column=0, sticky=(tk.W, tk.E), pady=(0, 10))
        f.columnconfigure(1, weight=1)
        f.columnconfigure(4, weight=1)

        # 文件选择行
        ttk.Label(f, text="发放记录表：").grid(
            row=0, column=0, padx=(0, 5), sticky=tk.W)
        self.check_disbursement_path_var = tk.StringVar()
        self.check_disbursement_entry = ttk.Entry(
            f, textvariable=self.check_disbursement_path_var)
        self.check_disbursement_entry.grid(row=0, column=1, padx=(0, 10),
                                            sticky=(tk.W, tk.E))
        ttk.Button(f, text="浏览", command=self._check_browse_disbursement,
                   width=8).grid(row=0, column=2, padx=(0, 10))
        ttk.Label(f, text="工作簿（可多选）：").grid(
            row=0, column=3, padx=(0, 5))

        lb_frame = ttk.Frame(f)
        lb_frame.grid(row=0, column=4, sticky=(tk.W, tk.E, tk.N, tk.S))
        lb_frame.columnconfigure(0, weight=1)
        lb_frame.rowconfigure(0, weight=1)
        self.check_disbursement_listbox = tk.Listbox(
            lb_frame, selectmode=tk.EXTENDED, height=4,
            font=('Microsoft YaHei UI', 9), exportselection=False)
        self.check_disbursement_listbox.grid(row=0, column=0,
                                              sticky=(tk.W, tk.E, tk.N, tk.S))
        self.check_disbursement_listbox.bind(
            '<<ListboxSelect>>', self._check_on_disbursement_sheet_select)
        sb = ttk.Scrollbar(lb_frame, orient=tk.VERTICAL,
                            command=self.check_disbursement_listbox.yview)
        sb.grid(row=0, column=1, sticky=(tk.N, tk.S))
        self.check_disbursement_listbox.configure(yscrollcommand=sb.set)

        btn_frame = ttk.Frame(f)
        btn_frame.grid(row=0, column=5, padx=(6, 0), sticky=(tk.N, tk.S))
        ttk.Button(btn_frame, text="全选",
                   command=self._check_select_all_disbursement_sheets,
                   width=6).pack(pady=(0, 2))
        ttk.Button(btn_frame, text="清空",
                   command=self._check_clear_disbursement_selection,
                   width=6).pack(pady=(2, 0))

        self.check_disbursement_tags_frame = tk.Frame(f)
        self.check_disbursement_tags_frame.grid(
            row=1, column=0, columnspan=6, sticky=(tk.W, tk.E), pady=(4, 0))

        # 打款金额列选择
        ttk.Label(f, text="打款金额列：").grid(
            row=2, column=0, padx=(0, 5), pady=(8, 0), sticky=tk.W)
        self.check_amount_col_var = tk.StringVar(value='C')
        ttk.Radiobutton(f, text="C列", variable=self.check_amount_col_var,
                        value='C').grid(row=2, column=1, padx=(0, 5),
                                        pady=(8, 0), sticky=tk.W)
        ttk.Radiobutton(f, text="H列", variable=self.check_amount_col_var,
                        value='H').grid(row=2, column=2, padx=(0, 10),
                                        pady=(8, 0), sticky=tk.W)

        # 月份选择
        ttk.Label(f, text="月份（备注1）：").grid(
            row=3, column=0, padx=(0, 5), pady=(8, 0), sticky=tk.W)
        self.check_month_vars = {}
        self.check_month_frame = ttk.Frame(f)
        self.check_month_frame.grid(row=3, column=1, columnspan=5,
                                     padx=(0, 5), pady=(8, 0), sticky=tk.W)

        # 打款类型选择
        ttk.Label(f, text="打款类型（备注2）：").grid(
            row=4, column=0, padx=(0, 5), pady=(8, 0), sticky=tk.W)
        self.check_type_vars = {}
        self.check_type_frame = ttk.Frame(f)
        self.check_type_frame.grid(row=4, column=1, columnspan=5,
                                    padx=(0, 5), pady=(8, 0), sticky=tk.W)

        # （汇总按钮已移至底部操作按钮栏，避免重复）

    def _check_browse_disbursement(self):
        from tkinter import messagebox
        file_path = filedialog.askopenfilename(
            title="选择发放记录表",
            filetypes=[("Excel文件", "*.xlsx *.xls"), ("所有文件", "*.*")])
        if not file_path:
            return
        self.check_disbursement_path_var.set(file_path)
        self.check_processor.df_disbursement = None
        self.check_processor.output_disbursement = None
        try:
            wb = load_workbook(file_path, read_only=True)
            sheets = wb.sheetnames
            wb.close()
            self.check_disbursement_listbox.delete(0, tk.END)
            for name in sheets:
                self.check_disbursement_listbox.insert(tk.END, name)
            if sheets:
                self.check_disbursement_listbox.selection_set(0)
                self.check_disbursement_listbox.activate(0)
                self._check_on_disbursement_sheet_select()
            self._log_status(
                f"已选择发放记录表：{os.path.basename(file_path)}，共{sheets}个工作簿")
        except Exception as e:
            messagebox.showerror("错误", f"读取文件失败：{e}", parent=self.root)

    def _check_on_disbursement_sheet_select(self, event=None):
        indices = self.check_disbursement_listbox.curselection()
        self.check_disbursement_selected = [
            self.check_disbursement_listbox.get(i) for i in indices]
        for w in self.check_disbursement_tags_frame.winfo_children():
            w.destroy()
        for name in self.check_disbursement_selected:
            tag = tk.Label(self.check_disbursement_tags_frame, text=name,
                           bg="#e0e0e0", relief="raised", bd=1,
                           padx=4, pady=2, font=('Microsoft YaHei UI', 8))
            tag.pack(side=tk.LEFT, padx=(0, 4), pady=2)
        if self.check_disbursement_selected:
            self._check_refresh_disbursement_filters()

    def _check_select_all_disbursement_sheets(self):
        self.check_disbursement_listbox.selection_set(0, tk.END)
        self._check_on_disbursement_sheet_select()

    def _check_clear_disbursement_selection(self):
        self.check_disbursement_listbox.selection_clear(0, tk.END)
        self.check_disbursement_selected = []
        self._check_on_disbursement_sheet_select()

    def _check_refresh_disbursement_filters(self):
        """预扫描所有选中工作簿，合并月份和打款类型"""
        if not hasattr(self, 'check_disbursement_selected') \
           or not self.check_disbursement_selected:
            return
        file_path = self.check_disbursement_path_var.get()
        if not file_path:
            return

        all_months = set()
        all_types = set()
        headers = None

        for sheet_name in self.check_disbursement_selected:
            success, msg = self.check_processor.load_disbursement(file_path, sheet_name)
            if not success:
                continue
            if headers is None:
                headers = self.check_processor.detect_disbursement_headers()
                self._log_status(
                    f"  检测表头：{', '.join(headers.get('all_columns', []))}")

            if 'remark1_col' in headers:
                df = self.check_processor.df_disbursement
                for val in df[headers['remark1_col']].dropna():
                    m = self.check_processor.extract_month_from_remark1(val)
                    if m:
                        all_months.add(m)

            if 'remark2_col' in headers:
                df = self.check_processor.df_disbursement
                for val in df[headers['remark2_col']].dropna():
                    _, t = self.check_processor.extract_company_and_type_from_remark2(val)
                    if t:
                        all_types.add(t)

        # 刷新月份复选框（不默认全选）
        for w in self.check_month_frame.winfo_children():
            w.destroy()
        self.check_month_vars = {}
        if all_months:
            months = sorted(all_months,
                            key=lambda x: int(re.match(r'(\d+)', x).group(1))
                            if re.match(r'(\d+)', x) else 0)
            for i, m in enumerate(months):
                var = tk.BooleanVar(value=False)
                self.check_month_vars[m] = var
                cb = ttk.Checkbutton(self.check_month_frame, text=m, variable=var)
                cb.grid(row=0, column=i, padx=(0, 8), sticky=tk.W)
        else:
            ttk.Label(self.check_month_frame, text="未检测到备注1列",
                      foreground='gray').grid(row=0, column=0)

        # 刷新打款类型复选框（不默认全选）
        for w in self.check_type_frame.winfo_children():
            w.destroy()
        self.check_type_vars = {}
        if all_types:
            order = {"代发打款": 0, "对公打款": 1, "刘先锋现金卡支付": 2}
            types_sorted = sorted(all_types, key=lambda x: order.get(x, 99))
            for i, t in enumerate(types_sorted):
                var = tk.BooleanVar(value=False)
                self.check_type_vars[t] = var
                cb = ttk.Checkbutton(self.check_type_frame, text=t, variable=var)
                cb.grid(row=0, column=i, padx=(0, 8), sticky=tk.W)
        else:
            ttk.Label(self.check_type_frame, text="未检测到备注2列",
                      foreground='gray').grid(row=0, column=0)

    # ────────────────────────────────────────────────
    #  B. 供应商对账单区域
    # ────────────────────────────────────────────────

    def _create_check_statement_frame(self, parent):
        f = ttk.LabelFrame(parent, text="B. 供应商对账单（汇总条件）", padding="12")
        f.grid(row=1, column=0, sticky=(tk.W, tk.E), pady=(0, 10))
        f.columnconfigure(1, weight=1)

        ttk.Label(f, text="供应商对账单：").grid(
            row=0, column=0, padx=(0, 5), sticky=tk.W)
        self.check_statement_path_var = tk.StringVar()
        self.check_statement_entry = ttk.Entry(
            f, textvariable=self.check_statement_path_var)
        self.check_statement_entry.grid(row=0, column=1, padx=(0, 10),
                                         sticky=(tk.W, tk.E))
        ttk.Button(f, text="浏览", command=self._check_browse_statement,
                   width=8).grid(row=0, column=2, padx=(0, 10))
        ttk.Label(f, text="工作簿：").grid(row=0, column=3, padx=(0, 5))
        self.check_statement_sheet_var = tk.StringVar()
        self.check_statement_combo = ttk.Combobox(
            f, textvariable=self.check_statement_sheet_var,
            state='readonly', width=20)
        self.check_statement_combo.grid(row=0, column=4, sticky=(tk.W, tk.E))
        self.check_statement_combo.bind(
            '<<ComboboxSelected>>', self._check_on_statement_sheet_select)

        ttk.Label(f, text="核算年月：").grid(
            row=1, column=0, padx=(0, 5), pady=(8, 0), sticky=tk.W)
        self.check_ym_vars = {}
        self.check_ym_frame = ttk.Frame(f)
        self.check_ym_frame.grid(row=1, column=1, columnspan=4,
                                  padx=(0, 5), pady=(8, 0), sticky=tk.W)

        # （汇总按钮已移至底部操作按钮栏，避免重复）

    def _check_browse_statement(self):
        from tkinter import messagebox
        file_path = filedialog.askopenfilename(
            title="选择供应商对账单",
            filetypes=[("Excel文件", "*.xlsx *.xls"), ("所有文件", "*.*")])
        if not file_path:
            return
        self.check_statement_path_var.set(file_path)
        self.check_processor.df_statement = None
        self.check_processor.output_statement = None
        try:
            wb = load_workbook(file_path, read_only=True)
            sheets = wb.sheetnames
            wb.close()
            self.check_statement_combo['values'] = sheets
            if sheets:
                self.check_statement_combo.current(0)
                self._check_on_statement_sheet_select()
            self._log_status(
                f"已选择供应商对账单：{os.path.basename(file_path)}，共{sheets}个工作簿")
        except Exception as e:
            messagebox.showerror("错误", f"读取文件失败：{e}", parent=self.root)

    def _check_on_statement_sheet_select(self, event=None):
        file_path = self.check_statement_path_var.get()
        sheet_name = self.check_statement_sheet_var.get()
        if not file_path or not sheet_name:
            return
        success, msg = self.check_processor.load_statement(file_path, sheet_name)
        if not success:
            self._log_status(f"  加载对账单失败：{msg}")
            return
        self._log_status(f"  {msg}")
        yms = self.check_processor.get_unique_year_months()
        for w in self.check_ym_frame.winfo_children():
            w.destroy()
        self.check_ym_vars = {}
        for i, ym in enumerate(yms):
            var = tk.BooleanVar(value=True)
            self.check_ym_vars[ym] = var
            cb = ttk.Checkbutton(self.check_ym_frame, text=ym, variable=var)
            cb.grid(row=0, column=i, padx=(0, 8), sticky=tk.W)

    # ────────────────────────────────────────────────
    #  C. 操作按钮栏（Tab 2 专用）
    # ────────────────────────────────────────────────

    def _create_check_button_frame(self, parent):
        f = ttk.Frame(parent)
        f.grid(row=2, column=0, sticky=(tk.W, tk.E), pady=(0, 10))
        ttk.Button(f, text="📊 汇总发放记录",
                   command=self._check_process_disbursement,
                   width=16).grid(row=0, column=0, padx=(0, 8))
        ttk.Button(f, text="📊 汇总对账单",
                   command=self._check_process_statement,
                   width=16).grid(row=0, column=1, padx=(0, 8))
        ttk.Button(f, text="📋 开始核对",
                   command=self._check_reconcile,
                   width=14).grid(row=0, column=2, padx=(0, 8))
        ttk.Button(f, text="📂 打开输出目录",
                   command=self._open_output_dir,
                   width=16).grid(row=0, column=3, padx=(0, 8))
        ttk.Button(f, text="🔄 重置",
                   command=self._clear_all_check,
                   width=16).grid(row=0, column=4, padx=(0, 8))

    def _clear_all_check(self):
        """重置 Tab 2 的所有状态"""
        self.check_disbursement_path_var.set("")
        self.check_disbursement_listbox.delete(0, tk.END)
        self.check_disbursement_selected = []
        for w in self.check_disbursement_tags_frame.winfo_children():
            w.destroy()
        for w in self.check_month_frame.winfo_children():
            w.destroy()
        for w in self.check_type_frame.winfo_children():
            w.destroy()
        self.check_month_vars = {}
        self.check_type_vars = {}
        self.check_statement_path_var.set("")
        self.check_statement_combo['values'] = []
        self.check_statement_sheet_var.set("")
        for w in self.check_ym_frame.winfo_children():
            w.destroy()
        self.check_ym_vars = {}
        self.check_processor = CheckProcessor()
        self.check_reconcile_path = None
        for item in self.tree.get_children():
            self.tree.delete(item)
        self.tree['columns'] = ()
        self.status_text.delete('1.0', tk.END)
        self._log_status("核对 Tab 已重置。")

    # ────────────────────────────────────────────────
    #  D. 汇总动作
    # ────────────────────────────────────────────────

    def _check_process_disbursement(self):
        from tkinter import messagebox
        if not hasattr(self, 'check_disbursement_selected') \
           or not self.check_disbursement_selected:
            messagebox.showwarning("警告", "请先选择发放记录表并选择工作簿",
                                  parent=self.root)
            return
        if not self.check_disbursement_path_var.get():
            messagebox.showwarning("警告", "请先选择发放记录表文件", parent=self.root)
            return

        file_path = self.check_disbursement_path_var.get()
        amount_choice = self.check_amount_col_var.get()
        selected_months = [m for m, var in self.check_month_vars.items() if var.get()]
        selected_types = [t for t, var in self.check_type_vars.items() if var.get()]

        self._log_status("=" * 50)
        self._log_status(
            f"开始汇总发放记录（共{len(self.check_disbursement_selected)}个工作簿）...")

        all_outputs = []
        headers = None
        for idx, sheet_name in enumerate(self.check_disbursement_selected):
            self._log_status(f"--- 处理 [{idx+1}]：{sheet_name} ---")
            success, msg = self.check_processor.load_disbursement(file_path, sheet_name)
            if not success:
                self._log_status(f"  跳过：{msg}")
                continue
            if headers is None:
                headers = self.check_processor.detect_disbursement_headers()
            amount_col = headers.get(
                'amount_col_c' if amount_choice == 'C' else 'amount_col_h', '')
            if not amount_col:
                self._log_status(
                    f"  未找到{'C' if amount_choice == 'C' else 'H'}列对应的表头")
                continue
            self._log_status(f"  使用金额列：{amount_col}")
            self._log_status(
                f"  过滤月份：{', '.join(selected_months) if selected_months else '全部'}")
            self._log_status(
                f"  过滤类型：{', '.join(selected_types) if selected_types else '全部'}")

            success2, msg2 = self.check_processor.process_disbursement(
                amount_col, selected_months, selected_types)
            if not success2:
                self._log_status(f"  处理失败：{msg2}")
                continue
            all_outputs.append(self.check_processor.output_disbursement)
            self._log_status(
                f"  ✓ 处理完成，{len(self.check_processor.output_disbursement)}行")

        if not all_outputs:
            messagebox.showwarning("警告", "没有成功处理任何工作簿")
            return

        merged = pd.concat(all_outputs, ignore_index=True)
        numeric_cols = merged.select_dtypes(include=['number']).columns.tolist()
        agg_dict = {col: 'sum' for col in numeric_cols}
        for col in merged.columns:
            if col not in numeric_cols and col != '供应商名称':
                agg_dict[col] = 'first'
        if '供应商名称' in merged.columns and numeric_cols:
            merged = merged.groupby('供应商名称', as_index=False).agg(agg_dict)
            merged = merged.reindex(columns=all_outputs[0].columns, fill_value=0)

        self.check_processor.output_disbursement = merged
        self._log_status(f"合并完成，共{len(merged)}行")
        self._log_status("=" * 50)
        # 不再自动保存单独文件，等「开始核对」后统一输出

    def _check_process_statement(self):
        from tkinter import messagebox
        if not self.check_statement_path_var.get():
            messagebox.showwarning("警告", "请先选择供应商对账单文件", parent=self.root)
            return
        selected_yms = [ym for ym, var in self.check_ym_vars.items() if var.get()]
        if not selected_yms:
            messagebox.showwarning("警告", "请至少选择一个核算年月", parent=self.root)
            return

        self._log_status("=" * 50)
        self._log_status(
            f"开始汇总对账单（核算年月：{', '.join(selected_yms)}）...")
        success, msg = self.check_processor.process_statement(selected_yms)
        if not success:
            messagebox.showerror("错误", msg, parent=self.root)
            self._log_status(f"错误：{msg}")
            return
        self._log_status(msg)
        self._log_status("=" * 50)
        # 不再自动保存单独文件，等「开始核对」后统一输出

    def _check_reconcile(self):
        from tkinter import messagebox
        if self.check_processor.output_disbursement is None:
            messagebox.showwarning("警告", "请先完成「汇总发放记录」", parent=self.root)
            return
        if self.check_processor.output_statement is None:
            messagebox.showwarning("警告", "请先完成「汇总对账单」", parent=self.root)
            return

        self._log_status("=" * 50)
        self._log_status("开始核对...")

        success, msg, result_df = self.check_processor.reconcile()
        if not success:
            messagebox.showerror("错误", msg, parent=self.root)
            self._log_status(f"错误：{msg}")
            return

        stats = self.check_processor.reconcile_stats
        self._log_status(msg)

        self._log_status(f"\n--- 核对统计 ---")
        self._log_status(f"  已核对：{stats['matched']} 条")
        if stats['pending'] > 0:
            self._log_status(
                f"  待核对：{stats['pending']} 条（差额合计：{stats['pending_diff']:.2f}）")
        else:
            self._log_status("  待核对：0 条")
        self._log_status(
            f"  未匹配 -> 发放记录表：{stats['unmatched_a']} 条"
            f"（金额合计：{stats['unmatched_a_amt']:.2f}）")
        self._log_status(
            f"  未匹配 -> 对账单表：{stats['unmatched_b']} 条"
            f"（金额合计：{stats['unmatched_b_amt']:.2f}）")
        total = stats['matched'] + stats['pending'] + stats['unmatched_a'] + stats['unmatched_b']
        self._log_status(f"  总计涉及：{total} 条")
        self._log_status("=" * 50)

        self._check_save_reconcile_excel(result_df)

        messagebox.showinfo(
            "提示",
            "已完成核对，请点击「打开输出目录」找到文件位置。\n\n文件：核对结果.xlsx（含3个Sheet）",
            parent=self.root
        )

    # ────────────────────────────────────────────────
    #  E. 预览 & 自动保存 & Excel 输出
    # ────────────────────────────────────────────────

    def _check_show_preview(self, df, title):
        for item in self.tree.get_children():
            self.tree.delete(item)
        if df is None or df.empty:
            return
        columns = list(df.columns)
        self.tree['columns'] = columns
        for col in columns:
            self.tree.heading(col, text=col)
            if col == '供应商名称':
                self.tree.column(col, width=180, anchor=tk.W)
            elif col == '公司':
                self.tree.column(col, width=250, anchor=tk.W)
            elif col == '总计':
                self.tree.column(col, width=100, anchor=tk.E)
            elif col == '备注':
                self.tree.column(col, width=30, anchor=tk.CENTER)
            else:
                self.tree.column(col, width=90, anchor=tk.E)
        display_rows = min(50, len(df))
        for row in df.head(display_rows).itertuples(index=False):
            processed_row = tuple(
                '' if (isinstance(v, (int, float)) and v == 0)
                else v if pd.notna(v)
                else ''
                for v in row
            )
            self.tree.insert('', tk.END, values=processed_row)
        self.tree.update_idletasks()
        self.tree.xview_moveto(0)
        self._log_status(f"  ✓ {title}预览显示前{display_rows}行，共{len(df)}行")

    def _check_auto_save_disbursement(self):
        if self.check_processor.output_disbursement is None:
            return
        file_path = self.check_disbursement_path_var.get() or os.getcwd()
        base_dir = os.path.dirname(file_path) if file_path else os.getcwd()
        save_path = os.path.join(base_dir, "发放记录汇总.xlsx")
        try:
            df = self.check_processor.output_disbursement.copy()
            for col in df.select_dtypes(include=['number']).columns:
                df[col] = df[col].replace(0, None)
            df.to_excel(save_path, engine='openpyxl',
                        sheet_name='发放记录汇总', index=False)
            wb = load_workbook(save_path)
            apply_excel_styles(wb['发放记录汇总'],
                               self.check_processor.output_disbursement)
            wb.save(save_path)
            wb.close()
            self._log_status(f"  ✓ 已自动保存：发放记录汇总.xlsx")
        except Exception as e:
            self._log_status(f"  ⚠ 自动保存失败：{e}")

    def _check_auto_save_statement(self):
        if self.check_processor.output_statement is None:
            return
        file_path = self.check_statement_path_var.get() or os.getcwd()
        base_dir = os.path.dirname(file_path) if file_path else os.getcwd()
        save_path = os.path.join(base_dir, "对账单汇总.xlsx")
        try:
            df = self.check_processor.output_statement.copy()
            for col in df.select_dtypes(include=['number']).columns:
                df[col] = df[col].replace(0, None)
            df.to_excel(save_path, engine='openpyxl',
                        sheet_name='对账单汇总', index=False)
            wb = load_workbook(save_path)
            apply_excel_styles(wb['对账单汇总'],
                               self.check_processor.output_statement)
            wb.save(save_path)
            wb.close()
            self._log_status(f"  ✓ 已自动保存：对账单汇总.xlsx")
        except Exception as e:
            self._log_status(f"  ⚠ 自动保存失败：{e}")

    def _check_save_reconcile_excel(self, result_df):
        """输出核对结果 Excel（3 个 Sheet）"""
        disbursement_path = self.check_disbursement_path_var.get()
        statement_path = self.check_statement_path_var.get()
        base_dir = os.path.dirname(disbursement_path or statement_path) \
            if (disbursement_path or statement_path) else os.getcwd()
        save_path = os.path.join(base_dir, "核对结果.xlsx")

        try:
            df_dis = self.check_processor.output_disbursement.copy()
            df_stmt = self.check_processor.output_statement.copy()

            for df in [df_dis, df_stmt]:
                for col in df.select_dtypes(include=['number']).columns:
                    df[col] = df[col].replace(0, None)

            with pd.ExcelWriter(save_path, engine='openpyxl') as writer:
                df_dis.to_excel(writer, sheet_name='发放记录汇总', index=False)
                df_stmt.to_excel(writer, sheet_name='对账单汇总', index=False)
                result_df.to_excel(writer, sheet_name='核对结果', index=False)

            wb = load_workbook(save_path)
            apply_excel_styles(wb['发放记录汇总'],
                               self.check_processor.output_disbursement)
            apply_excel_styles(wb['对账单汇总'],
                               self.check_processor.output_statement)

            ws_result = wb['核对结果']
            apply_excel_styles(ws_result, result_df)

            fill_green = PatternFill(start_color="C6EFCE", end_color="C6EFCE",
                                      fill_type="solid")
            fill_yellow = PatternFill(start_color="FFEB9C", end_color="FFEB9C",
                                       fill_type="solid")

            header_cells = list(ws_result[1])
            status_col_idx = None
            for ci, cell in enumerate(header_cells, 1):
                if cell.value == '状态':
                    status_col_idx = ci
                    break

            if status_col_idx:
                for row in range(2, ws_result.max_row + 1):
                    status_val = ws_result.cell(row=row,
                                                 column=status_col_idx).value
                    if status_val == '已核对':
                        fill = fill_green
                        for col in range(1, len(header_cells) + 1):
                            ws_result.cell(row=row, column=col).fill = fill
                    elif status_val == '待核对':
                        fill = fill_yellow
                        for col in range(1, len(header_cells) + 1):
                            ws_result.cell(row=row, column=col).fill = fill

            wb.save(save_path)
            wb.close()
            self.check_reconcile_path = save_path
            self._log_status(f"  ✓ 已保存核对结果：{os.path.basename(save_path)}")
        except Exception as e:
            import traceback
            self._log_status(f"  ⚠ 保存失败：{e}")
            self._log_status(traceback.format_exc())
