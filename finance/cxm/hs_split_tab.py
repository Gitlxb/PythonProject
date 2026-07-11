# -*- coding: utf-8 -*-
"""
Tab 1「拆分汇总比较」UI Mixin
包含：文件选择、工作簿多选、按钮栏、文件浏览等所有 Tab 1 专属方法

通过多重继承混入 SummaryProcessorApp：
  class SummaryProcessorApp(GuiActionsMixin, SplitTabMixin, CheckTabMixin):
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import os
import pandas as pd
from openpyxl import load_workbook


class SplitTabMixin:
    """Tab 1「拆分汇总比较」的 UI 创建和事件处理"""

    # ==================== Tab 1 内容组装 ====================

    def _create_tab1_content(self, parent):
        """构建 Tab 1（拆分汇总比较）的全部 UI

        布局 grid 行分配：
          row=0  文件选择区
          row=1  按钮栏
          row=2  数据预览（共用 Treeview）
          row=3  状态信息（共用 ScrolledText）
        """
        parent.columnconfigure(0, weight=1)
        parent.rowconfigure(2, weight=3)
        parent.rowconfigure(3, weight=1)

        self._create_file_frame(parent)                # row=0
        self._create_button_frame(parent)               # row=1
        self._create_preview_frame(parent, row=2)       # row=2
        self._create_status_frame(parent, row=3)        # row=3

    # ==================== 文件选择区 ====================

    def _create_file_frame(self, parent):
        file_frame = ttk.LabelFrame(parent, text="文件选择", padding="12")
        file_frame.grid(row=0, column=0, sticky=(tk.W, tk.E), pady=(0, 12))
        file_frame.columnconfigure(1, weight=1)
        file_frame.columnconfigure(4, weight=1)

        # 第1行：代工费发放记录表
        ttk.Label(file_frame, text="代工费发放记录表：").grid(
            row=0, column=0, padx=(0, 5), sticky=tk.W)
        self.file_entry = ttk.Entry(file_frame, textvariable=self.file_path_var)
        self.file_entry.grid(row=0, column=1, padx=(0, 10), sticky=(tk.W, tk.E))
        ttk.Button(file_frame, text="浏览", command=self._browse_file, width=8).grid(
            row=0, column=2, padx=(0, 10))
        ttk.Label(file_frame, text="工作簿（可多选）：").grid(
            row=0, column=3, padx=(0, 5))

        sheet_list_frame = ttk.Frame(file_frame)
        sheet_list_frame.grid(row=0, column=4, sticky=(tk.W, tk.E, tk.N, tk.S))
        sheet_list_frame.columnconfigure(0, weight=1)
        sheet_list_frame.rowconfigure(0, weight=1)

        self.sheet_listbox = tk.Listbox(
            sheet_list_frame, selectmode=tk.EXTENDED, height=5,
            font=('Microsoft YaHei UI', 9), exportselection=False)
        self.sheet_listbox.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        self.sheet_listbox.bind('<<ListboxSelect>>', self._on_sheet_select)

        self.sheet_scrollbar = ttk.Scrollbar(sheet_list_frame, orient=tk.VERTICAL,
                                              command=self.sheet_listbox.yview)
        self.sheet_scrollbar.grid(row=0, column=1, sticky=(tk.N, tk.S))
        self.sheet_listbox.configure(yscrollcommand=self.sheet_scrollbar.set)

        sheet_btn_frame = ttk.Frame(file_frame)
        sheet_btn_frame.grid(row=0, column=5, padx=(6, 0), sticky=(tk.N, tk.S))
        ttk.Button(sheet_btn_frame, text="全选",
                   command=self._select_all_sheets, width=6).pack(pady=(0, 2))
        ttk.Button(sheet_btn_frame, text="清空",
                   command=self._clear_sheet_selection, width=6).pack(pady=(2, 0))

        self.selected_tags_frame = tk.Frame(file_frame)
        self.selected_tags_frame.grid(row=1, column=0, columnspan=6,
                                     sticky=(tk.W, tk.E), pady=(4, 0))

        # 第2行：代工费汇总表
        ttk.Label(file_frame, text="代工费汇总表：").grid(
            row=2, column=0, padx=(0, 5), pady=(8, 0), sticky=tk.W)
        self.check_file_entry = ttk.Entry(file_frame,
                                          textvariable=self.check_file_path_var)
        self.check_file_entry.grid(row=2, column=1, padx=(0, 10),
                                   pady=(8, 0), sticky=(tk.W, tk.E))
        ttk.Button(file_frame, text="浏览", command=self._browse_check_file,
                   width=8).grid(row=2, column=2, padx=(0, 10), pady=(8, 0))
        ttk.Label(file_frame, text="工作簿：").grid(
            row=2, column=3, padx=(0, 5), pady=(8, 0))
        self.check_sheet_combo = ttk.Combobox(file_frame,
                                              textvariable=self.check_sheet_var,
                                              state='readonly', width=20)
        self.check_sheet_combo.grid(row=2, column=4, pady=(8, 0),
                                    sticky=(tk.W, tk.E))

        # 第3行：追加到目标表
        ttk.Label(file_frame, text="追加到目标表：").grid(
            row=3, column=0, padx=(0, 5), pady=(8, 0), sticky=tk.W)
        self.target_file_entry = ttk.Entry(file_frame,
                                           textvariable=self.target_file_path_var)
        self.target_file_entry.grid(row=3, column=1, padx=(0, 10),
                                    pady=(8, 0), sticky=(tk.W, tk.E))
        ttk.Button(file_frame, text="浏览", command=self._browse_target_file,
                   width=8).grid(row=3, column=2, padx=(0, 10), pady=(8, 0))

    # ==================== Sheet 多选辅助 ====================

    def _on_sheet_select(self, event=None):
        selected_indices = self.sheet_listbox.curselection()
        self.selected_sheets = [self.sheet_listbox.get(i) for i in selected_indices]
        self._update_selected_tags()

    def _update_selected_tags(self):
        for widget in self.selected_tags_frame.winfo_children():
            widget.destroy()
        if not self.selected_sheets:
            return
        for name in self.selected_sheets:
            tag = tk.Label(self.selected_tags_frame, text=name,
                           bg="#e0e0e0", relief="raised", bd=1,
                           padx=4, pady=2, font=('Microsoft YaHei UI', 8))
            tag.pack(side=tk.LEFT, padx=(0, 4), pady=2)

    def _select_all_sheets(self):
        self.sheet_listbox.selection_set(0, tk.END)
        self._on_sheet_select()

    def _clear_sheet_selection(self):
        self.sheet_listbox.selection_clear(0, tk.END)
        self.selected_sheets = []
        self._update_selected_tags()

    # ==================== 按钮栏 ====================

    def _create_button_frame(self, parent):
        button_frame = ttk.Frame(parent)
        button_frame.grid(row=1, column=0, sticky=(tk.W, tk.E), pady=(0, 12))

        ttk.Button(button_frame, text="📊数据汇总",
                   command=self._process_data, width=14).grid(
                       row=0, column=0, padx=(0, 8))
        ttk.Button(button_frame, text="✅核对",
                   command=self._check_data, width=14).grid(
                       row=0, column=1, padx=(0, 8))
        ttk.Button(button_frame, text="🔗追加到目标表",
                   command=self._append_to_target, width=14).grid(
                       row=0, column=2, padx=(0, 8))
        ttk.Button(button_frame, text="📂打开输出目录",
                   command=self._open_output_dir, width=14).grid(
                       row=0, column=3, padx=(0, 8))
        ttk.Button(button_frame, text="🔄重置",
                   command=self._clear_all, width=14).grid(
                       row=0, column=5)

    # ==================== 文件浏览 ====================

    def _browse_file(self):
        file_path = filedialog.askopenfilename(
            title="选择Excel文件",
            filetypes=[("Excel文件", "*.xlsx *.xls"), ("所有文件", "*.*")])
        if not file_path:
            return
        self._log_status(f"正在读取文件：{os.path.basename(file_path)}...")
        success, message, sheets = self.processor.load_file(file_path)
        if success:
            self.file_path_var.set(file_path)
            self.sheet_listbox.delete(0, tk.END)
            for name in sheets:
                self.sheet_listbox.insert(tk.END, name)
            if sheets:
                self.sheet_listbox.selection_set(0)
                self.sheet_listbox.activate(0)
                self._on_sheet_select()
            self._log_status(message)
            self._log_status(f"包含{len(sheets)}个工作表，默认选中第一个")
        else:
            messagebox.showerror("错误", message)
            self._log_status(f"错误：{message}")

    def _browse_check_file(self):
        file_path = filedialog.askopenfilename(
            title="选择代工费汇总表",
            filetypes=[("Excel文件", "*.xlsx *.xls"), ("所有文件", "*.*")])
        if not file_path:
            return
        self.check_file_path = file_path
        self.check_file_path_var.set(file_path)
        try:
            wb = load_workbook(file_path, read_only=True)
            sheets = wb.sheetnames
            wb.close()
            self.check_sheet_combo['values'] = sheets
            if sheets:
                self.check_sheet_combo.current(0)
            self._log_status(f"已选择代工费汇总表：{os.path.basename(file_path)}")
            self._log_status(f"  包含 {len(sheets)} 个工作表")
        except Exception as e:
            messagebox.showerror("错误", f"读取文件失败：{e}", parent=self.root)
            self._log_status(f"错误：读取代工费汇总表失败 {e}")

    def _browse_target_file(self):
        file_path = filedialog.askopenfilename(
            title="选择要追加到的目标Excel文件",
            filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")])
        if not file_path:
            return
        self.target_file_path = file_path
        self.target_file_path_var.set(file_path)
        self.appended_types = set()
        self._log_status(f"已选择目标文件：{os.path.basename(file_path)}")
