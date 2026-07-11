# -*- coding: utf-8 -*-
"""
GUI 业务逻辑模块 - SummaryProcessorApp 的数据处理动作
从 hs_gui.py 拆分，包含：数据汇总、核对、预览、追加、自动保存

使用方式：在 hs_gui.py 中通过 Mixin 或直接将方法挂载到 SummaryProcessorApp 类上
推荐：from finance.cxm.hs_gui_actions import *
"""

import pandas as pd
import os
import re
from typing import List, Optional
from openpyxl.utils import get_column_letter

from finance.cxm.hs_huizong_cxm import SummaryProcessor
from finance.cxm.hs_checker import DataChecker
from finance.cxm.hs_utils import apply_excel_styles
from openpyxl import load_workbook
from finance.cxm.dialogs import show_selection_dialog, show_check_condition_dialog


# ==================== 工具函数 ====================

def _month_sort_key(x: str):
    """月份排序 key：单月份 (0, num)，范围月份 (1, start_num)"""
    if '-' in x:
        start = int(x.split('-')[0])
        return (1, start)
    return (0, int(x.replace('月', '')))


# ==================== Mixin 类 ====================

class GuiActionsMixin:
    """业务逻辑 Mixin 类，与 hs_gui.py 中的 SummaryProcessorApp 配合使用
    
    将此类的方法混入 SummaryProcessorApp 即可：
    class SummaryProcessorApp(GuiActionsMixin):
        ...
    """

    # ==================== 数据汇总（支持多工作簿） ====================

    def _process_data(self):
        file_path = self.file_path_var.get()

        if not file_path:
            from tkinter import messagebox
            messagebox.showwarning("警告", "请先选择Excel文件")
            return
        if not self.selected_sheets:
            from tkinter import messagebox
            messagebox.showwarning("警告", "请至少选择一个工作簿")
            return

        self._log_status("=" * 50)
        self._log_status(f"开始数据汇总（共{len(self.selected_sheets)}个工作簿）...")

        # ===== 预扫描：合并所有工作簿的可选数据 =====
        self._log_status("预检测所有工作簿的可选数据...")
        all_payment_types = set()
        all_months = set()
        first_amount_columns = []

        for idx, sheet_name in enumerate(self.selected_sheets):
            success, msg = self.processor.load_sheet(sheet_name)
            if not success:
                self._log_status(f"   跳过工作簿 [{idx+1}]：{sheet_name}（{msg}）")
                continue

            pts = self.processor.detect_payment_types()
            all_payment_types.update(pts)

            ms = self.processor.detect_months()
            all_months.update(ms)

            if idx == 0:
                first_amount_columns = self.processor.detect_amount_columns()

        # 排序合并结果
        order = {"刘先锋现金卡": 0, "对公打款": 1, "代发打款": 2}
        payment_types = sorted(list(all_payment_types), key=lambda x: order.get(x, 99))
        months = sorted(list(all_months), key=_month_sort_key)
        amount_columns = first_amount_columns

        # ===== 弹窗选择（仅一次）=====
        has_multi_amount = len(amount_columns) >= 2
        needs_dialog = payment_types or months or has_multi_amount

        selected_types = None
        selected_months = None
        selected_amount_col_idx = None

        if needs_dialog:
            self._log_status(f"   检测到打款类型：{', '.join(payment_types) if payment_types else '无'}")
            self._log_status(f"   检测到月份：{', '.join(months) if months else '无'}")
            if amount_columns:
                cols_desc = ', '.join([f'{col_letter}列' for _, col_letter in amount_columns])
                self._log_status(f"   检测到打款金额列：{cols_desc}")

            selected_types, selected_months, selected_amount_col_idx = show_selection_dialog(
                self.root, payment_types, months, amount_columns
            )
            if selected_types is None:
                self._log_status("用户取消了选择，中止处理")
                return

            self.selected_months_for_save = selected_months
            self._log_status(f"   已选择打款类型：{', '.join(selected_types) if selected_types else '无'}")
            self._log_status(f"   已选择月份：{', '.join(selected_months) if selected_months else '无'}")
            if selected_amount_col_idx is not None:
                self._log_status(f"   已选择打款金额列：{get_column_letter(selected_amount_col_idx + 1)}列")
        else:
            self._log_status("   未检测到可选项，将处理全部数据")

        # ===== 正式处理每个工作簿 =====
        all_output_dfs = []
        for idx, sheet_name in enumerate(self.selected_sheets):
            self._log_status(f"\n--- 处理工作簿 [{idx+1}/{len(self.selected_sheets)}]：{sheet_name} ---")

            self._log_status(f"1. 加载工作表：{sheet_name}")
            success, message = self.processor.load_sheet(sheet_name)
            if not success:
                self._log_status(f"   错误：{message}，跳过此工作簿")
                continue
            self._log_status(f"   ✓ {message}")

            # 检测必要列
            self._log_status("2. 检测列...")
            cols_ok = True
            for col in ["备注2", "供应商名称", "收款姓名"]:
                if col not in self.processor.data_frame.columns:
                    self._log_status(f"   错误：未找到'{col}'列，跳过此工作簿")
                    cols_ok = False
                    break
                self._log_status(f"   ✓ 找到'{col}'列")

            if not cols_ok:
                continue

            self._log_status("3. 应用缩写规则...")
            success, message = self.processor.process_data(
                selected_types, selected_months, selected_amount_col_idx
            )
            if not success:
                self._log_status(f"   错误：{message}，跳过此工作簿")
                continue
            self._log_status(f"   ✓ {message}")

            # 收集本工作簿的处理结果
            if self.processor.output_data is not None:
                all_output_dfs.append(self.processor.output_data.copy())
                self._log_status(f"   已收集 {len(self.processor.output_data)} 行数据")

        if not all_output_dfs:
            from tkinter import messagebox
            messagebox.showwarning("警告", "没有成功处理任何工作簿")
            self._log_status("错误：没有成功处理任何工作簿")
            return

        # 合并所有工作簿的结果
        self._log_status(f"\n4. 合并 {len(all_output_dfs)} 个工作簿的结果...")
        merged_df = pd.concat(all_output_dfs, ignore_index=True)

        # 同名供应商+相同公司金额累加（同名不同公司的保留为不同行）
        self._log_status("5. 同名供应商金额累加...")
        numeric_cols = merged_df.select_dtypes(include=['number']).columns.tolist()
        if '供应商名称' in merged_df.columns and '公司' in merged_df.columns and numeric_cols:
            agg_dict = {col: 'sum' for col in numeric_cols}
            for col in merged_df.columns:
                if col not in numeric_cols and col not in ['供应商名称', '公司']:
                    agg_dict[col] = 'first'
            merged_df = merged_df.groupby(['供应商名称', '公司'], as_index=False).agg(agg_dict)
            self._log_status(f"   ✓ 合并后共 {len(merged_df)} 行（同名同公司已累加，不同公司保留独立行）")
        else:
            self._log_status(f"   ✓ 合并后共 {len(merged_df)} 行")

        self.processor.output_data = merged_df

        # 列标准化：确保列顺序正确，处理备注列
        self._log_status("6. 标准化输出表结构...")
        df = self.processor.output_data

        # 处理备注列：将空列名统一改为"备注"
        if '' in df.columns:
            df = df.rename(columns={'': '备注'})
        elif '（空白）' in df.columns:
            df = df.rename(columns={'（空白）': '备注'})
        elif '备注' not in df.columns:
            df['备注'] = ''
        # 确保备注列值为空白
        df['备注'] = ''

        # 标准化列顺序：供应商名称 + 公司列 + 总计 + 备注 + 公司
        fixed_cols = ['供应商名称', '总计', '备注', '公司']
        company_cols = [c for c in df.columns if c not in fixed_cols]
        company_cols = [c for c in company_cols if c not in ['', '（空白）']]
        final_columns = ['供应商名称'] + company_cols + ['总计', '备注', '公司']
        final_columns = [c for c in final_columns if c in df.columns]
        df = df.reindex(columns=final_columns)
        self.processor.output_data = df
        self._log_status(f"   ✓ 列标准化完成，共 {len(df.columns)} 列")

        self._log_status("7. 生成汇总表...")
        self._show_preview()

        # 自动保存供应商拆分汇总
        self._log_status("8. 自动保存结果...")
        self._auto_save_all()
        self.last_operation = 'summary'
        self._log_status("数据汇总完成！")
        self._log_status("=" * 50)

    # ==================== 数据预览 ====================

    def _show_preview(self):
        import tkinter as tk
        for item in self.tree.get_children():
            self.tree.delete(item)
        if self.processor.output_data is None:
            return

        df = self.processor.output_data
        columns = list(df.columns)
        self.tree['columns'] = columns

        for col in columns:
            self.tree.heading(col, text=col)
            if col == '供应商名称':
                self.tree.column(col, width=180, anchor=tk.W)
            elif col == '公司':
                self.tree.column(col, width=250, anchor=tk.W)
            elif col == '总计':
                self.tree.column(col, width=100, anchor=tk.E)
            elif col == '备注':
                self.tree.column(col, width=30, anchor=tk.CENTER)
            else:
                self.tree.column(col, width=90, anchor=tk.E)

        display_rows = min(50, len(df))
        for row in df.head(display_rows).itertuples(index=False):
            processed_row = tuple(
                '' if (isinstance(v, (int, float)) and v == 0)
                else v if pd.notna(v)
                else ''
                for v in row
            )
            self.tree.insert('', tk.END, values=processed_row)

        self.tree.update_idletasks()
        self.tree.xview_moveto(0)
        self._log_status(f"   ✓ 预览显示前{display_rows}行数据")
        self._log_status(f"   ✓ 共{len(df)}行，{len(columns)}列")

    def _show_check_preview(self, result_df: pd.DataFrame):
        import tkinter as tk
        for item in self.tree.get_children():
            self.tree.delete(item)
        if result_df is None or result_df.empty:
            return

        columns = list(result_df.columns)
        self.tree['columns'] = columns
        for col in columns:
            self.tree.heading(col, text=col)
            if col in ('介绍人', '供应商名称'):
                self.tree.column(col, width=150, anchor=tk.W)
            elif col == '状态':
                self.tree.column(col, width=80, anchor=tk.CENTER)
            else:
                self.tree.column(col, width=120, anchor=tk.E)

        display_rows = min(50, len(result_df))
        for row in result_df.head(display_rows).itertuples(index=False):
            processed_row = tuple(str(v) if pd.notna(v) else '' for v in row)
            self.tree.insert('', tk.END, values=processed_row)

        self.tree.update_idletasks()
        self.tree.xview_moveto(0)
        self._log_status(f"   ✓ 核对结果预览显示前{display_rows}行")

    # ==================== 核对数据 ====================

    def _check_data(self):
        from tkinter import messagebox
        if self.processor.output_data is None:
            messagebox.showwarning("警告", "请先进行数据汇总", parent=self.root)
            return
        if not self.check_file_path or not os.path.exists(self.check_file_path):
            messagebox.showwarning("警告", "请先选择代工费汇总表", parent=self.root)
            return

        self._log_status("=" * 50)
        self._log_status("开始核对数据...")

        check_sheet_name = self.check_sheet_var.get()
        if not check_sheet_name:
            messagebox.showwarning("警告", "请选择代工费汇总表的工作簿", parent=self.root)
            return

        self._log_status(f"使用代工费汇总表：{os.path.basename(self.check_file_path)} [{check_sheet_name}]")

        # 加载待核对文件
        self._log_status("正在加载代工费汇总表数据（快速模式）...")
        df, month_col, person_col, amount_col, msg = DataChecker.load_check_file(
            self.check_file_path, check_sheet_name)
        if df is None:
            messagebox.showerror("错误", msg, parent=self.root)
            self._log_status(f"错误：{msg}")
            return
        self._log_status(f"加载成功，{msg}")
        self._log_status(f"识别到列：月份={month_col}, 介绍人={person_col}, 金额={amount_col}")

        # 加载基准表
        self._log_status("正在加载基准表（供应商汇总）...")
        checker = DataChecker("")
        success, msg = checker.load_reference_from_dataframe(self.processor.output_data)
        if not success:
            messagebox.showerror("错误", msg, parent=self.root)
            self._log_status(f"错误：{msg}")
            return
        self._log_status(msg)

        # 获取唯一值
        self._log_status("正在检测月份和介绍人...")
        months = checker.get_unique_months(df, month_col)
        persons = checker.get_unique_persons(df, person_col)
        self._log_status(f"检测到月份：{', '.join(months) if months else '无'}")
        self._log_status(f"检测到介绍人：{len(persons)}个")

        if not months:
            messagebox.showwarning("警告", "未检测到任何月份数据", parent=self.root)
            return
        if not persons:
            messagebox.showwarning("警告", "未检测到任何介绍人数据", parent=self.root)
            return

        # 计算月份-介绍人映射
        month_person_map = {m: set() for m in months}
        for val, person in zip(df[month_col], df[person_col]):
            if pd.isna(val) or pd.isna(person):
                continue
            val_str = str(val).strip()
            person_str = str(person).strip()
            if not person_str:
                continue
            for m in months:
                if re.search(rf'(?<!\d){re.escape(m)}(?!\d)', val_str):
                    month_person_map[m].add(person_str)
        month_person_map = {m: sorted(list(ps)) for m, ps in month_person_map.items()}

        # 弹窗选择
        self._log_status("打开条件选择弹窗...")
        selected_month, selected_persons = show_check_condition_dialog(
            self.root, months, month_person_map)
        if selected_month is None:
            self._log_status("用户取消了条件选择")
            return

        self._log_status(f"已选择月份：{selected_month}")
        self._log_status(f"已选择介绍人：{', '.join(selected_persons)}")

        # 执行核对
        self._log_status("执行核对...")
        success, msg, result_df = checker.check(
            df, selected_month, selected_persons, month_col, person_col, amount_col)
        if not success:
            messagebox.showerror("错误", msg, parent=self.root)
            self._log_status(f"错误：{msg}")
            return

        self._log_status(msg)
        checked = len(result_df[result_df['状态'] == '已核对'])
        pending = len(result_df[result_df['状态'] == '待核对'])
        unmatched = len(result_df[result_df['状态'] == '未匹配'])
        self._log_status(f"核对结果：已核对{checked}条，待核对{pending}条，未匹配{unmatched}条")

        self.check_result_df = result_df
        self._show_check_preview(result_df)
        self._log_status("核对完成！核对结果已显示在预览区。")
        # 自动保存核对结果
        self._log_status("自动保存核对结果...")
        self._auto_save_check_only()
        self.last_operation = 'check'
        self._log_status("=" * 50)

    # ==================== 追加到目标表 ====================

    def _append_to_target(self):
        from tkinter import messagebox, filedialog

        # 构建当前可用的数据（按工作表名称）
        all_available = {}
        if self.processor.output_data is not None:
            all_available['供应商拆分汇总'] = self.processor.output_data
        if self.check_result_df is not None and not self.check_result_df.empty:
            all_available['核对结果'] = self.check_result_df

        if not all_available:
            messagebox.showwarning("警告", "没有可追加的数据，请先进行数据汇总或核对", parent=self.root)
            return

        # 只追加尚未追加过的数据（针对当前目标文件）
        data_to_append = {
            name: df for name, df in all_available.items()
            if name not in self.appended_types
        }

        if not data_to_append:
            messagebox.showinfo("提示", "所有可用数据均已追加到目标表，无需重复追加。", parent=self.root)
            return

        target_path = self.target_file_path_var.get()
        if not target_path:
            target_path = filedialog.askopenfilename(
                title="选择要追加到的目标Excel文件",
                filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")]
            )
            if not target_path:
                self._log_status("用户取消了追加操作")
                return
            self.target_file_path = target_path
            self.target_file_path_var.set(target_path)
            self.appended_types = set()  # 新选择目标文件，清空记录
            data_to_append = all_available  # 新文件，重新计算可追加数据

        self._log_status("=" * 50)
        self._log_status(f"开始追加到目标表（{', '.join(data_to_append.keys())}）...")

        try:
            appended = []

            if os.path.exists(target_path):
                self._log_status(f"目标文件已存在：{os.path.basename(target_path)}")
                target_wb = load_workbook(target_path)
                existing_sheets = set(target_wb.sheetnames)

                def get_unique_name(base_name):
                    if base_name not in existing_sheets:
                        return base_name
                    i = 1
                    while f"{base_name}_{i}" in existing_sheets:
                        i += 1
                    return f"{base_name}_{i}"

                for sheet_name, df in data_to_append.items():
                    name = get_unique_name(sheet_name)
                    if name != sheet_name:
                        self._log_status(f"  工作表'{sheet_name}'已存在，使用新名称'{name}'")

                    df_to_write = df.copy()
                    for col in df_to_write.select_dtypes(include=['number']).columns:
                        df_to_write[col] = df_to_write[col].replace(0, None)

                    with pd.ExcelWriter(target_path, engine='openpyxl', mode='a',
                                        if_sheet_exists='new') as writer:
                        df_to_write.to_excel(writer, sheet_name=name, index=False)

                    wb = load_workbook(target_path)
                    apply_excel_styles(wb[name], df)
                    wb.save(target_path)
                    wb.close()

                    appended.append(name)

                target_wb.close()
            else:
                self._log_status(f"目标文件不存在，将创建新文件：{os.path.basename(target_path)}")
                with pd.ExcelWriter(target_path, engine='openpyxl') as writer:
                    for sheet_name, df in data_to_append.items():
                        df_to_write = df.copy()
                        for col in df_to_write.select_dtypes(include=['number']).columns:
                            df_to_write[col] = df_to_write[col].replace(0, None)
                        df_to_write.to_excel(writer, sheet_name=sheet_name, index=False)

                wb = load_workbook(target_path)
                for sheet_name, df in data_to_append.items():
                    if sheet_name in wb.sheetnames:
                        apply_excel_styles(wb[sheet_name], df)
                wb.save(target_path)
                wb.close()
                appended = list(data_to_append.keys())

            # 记录已追加的数据类型（用基础名，不是带后缀的实际表名）
            for sheet_name in data_to_append.keys():
                self.appended_types.add(sheet_name)

            self._log_status(f"追加完成！新增工作表：{', '.join(appended)}")
            self._log_status("=" * 50)
            messagebox.showinfo("成功", f"已成功追加以下工作表：\n" + "\n".join(appended), parent=self.root)

        except Exception as e:
            import traceback
            messagebox.showerror("错误", f"追加失败：{str(e)}", parent=self.root)
            self._log_status(f"错误：追加失败 {str(e)}")
            self._log_status(traceback.format_exc())

    # ==================== 自动保存方法 ====================

    def _save_df_to_excel(self, df: pd.DataFrame, save_path: str, sheet_title: str,
                           reference_df: pd.DataFrame = None) -> bool:
        """将DataFrame保存为Excel文件 - 使用pd.to_excel()快速写入 + 应用样式"""
        try:
            df_to_write = df.copy()
            for col in df_to_write.select_dtypes(include=['number']).columns:
                df_to_write[col] = df_to_write[col].replace(0, None)

            df_to_write.to_excel(save_path, engine='openpyxl',
                                 sheet_name=sheet_title, index=False)

            wb = load_workbook(save_path)
            # 使用工作表名称获取，而不是wb.active（更安全）
            ws = wb[sheet_title]
            apply_excel_styles(ws, reference_df if reference_df is not None else df)
            wb.save(save_path)
            wb.close()
            return True
        except Exception as e:
            import traceback
            self._log_status(f"   ⚠ 保存失败：{e}")
            self._log_status(traceback.format_exc())
            return False

    def _auto_save_all(self):
        """数据汇总完成后自动保存为独立文件：供应商拆分汇总.xlsx"""
        file_path = self.file_path_var.get()
        base_dir = os.path.dirname(file_path) if file_path else os.getcwd()
        save_name = '供应商拆分汇总.xlsx'
        save_path = os.path.join(base_dir, save_name)

        if self._save_df_to_excel(self.processor.output_data, save_path,
                                   '供应商拆分汇总', self.processor.output_data):
            self.last_summary_path = save_path
            self._log_status(f"   ✓ 已自动保存：{save_name}")

    def _auto_save_check_only(self):
        """核对完成后自动保存为独立文件：核对结果.xlsx"""
        if self.check_result_df is None or self.check_result_df.empty:
            return
        file_path = self.file_path_var.get() or (self.check_file_path or '')
        base_dir = os.path.dirname(file_path) or os.getcwd()
        save_name = '核对结果.xlsx'
        save_path = os.path.join(base_dir, save_name)

        if self._save_df_to_excel(self.check_result_df, save_path,
                                   '核对结果', self.check_result_df):
            self.last_check_path = save_path
            self._log_status(f"   ✓ 已自动保存：{save_name}")
