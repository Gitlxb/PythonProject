# -*- coding: utf-8 -*-
"""
多表合并（删列+公司命名）核心引擎

流程：
  1. 用户上传多个 Excel 文件，在 GUI 中逐个勾选工作表
  2. 对每个被勾选的工作表，删除「起始列及之后」的所有列（如 AT 及之后）
  3. 对每个工作表，在指定列（公司名列）的所有数据行写入 sheet_name
     —— 首表跳过保留表头，其余已 skip，全部数据行写入
  4. 合并：保留第一个工作表全部行（含表头），其余跳过 N 行表头，
     纵向堆叠；输出独立文件「汇总表.xlsx」

格式处理：
  - 纯数据合并，不保留任何单元格样式（字体/填充/边框/对齐/数字格式/合并/列宽/行高）
  - 输出独立文件，源文件零修改、零风险

性能策略（2026-08-11 回归纯数据版）：
  - 所有文件统一用 openpyxl read_only 流式读取（低内存、低开销）
  - 数据读取时通过 max_col=keep_cols 直接截断删除列，并强制统一列宽（修复列错位隐患）
  - 尾空 200 行截断：防止整列格式化导致 max_row 被撑高到 1048576 引发的内存/耗时爆炸
  - 写盘用独立文件（pd.ExcelWriter 新建），不再重写 27MB 源文件
"""

import os
import pandas as pd
from openpyxl import load_workbook

try:
    from .hbbg_process import ProcessMergeEngine
except ImportError:
    from hbbg_process import ProcessMergeEngine


class CustomMergeEngine:
    """多表合并（删列+公司命名）引擎 — 纯数据（openpyxl 流式读取 + pandas 合并），无 GUI 依赖"""

    def __init__(self, progress_callback=None, status_callback=None):
        self._progress_cb = progress_callback or (lambda p, t: None)
        self._status_cb = status_callback or (lambda t: None)

    # ================================================================
    #  静态工具方法
    # ================================================================

    @staticmethod
    def col_letter_to_index(col_letter):
        """列字母 → 1-based 列索引（A→1, J→10, AA→27）"""
        return ProcessMergeEngine.col_letter_to_index(col_letter)

    @staticmethod
    def list_sheets(file_path):
        """列出文件中的所有工作表（含隐藏标记）"""
        try:
            wb = load_workbook(file_path, read_only=True)
            sheets = [(ws.title, ws.sheet_state == 'hidden') for ws in wb.worksheets]
            wb.close()
            return sheets
        except Exception as e:
            print(f"[自定义合并] 读取工作表列表失败 {os.path.basename(file_path)}: {e}")
            return []

    # ================================================================
    #  主流程
    # ================================================================

    def run(self, selections, first_file_path, delete_col_letter,
            sheet_name_col_letter, header_rows):
        """
        执行多表合并（删列+公司命名），纯数据输出独立汇总表。

        参数:
            selections:           List[Tuple[file_path, sheet_name]]
            first_file_path:      selections[0][0]，表头来源 + 数据来源之一
            delete_col_letter:    删除起始列字母，如 "AT"
            sheet_name_col_letter: 公司名列字母，如 "AR"
            header_rows:          表头行数 N（int，>=1）
        """
        total = len(selections)
        if total == 0:
            return {"status": "error", "message": "未选择任何工作表"}

        try:
            start_col_idx = self.col_letter_to_index(delete_col_letter)
            sheet_name_col_idx = self.col_letter_to_index(sheet_name_col_letter)
        except Exception as e:
            return {"status": "error", "message": f"列字母解析失败：{e}"}

        keep_cols = start_col_idx - 1          # 保留列数（1..keep_cols）
        output_path = os.path.join(os.path.dirname(first_file_path), "汇总表.xlsx")

        # ---- 校验：公司名列必须落在保留范围内 ----
        if sheet_name_col_idx > keep_cols:
            return {
                "status": "error",
                "message": (f"公司名列 {sheet_name_col_letter}（第 {sheet_name_col_idx} 列）"
                            f"在删除列 {delete_col_letter} 及之后，请将删除起始列改小"
                            f"或选择更靠前的列")
            }

        # ============================================================
        #  数据读取（所有文件统一 read_only 流式）
        # ============================================================
        wb_cache = {}
        all_frames = []

        for idx, (file_path, sheet_name) in enumerate(selections):
            filename = os.path.basename(file_path)
            base_pct = (idx / total) * 85
            self._progress_cb(base_pct,
                              f"[{idx+1}/{total}] 正在读取: {filename} / {sheet_name}")
            self._status_cb(f"读取 [{idx+1}/{total}]: {sheet_name}")

            is_first = (idx == 0)

            # ---- 获取工作表对象（按文件缓存，每文件只读一次）----
            if file_path in wb_cache:
                wb = wb_cache[file_path]
                try:
                    ws = wb[sheet_name]
                except Exception:
                    print(f"[自定义合并] 工作表不存在: {sheet_name}")
                    continue
            else:
                try:
                    wb = load_workbook(file_path, read_only=True, data_only=False)
                    wb_cache[file_path] = wb
                    ws = wb[sheet_name]
                except Exception as e:
                    print(f"[自定义合并] 打开失败 {file_path}/{sheet_name}: {e}")
                    continue

            # ---- 读取数据：max_col=keep_cols 直接截断删除列 + 强制列宽 ----
            max_row = ws.max_row
            use_early_stop = (max_row >= 1_000_000)  # 表格被格式化撑高时启用尾空截断
            rows = []
            empty_run = 0
            for row in ws.iter_rows(min_row=1, max_row=max_row,
                                    max_col=keep_cols, values_only=True):
                all_none = all(v is None for v in row)
                if all_none:
                    if use_early_stop:
                        empty_run += 1
                        if empty_run >= 200:
                            break
                        continue
                    else:
                        rows.append(list(row))
                        continue
                empty_run = 0
                rows.append(list(row))

            if not rows:
                print(f"[自定义合并] 跳过空表: {sheet_name}")
                continue

            df = pd.DataFrame(rows)

            # ---- 非首表：跳过表头行 ----
            if not is_first:
                if header_rows < len(df):
                    df = df.iloc[header_rows:]
                else:
                    print(f"[自定义合并] 跳过（表头行数超出数据）: {sheet_name}")
                    continue

            # ---- 公司命名：每个数据行整列写入 sheet_name ----
            #   首表：df 含表头，从 header_rows 行起写（已加 if 守卫）
            #   其他表：df 已 skip 表头，从第 0 行起写
            start_idx = header_rows if is_first else 0
            col_pos = sheet_name_col_idx - 1  # pandas 0-based 列索引
            if start_idx < len(df) and col_pos < df.shape[1]:
                df.iloc[start_idx:, col_pos] = sheet_name

            all_frames.append(df)

            self._progress_cb(((idx + 1) / total) * 85,
                              f"[{idx+1}/{total}] ✓ {sheet_name} ({len(df)} 行)")

        # ---- 释放所有工作簿（节省内存）----
        for wb in wb_cache.values():
            try:
                wb.close()
            except Exception:
                pass

        if not all_frames:
            return {"status": "error", "message": "所有工作表处理后无数据"}

        # ---- 合并 ----
        self._progress_cb(88, "正在合并数据...")
        result = pd.concat(all_frames, ignore_index=True)

        # ---- 写盘：独立文件（纯数据，无样式，极快）----
        self._progress_cb(92, "正在保存汇总表...")
        try:
            with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
                result.to_excel(writer, sheet_name='汇总表', index=False, header=False)
        except Exception as e:
            return {"status": "error", "message": f"保存汇总表失败：{e}"}

        self._progress_cb(100, "完成！")

        return {
            "status": "success",
            "output_path": output_path,
            "summary": (f"已完成多表合并（删列+公司命名）。\n\n"
                        f"共处理 {total} 个工作表。\n"
                        f"删除起始列：{delete_col_letter} 及之后\n"
                        f"公司命名列：{sheet_name_col_letter}（每张工作表数据行写入 sheet_name）\n"
                        f"表头行数：{header_rows}\n\n"
                        f"汇总表已生成为独立文件（纯数据，不保留原格式；源文件未修改）：\n{output_path}")
        }
