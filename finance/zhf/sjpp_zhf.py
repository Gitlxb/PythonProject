import os
import traceback
from copy import copy
from typing import Optional, Tuple

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.workbook import Workbook

import tkinter as tk
from tkinter import filedialog, messagebox, ttk


def normalize_text(value) -> str:
    """
    规范化文本输入，去除空白字符并处理特殊类型

    Args:
        value: 待处理的输入值，可以是 None、字符串、数字或其他类型

    Returns:
        str: 规范化后的字符串。None 返回空字符串；字符串去除首尾空白和换行符；
             整数形式的浮点数转为整数字符串；其他类型转为字符串并去除首尾空白
    """
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip().replace("\n", "").replace("\r", "")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def find_last_used_row(ws, check_max_col: int = 60) -> int:
    """
    查找工作表中最后一个包含数据的行号

    Args:
        ws: 工作表对象
        check_max_col (int): 检查的最大列数，默认为60

    Returns:
        int: 最后一个包含数据的行号，如果工作表为空则返回1
    """
    max_check_col = min(check_max_col, ws.max_column)
    for r in range(ws.max_row, 0, -1):
        for c in range(1, max_check_col + 1):
            if ws.cell(r, c).value not in (None, ""):
                return r
    return 1


def copy_cell_style(src, dst):
    """
    将源单元格的样式属性复制到目标单元格

    Args:
        src: 源单元格对象，包含字体、填充、边框、对齐等样式属性
        dst: 目标单元格对象，将被应用源单元格的样式

    Returns:
        None
    """
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format
    dst.protection = copy(src.protection)


def apply_column_style_from_reference(ws, new_col: int, header_ref_col: int, data_ref_col: int, max_row: int):
    """
    从参考列复制单元格样式到新列

    Args:
        ws: 工作表对象
        new_col (int): 目标列号，将应用复制后的样式
        header_ref_col (int): 表头参考列号，用于复制表头样式
        data_ref_col (int): 数据参考列号，用于复制数据行样式
        max_row (int): 工作表的最大行数，用于判断是否需要复制数据行样式
    """
    copy_cell_style(ws.cell(1, header_ref_col), ws.cell(1, new_col))
    if max_row >= 2:
        copy_cell_style(ws.cell(2, data_ref_col), ws.cell(2, new_col))


def ensure_column_width(ws, col_letter: str, width: float):
    """
    设置工作表中指定列的宽度

    Args:
        ws: 工作表对象
        col_letter (str): 列字母标识，如 'A'、'B' 等
        width (float): 列宽度值
    """
    ws.column_dimensions[col_letter].width = width


def safe_sheet(wb: Workbook, name: str):
    """
    安全获取工作表，如果工作表不存在则抛出异常

    Args:
        wb (Workbook): Excel工作簿对象
        name (str): 工作表名称

    Returns:
        Worksheet: 指定名称的工作表对象

    Raises:
        ValueError: 当工作表不存在时抛出异常
    """
    if name not in wb.sheetnames:
        raise ValueError(f"工作表不存在：{name}")
    return wb[name]


def find_sheet_by_prefix(wb: Workbook, prefix: str) -> str:
    """
    根据前缀查找工作簿中的工作表名称

    Args:
        wb (Workbook): Excel 工作簿对象
        prefix (str): 工作表名称的前缀

    Returns:
        str: 匹配前缀的工作表名称

    Raises:
        ValueError: 当未找到以指定前缀开头的工作表时抛出异常
    """
    for name in wb.sheetnames:
        if name.startswith(prefix):
            return name
    raise ValueError(f"未找到以“{prefix}”开头的工作表。当前工作表：{wb.sheetnames}")


def quote_sheet_name(name: str) -> str:
    """
    对工作表名称进行引号转义处理

    Args:
        name (str): 原始工作表名称

    Returns:
        str: 转义后的工作表名称，将单引号替换为两个单引号并用单引号包裹
    """
    safe_name = name.replace("'", "''")
    return f"'{safe_name}'"


def build_external_sheet_ref(file_path: str, sheet_name: str) -> str:
    """
    构建Excel外部工作表引用字符串

    Args:
        file_path (str): Excel文件的路径
        sheet_name (str): 工作表名称

    Returns:
        str: Excel外部引用格式的工作表引用字符串，格式为：'路径\[文件名]工作表名'
    """
    abs_path = os.path.abspath(file_path)
    folder = os.path.dirname(abs_path).replace("/", "\\")
    file_name = os.path.basename(abs_path)
    ext_ref = f"{folder}\\[{file_name}]{sheet_name}"
    ext_ref = ext_ref.replace("'", "''")
    return f"'{ext_ref}'"


def set_formula_column(ws, start_row: int, end_row: int, col_letter: str, formula_builder):
    """
    为指定列的单元格范围设置公式

    Args:
        ws: 工作表对象
        start_row (int): 起始行号
        end_row (int): 结束行号
        col_letter (str): 列字母（如 'A'、'B' 等）
        formula_builder: 公式构建函数，接收行号作为参数，返回公式字符串
    """
    for r in range(start_row, end_row + 1):
        ws[f"{col_letter}{r}"] = formula_builder(r)


def detect_sheet_names(source_wb: Workbook, prev_wb: Optional[Workbook] = None) -> Tuple[
    str, str, str, str, Optional[str]]:
    """
    检测并获取工作簿中的目标工作表名称

    Args:
        source_wb (Workbook): 源工作簿对象，用于查找目标工作表
        prev_wb (Optional[Workbook]): 可选的上期工作簿对象，用于查找上期百事通报备名单，默认为 None

    Returns:
        Tuple[str, str, str, str, Optional[str]]: 包含五个元素的元组，依次为：
            - 项目经理工作表名称（固定为"运营中心-项目经理"）
            - 代工费工作表名称
            - 百事通报备名单工作表名称
            - 员工信息汇总工作表名称
            - 上期百事通报备名单工作表名称（如果 prev_wb 为 None 则返回 None）
    """
    sheet_pm = "运营中心-项目经理"
    sheet_daigong = find_sheet_by_prefix(source_wb, "代工费-")
    sheet_baishi = find_sheet_by_prefix(source_wb, "百事通报备名单-")
    sheet_emp = find_sheet_by_prefix(source_wb, "员工信息汇总-")
    prev_baishi = find_sheet_by_prefix(prev_wb, "百事通报备名单-") if prev_wb else None
    return sheet_pm, sheet_daigong, sheet_baishi, sheet_emp, prev_baishi


def process_daigong_sheet(ws_dg, baishi_sheet_name: str):
    """
    处理代工工作表，设置项目名称+姓名列和报备状态验证列

    Args:
        ws_dg: 代工工作表对象
        baishi_sheet_name (str): 报备工作表的名称，用于VLOOKUP公式引用

    功能说明:
        - 设置K列为"项目名称+姓名"，公式为D列和E列的拼接
        - 设置L列为"是否在报备名单"，通过VLOOKUP函数从报备表中验证
        - 应用统一的列样式和列宽设置
    """
    last_row = find_last_used_row(ws_dg, check_max_col=20)

    for col_letter in ["K", "L"]:
        col_idx = column_index_from_string(col_letter)
        apply_column_style_from_reference(ws_dg, col_idx, 10, 10, last_row)

    ws_dg["K2"] = "项目名称+姓名"
    ws_dg["L2"] = "是否在报备名单"

    ensure_column_width(ws_dg, "K", 24)
    ensure_column_width(ws_dg, "L", 24)

    quoted_baishi = quote_sheet_name(baishi_sheet_name)

    set_formula_column(ws_dg, 3, last_row, "K", lambda r: f"=D{r}&E{r}")
    set_formula_column(ws_dg, 3, last_row, "L", lambda r: f"=VLOOKUP(K{r},{quoted_baishi}!$AB:$AB,1,0)")


def process_baishi_sheet(ws_bs, daigong_sheet_name: str, pm_sheet_name: str, emp_sheet_name: str,
                         prev_month_path: Optional[str], prev_baishi_sheet_name: Optional[str]):
    """
    处理百事工作表，设置列样式、标题、宽度和公式

    Args:
        ws_bs: 百事工作表对象
        daigong_sheet_name (str): 代工工作表名称
        pm_sheet_name (str): 项目经理工作表名称
        emp_sheet_name (str): 员工工作表名称
        prev_month_path (Optional[str]): 上月文件路径，用于引用上月数据
        prev_baishi_sheet_name (Optional[str]): 上月百世工作表名称
    """
    last_row = find_last_used_row(ws_bs, check_max_col=35)

    for col_letter in ["Z", "AA", "AB", "AC", "AD", "AE", "AG"]:
        col_idx = column_index_from_string(col_letter)
        apply_column_style_from_reference(ws_bs, col_idx, 23, 23, last_row)

    ws_bs["Z1"] = "姓名"
    ws_bs["AA1"] = "企业+姓名"
    ws_bs["AB1"] = "工厂实际结算"
    ws_bs["AC1"] = "总人数（上月延续）"
    ws_bs["AD1"] = "项目驻场"
    ws_bs["AE1"] = "项目经理"
    ws_bs["AG1"] = "员工信息匹配"

    ensure_column_width(ws_bs, "Z", 12)
    ensure_column_width(ws_bs, "AA", 24)
    ensure_column_width(ws_bs, "AB", 24)
    ensure_column_width(ws_bs, "AC", 24)
    ensure_column_width(ws_bs, "AD", 12)
    ensure_column_width(ws_bs, "AE", 12)
    ensure_column_width(ws_bs, "AG", 24)

    q_daigong = quote_sheet_name(daigong_sheet_name)
    q_pm = quote_sheet_name(pm_sheet_name)
    q_emp = quote_sheet_name(emp_sheet_name)

    set_formula_column(ws_bs, 2, last_row, "Z", lambda r: f"=C{r}")
    set_formula_column(ws_bs, 2, last_row, "AA", lambda r: f"=K{r}&Z{r}")
    set_formula_column(ws_bs, 2, last_row, "AB", lambda r: f"=VLOOKUP(AA{r},{q_daigong}!$K:$K,1,0)")

    if prev_month_path and prev_baishi_sheet_name:
        ext_prev_baishi = build_external_sheet_ref(prev_month_path, prev_baishi_sheet_name)
        set_formula_column(
            ws_bs, 2, last_row, "AC",
            lambda r: f"=VLOOKUP(AA{r},{ext_prev_baishi}!$AC:$AC,1,0)"
        )
    else:
        set_formula_column(ws_bs, 2, last_row, "AC", lambda r: "")

    set_formula_column(ws_bs, 2, last_row, "AD", lambda r: f"=R{r}")
    set_formula_column(ws_bs, 2, last_row, "AE", lambda r: f"=VLOOKUP(K{r},{q_pm}!$B:$C,2,0)")
    set_formula_column(ws_bs, 2, last_row, "AG", lambda r: f"=VLOOKUP(AB{r},{q_emp}!$AB:$AB,1,0)")


def process_employee_sheet(ws_emp, baishi_sheet_name: str):
    """
    处理员工工作表，设置指定列的样式、列宽和公式

    Args:
        ws_emp: 员工工作表对象
        baishi_sheet_name (str): 白石工作表的名称，用于VLOOKUP公式引用

    功能说明:
        - 为AA到AF列设置列样式和列宽
        - 设置表头名称（姓名、项目+姓名、工厂实际结算、项目驻场、项目经理、总人数）
        - 为每行设置公式，从白石工作表中查找对应数据
    """
    last_row = find_last_used_row(ws_emp, check_max_col=40)

    for col_letter in ["AA", "AB", "AC", "AD", "AE", "AF"]:
        col_idx = column_index_from_string(col_letter)
        apply_column_style_from_reference(ws_emp, col_idx, 26, 26, last_row)

    ws_emp["AA1"] = "姓名"
    ws_emp["AB1"] = "项目+姓名"
    ws_emp["AC1"] = "工厂实际结算"
    ws_emp["AD1"] = "项目驻场"
    ws_emp["AE1"] = "项目经理"
    ws_emp["AF1"] = "总人数（上月延续）"

    ensure_column_width(ws_emp, "AA", 13)
    ensure_column_width(ws_emp, "AB", 24)
    ensure_column_width(ws_emp, "AC", 24)
    ensure_column_width(ws_emp, "AD", 13)
    ensure_column_width(ws_emp, "AE", 13)
    ensure_column_width(ws_emp, "AF", 18)


    q_baishi = quote_sheet_name(baishi_sheet_name)

    set_formula_column(ws_emp, 2, last_row, "AA", lambda r: f"=C{r}")
    set_formula_column(ws_emp, 2, last_row, "AB", lambda r: f"=F{r}&AA{r}")
    set_formula_column(ws_emp, 2, last_row, "AC", lambda r: f"=VLOOKUP(AB{r},{q_baishi}!$AB:$AE,1,0)")
    set_formula_column(ws_emp, 2, last_row, "AD", lambda r: f"=VLOOKUP(AB{r},{q_baishi}!$AC:$AE,2,0)")
    set_formula_column(ws_emp, 2, last_row, "AE", lambda r: f"=VLOOKUP(AB{r},{q_baishi}!$AC:$AE,3,0)")
    set_formula_column(ws_emp, 2, last_row, "AF", lambda r: f"=VLOOKUP(AB{r},{q_baishi}!$AC:$AE,1,0)")


def generate_full_workbook(source_path: str, prev_month_path: Optional[str] = None) -> str:
    """
    生成完整的工作簿文件，处理员工、白世和代工相关数据并保存

    Args:
        source_path (str): 源Excel文件路径
        prev_month_path (Optional[str]): 上月Excel文件路径，用于数据对比，默认为None

    Returns:
        str: 生成的输出文件路径，文件名格式为"{原文件名}_自动生成（已核对）.xlsx"
    """
    wb = load_workbook(source_path)
    prev_wb = load_workbook(prev_month_path, data_only=False) if prev_month_path else None

    sheet_pm, sheet_daigong, sheet_baishi, sheet_emp, prev_baishi = detect_sheet_names(wb, prev_wb)

    ws_pm = safe_sheet(wb, sheet_pm)
    ws_dg = safe_sheet(wb, sheet_daigong)
    ws_bs = safe_sheet(wb, sheet_baishi)
    ws_emp = safe_sheet(wb, sheet_emp)

    _ = ws_pm

    process_employee_sheet(ws_emp, sheet_baishi)
    process_baishi_sheet(
        ws_bs,
        daigong_sheet_name=sheet_daigong,
        pm_sheet_name=sheet_pm,
        emp_sheet_name=sheet_emp,
        prev_month_path=prev_month_path,
        prev_baishi_sheet_name=prev_baishi,
    )
    process_daigong_sheet(ws_dg, baishi_sheet_name=sheet_baishi)
    
    # 设置自动重算
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcOnSave = True
    
    # 保存文件
    base_dir = os.path.dirname(source_path)
    base_name = os.path.splitext(os.path.basename(source_path))[0]
    out_path = os.path.join(base_dir, f"{base_name}_自动生成（已核对）.xlsx")
    wb.save(out_path)
    return out_path


class App:
    def __init__(self, root):
        self.root = root
        self.root.title("完整版表格自动生成工具")
        self.root.geometry("820x530")

        self.source_var = tk.StringVar()
        self.prev_var = tk.StringVar()
        self.output_var = tk.StringVar(value="等待生成...")

        self.build_ui()

    def build_ui(self):
        main = ttk.Frame(self.root, padding=16)
        main.pack(fill="both", expand=True)

        ttk.Label(
            main,
            text="从原始表自动生成“完整版”表格",
            font=("Microsoft YaHei UI", 14, "bold"),
        ).pack(anchor="w", pady=(0, 14))

        tip_text = (
            "使用说明：\n"
            "1）先选择“本月原始表”。\n"
            "2）再选择“上月（已核对）表”。\n"
            "3）程序会按公式链生成完整版，其中“总人数（上月延续）”来自上月已核对表。\n\n"
        )
        ttk.Label(main, text=tip_text, justify="left").pack(anchor="w", pady=(0, 12))

        row1 = ttk.Frame(main)
        row1.pack(fill="x", pady=6)
        ttk.Label(row1, text="本月原始表：", width=12).pack(side="left")
        ttk.Entry(row1, textvariable=self.source_var).pack(side="left", fill="x", expand=True, padx=6)
        ttk.Button(row1, text="选择文件", command=self.choose_source).pack(side="left")

        row2 = ttk.Frame(main)
        row2.pack(fill="x", pady=6)
        ttk.Label(row2, text="上月（已核对）：", width=12).pack(side="left")
        ttk.Entry(row2, textvariable=self.prev_var).pack(side="left", fill="x", expand=True, padx=6)
        ttk.Button(row2, text="选择文件", command=self.choose_prev).pack(side="left")

        row3 = ttk.Frame(main)
        row3.pack(fill="x", pady=18)
        ttk.Button(row3, text="开始生成", command=self.run_generate).pack(side="left")
        ttk.Button(row3, text="打开输出目录", command=self.open_output_dir).pack(side="left", padx=8)

        ttk.Separator(main).pack(fill="x", pady=10)
        ttk.Label(main, text="输出结果：", font=("Microsoft YaHei UI", 10, "bold")).pack(anchor="w")
        ttk.Label(main, textvariable=self.output_var, foreground="#0a6").pack(anchor="w", pady=6)

        note = (
            "说明：\n"
            "1）程序会自动识别以“代工费- / 百事通报备名单- / 员工信息汇总-”开头的工作表。\n"
            "2）生成后请打开 Excel 表格，让外部引用公式完成重算。\n"
            "3）如果上月文件路径发生变化，Excel 首次打开时可能提示更新链接，属于正常现象。"
        )
        ttk.Label(main, text=note, justify="left", foreground="#555").pack(anchor="w", pady=(12, 0))

    def choose_source(self):
        path = filedialog.askopenfilename(
            title="选择本月原始表",
            filetypes=[("Excel 文件", "*.xlsx")],
        )
        if path:
            self.source_var.set(path)

    def choose_prev(self):
        path = filedialog.askopenfilename(
            title="选择上月（已核对）表",
            filetypes=[("Excel 文件", "*.xlsx")],
        )
        if path:
            self.prev_var.set(path)

    def run_generate(self):
        source_path = self.source_var.get().strip()
        prev_path = self.prev_var.get().strip()

        if not source_path:
            messagebox.showwarning("提示", "请先选择本月原始表。")
            return
        if not prev_path:
            messagebox.showwarning("提示", "请再选择上月（已核对）表；“总人数（上月延续）”必须依赖这个文件。")
            return
        if not os.path.exists(source_path):
            messagebox.showerror("错误", "本月原始表文件不存在。")
            return
        if not os.path.exists(prev_path):
            messagebox.showerror("错误", "你选择的上月（已核对）文件不存在。")
            return

        try:
            out_path = generate_full_workbook(source_path, prev_path)
            self.output_var.set(out_path)
            messagebox.showinfo("完成", f"已生成文件：\n{out_path}")
        except Exception as e:
            self.output_var.set("生成失败，请查看报错信息。")
            messagebox.showerror("生成失败", f"{e}\n\n详细报错：\n{traceback.format_exc()}")

    def open_output_dir(self):
        out_path = self.output_var.get().strip()
        if not out_path or out_path == "等待生成...":
            messagebox.showinfo("提示", "还没有生成文件。")
            return
        folder = os.path.dirname(out_path)
        if os.path.exists(folder):
            try:
                os.startfile(folder)
            except AttributeError:
                import subprocess
                subprocess.Popen(["xdg-open", folder])


def main():
    root = tk.Tk()
    try:
        style = ttk.Style()
        if "vista" in style.theme_names():
            style.theme_use("vista")
    except Exception:
        pass

    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
