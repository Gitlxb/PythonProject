# -*- coding: utf-8 -*-
"""
通用合并核心引擎

将多个 Excel 文件原样堆叠合并，保留全部格式。
支持：文件间空行、跳过表头、保留表头文件、列移动设置。

用法:
    from hbbg_generic import GenericMergeEngine

    engine = GenericMergeEngine(
        gap_rows=0,
        progress_callback=lambda pct, text: ...,
        status_callback=lambda text: ...,
    )
    engine.do_merge(ordered_files, output_path,
                    keep_header_file=None, skip_rows=1,
                    col_move_rules=None)
"""
import os
import traceback

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from copy import copy

# 列移动设置（来自同目录）
try:
    from .hbbg_column_move import apply_column_moves, get_rules_for_file
except ImportError:
    from hbbg_column_move import apply_column_moves, get_rules_for_file


class GenericMergeEngine:
    """通用合并引擎：原样堆叠合并，无 GUI 依赖"""

    def __init__(self, gap_rows=0, progress_callback=None, status_callback=None):
        """
        参数:
            gap_rows: 文件间空行数（默认0）
            progress_callback: Callable(pct: float, text: str)
            status_callback:   Callable(text: str)
        """
        self.gap_rows = gap_rows
        self._progress_cb = progress_callback or (lambda p, t: None)
        self._status_cb = status_callback or (lambda t: None)

    # ================================================================
    #  静态工具方法
    # ================================================================

    @staticmethod
    def get_actual_max_row(ws):
        """检测实际有数据的最后一行，避免 max_row 因格式虚高"""
        max_row = ws.max_row
        max_col = ws.max_column
        for row_idx in range(max_row, 0, -1):
            for col_idx in range(1, max_col + 1):
                if ws.cell(row=row_idx, column=col_idx).value is not None:
                    return row_idx
        return 0

    # ================================================================
    #  核心复制引擎
    # ================================================================

    def copy_sheet_content(self, file_path, target_ws, start_row,
                           skip_header_rows=0, col_move_rules=None):
        """
        将 Excel 文件原样复制到目标 worksheet，保留所有格式。

        参数:
            file_path:       源文件完整路径
            target_ws:       目标 openpyxl worksheet
            start_row:       写入起始行（1-based）
            skip_header_rows: 跳过源文件前 N 行
            col_move_rules:   List[ColumnMoveRule] — 列移动规则

        返回: 实际写入的行数
        """
        # 当有列移动规则时，使用 data_only=True 获取公式缓存值，
        # 避免 insert_cols() 后公式引用错乱
        use_data_only = bool(col_move_rules)
        try:
            source_wb = load_workbook(file_path, data_only=use_data_only)
        except Exception:
            df = pd.read_excel(file_path, header=None, engine='xlrd')
            temp_path = file_path + '_temp_convert.xlsx'
            df.to_excel(temp_path, index=False, header=False, engine='openpyxl')
            source_wb = load_workbook(temp_path, data_only=use_data_only)
            os.remove(temp_path)

        source_ws = source_wb.active

        # ---- 列移动（在复制前执行）----
        if col_move_rules:
            basename = os.path.basename(file_path)
            file_rules = get_rules_for_file(col_move_rules, basename)
            if file_rules:
                inserted = apply_column_moves(source_ws, file_rules)
                print(f"    [{basename}] 列移动: {inserted} 列被插入")

        # 行方向：精确检测实际数据范围
        raw_max_row = source_ws.max_row
        actual_max_row = self.get_actual_max_row(source_ws)
        max_col = source_ws.max_column

        if actual_max_row == 0 or max_col == 0:
            source_wb.close()
            return 0

        if skip_header_rows >= actual_max_row:
            source_wb.close()
            return 0

        print(f"  {os.path.basename(file_path)}: 实际数据 = {actual_max_row}行 x {max_col}列"
              f"（原始max_row={raw_max_row}{' ← 已裁减空行' if actual_max_row < raw_max_row else ''}，"
              f"跳过{skip_header_rows}行）")

        row_start = 1 + skip_header_rows
        written_rows = actual_max_row - skip_header_rows

        # ---- 复制行高 + 单元格（值 & 样式）----
        for row_idx in range(row_start, actual_max_row + 1):
            src_rd = source_ws.row_dimensions[row_idx]
            tgt_row_idx = start_row + row_idx - row_start
            if src_rd.height:
                target_ws.row_dimensions[tgt_row_idx].height = src_rd.height

            for col_idx in range(1, max_col + 1):
                src_cell = source_ws.cell(row=row_idx, column=col_idx)
                tgt_cell = target_ws.cell(row=tgt_row_idx, column=col_idx)
                tgt_cell.value = src_cell.value
                if src_cell.has_style:
                    tgt_cell.font = copy(src_cell.font)
                    tgt_cell.border = copy(src_cell.border)
                    tgt_cell.fill = copy(src_cell.fill)
                    tgt_cell.number_format = src_cell.number_format
                    tgt_cell.protection = copy(src_cell.protection)
                    tgt_cell.alignment = copy(src_cell.alignment)

        # ---- 复制合并单元格 ----
        for merged_range in source_ws.merged_cells.ranges:
            min_col, min_row, max_col_m, max_row_m = merged_range.bounds
            if min_row > actual_max_row:
                continue
            if max_row_m <= skip_header_rows:
                continue
            adj_min_row = max(min_row, row_start)
            adj_max_row = min(max_row_m, actual_max_row)
            new_min_row = adj_min_row + start_row - row_start
            new_max_row = adj_max_row + start_row - row_start
            new_ref = f"{get_column_letter(min_col)}{new_min_row}:{get_column_letter(max_col_m)}{new_max_row}"
            try:
                target_ws.merge_cells(new_ref)
            except Exception as e:
                print(f"    合并单元格跳过 {new_ref}: {e}")

        # ---- 复制列宽 ----
        for col_idx in range(1, max_col + 1):
            col_letter = get_column_letter(col_idx)
            if col_letter in source_ws.column_dimensions:
                src_cd = source_ws.column_dimensions[col_letter]
                if src_cd.width:
                    tgt_cd = target_ws.column_dimensions[col_letter]
                    if not tgt_cd.width or src_cd.width > tgt_cd.width:
                        tgt_cd.width = src_cd.width

        source_wb.close()
        return written_rows

    # ================================================================
    #  合并执行
    # ================================================================

    def do_merge(self, ordered_files, output_path,
                 keep_header_file=None, skip_rows=1,
                 col_move_rules=None):
        """
        执行通用合并。

        参数:
            ordered_files:     有序文件路径列表
            output_path:       输出文件路径
            keep_header_file:  保留表头的文件路径（None=所有文件都保留表头）
            skip_rows:         非表头文件跳过的行数
            col_move_rules:    List[ColumnMoveRule] — 列移动规则
        """
        total = len(ordered_files)
        print(f"\n[通用合并] 共 {total} 个文件，间隔 {self.gap_rows} 行，"
              f"保留表头: {'无' if not keep_header_file else os.path.basename(keep_header_file)}")

        new_wb = Workbook()
        new_ws = new_wb.active
        new_ws.title = "合并结果"
        current_row = 1

        for idx, file_path in enumerate(ordered_files):
            filename = os.path.basename(file_path)
            self._progress_cb(((idx + 1) / total) * 100, f"[{idx+1}/{total}] 正在合并: {filename}")
            self._status_cb(f"合并 [{idx+1}/{total}]: {filename}")

            # 判断当前文件是否需要跳过表头
            if keep_header_file is None:
                current_skip = 0
            elif file_path == keep_header_file:
                current_skip = 0
            else:
                current_skip = skip_rows

            try:
                written = self.copy_sheet_content(
                    file_path, new_ws, current_row,
                    skip_header_rows=current_skip,
                    col_move_rules=col_move_rules,
                )
                if written > 0:
                    current_row += written + self.gap_rows
                    print(f"  ✓ 已合并 [{idx+1}/{total}]: {filename}（{written}行）")
                else:
                    print(f"  ✗ 跳过（无内容）: {filename}")
            except Exception as e:
                print(f"  ✗ 合并失败: {filename} - {e}")
                traceback.print_exc()

        new_wb.save(output_path)
        new_wb.close()

        self._progress_cb(100, f"合并完成！共 {total} 个文件")
