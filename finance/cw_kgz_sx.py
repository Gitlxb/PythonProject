# -- coding: utf-8 --
# @Time : 2025-07-28 16:54
# @Author : 贝特利
# @Email : 1356087739@qq.com
# @File : cw_kgz_sx.py
# @Software: PyCharm

import os
import pandas as pd
from datetime import datetime, timedelta
from tkinter import Tk, filedialog, simpledialog, messagebox
from tkinter.ttk import Combobox
import tkinter as tk
from openpyxl import load_workbook
from openpyxl.styles import numbers


def select_excel_file(parent=None):
    """选择Excel文件"""
    if parent and (isinstance(parent, tk.Tk) or isinstance(parent, tk.Toplevel)):
        temp = tk.Toplevel(parent)
        temp.withdraw()
        file_path = filedialog.askopenfilename(
            title="选择Excel文件",
            filetypes=[("Excel文件", "*.xlsx *.xls"), ("所有文件", "*.*")],
            parent=temp
        )
        temp.destroy()
    else:
        root = Tk()
        root.withdraw()
        file_path = filedialog.askopenfilename(
            title="选择Excel文件",
            filetypes=[("Excel文件", "*.xlsx *.xls"), ("所有文件", "*.*")]
        )
        root.destroy()
    return file_path


def get_system_date():
    """获取当前系统日期"""
    return datetime.now().date()


def select_date_column(columns, parent=None):
    """选择日期列"""
    if parent and (isinstance(parent, tk.Tk) or isinstance(parent, tk.Toplevel)):
        root = tk.Toplevel(parent)
    else:
        root = tk.Tk()
    root.title("选择日期列")
    root.geometry("300x150")
    root.transient(parent) if parent else None
    root.grab_set()

    tk.Label(root, text="请选择日期列：").pack(pady=10)

    combo = Combobox(root, values=columns, state="readonly")
    combo.pack(pady=10)
    combo.current(0)

    selected_column = [None]

    def on_ok():
        selected_column[0] = combo.get()
        root.destroy()

    tk.Button(root, text="确定", command=on_ok).pack(pady=10)

    root.lift()
    root.focus_force()
    root.mainloop()
    return selected_column[0]


def calculate_date_range(user_date):
    """根据规则计算日期范围"""
    if user_date.day <= 20:  # 20号及之前
        # 本月、上个月、上上个月的1号
        start_date = (datetime(user_date.year, user_date.month, 1) - pd.DateOffset(months=2)).date()
    else:  # 20号之后
        # 本月和上个月的1号
        start_date = (datetime(user_date.year, user_date.month, 1) - pd.DateOffset(months=1)).date()

    end_date = user_date
    return start_date.replace(day=1), end_date


def filter_data(df, date_column, user_date):
    """筛选数据并返回符合和不符合的数据"""
    try:
        # 计算日期范围
        start_date, end_date = calculate_date_range(user_date)

        # 保存原始数据类型
        original_dtypes = df.dtypes.to_dict()

        # 转换日期列为datetime.date类型
        df[date_column] = pd.to_datetime(df[date_column]).dt.date

        # 确保比较的类型一致
        start_date = pd.to_datetime(start_date).date()
        end_date = pd.to_datetime(end_date).date()

        # 筛选符合条件的数据
        matched_df = df[(df[date_column] >= start_date) & (df[date_column] <= end_date)]

        # 筛选不符合条件的数据
        unmatched_df = df[~((df[date_column] >= start_date) & (df[date_column] <= end_date))]

        # 恢复原始数据类型
        for df_to_process in [matched_df, unmatched_df]:
            for col, dtype in original_dtypes.items():
                if col in df_to_process.columns:
                    if dtype == 'object':
                        df_to_process[col] = df_to_process[col].astype(str)
                    else:
                        df_to_process[col] = df_to_process[col].astype(dtype)

        return matched_df, unmatched_df
    except Exception as e:
        messagebox.showerror("错误", f"数据处理出错: {str(e)}")
        return None, None


def main(parent=None):
    try:
        # 1. 选择文件
        excel_path = select_excel_file(parent)
        if not excel_path:
            return

        # 2. 读取数据，将所有列作为文本读取以保留原始格式
        df = pd.read_excel(excel_path, dtype=str)

        # 3. 获取系统日期
        user_date = get_system_date()
        messagebox.showinfo("使用日期", f"将使用当前系统日期: {user_date.strftime('%Y-%m-%d')}", parent=parent)

        # 4. 选择日期列
        date_column = select_date_column(df.columns.tolist(), parent)
        if not date_column:
            return

        # 5. 筛选数据
        matched_df, unmatched_df = filter_data(df, date_column, user_date)
        if matched_df is None or unmatched_df is None:
            return

        # 6. 保存结果
        dir_name = os.path.dirname(excel_path)
        base_name = os.path.splitext(os.path.basename(excel_path))[0]
        new_filename = f"{base_name}_filtered_{user_date.strftime('%Y%m%d')}.xlsx"
        new_path = os.path.join(dir_name, new_filename)

        # 使用ExcelWriter确保数据类型正确
        with pd.ExcelWriter(new_path, engine='openpyxl') as writer:
            # 写入符合日期的数据
            matched_df.to_excel(writer, sheet_name='符合日期', index=False)

            # 写入不符合日期的数据
            unmatched_df.to_excel(writer, sheet_name='不符合日期', index=False)

            # 获取工作簿和工作表对象
            workbook = writer.book
            matched_ws = writer.sheets['符合日期']
            unmatched_ws = writer.sheets['不符合日期']

            # 设置所有列为文本格式，防止科学计数法
            for ws in [matched_ws, unmatched_ws]:
                for column in ws.columns:
                    for cell in column:
                        cell.number_format = numbers.FORMAT_TEXT

            # 获取原始工作表的列宽
            wb = load_workbook(excel_path)
            original_ws = wb.active

            # 复制列宽到两个工作表
            for ws in [matched_ws, unmatched_ws]:
                for col in range(1, original_ws.max_column + 1):
                    col_letter = original_ws.cell(row=1, column=col).column_letter
                    if col_letter in original_ws.column_dimensions:
                        ws.column_dimensions[col_letter].width = original_ws.column_dimensions[col_letter].width

        messagebox.showinfo("完成", f"文件已保存至:\n{new_path}\n"
                                  f"符合日期数据: {len(matched_df)}条\n"
                                  f"不符合日期数据: {len(unmatched_df)}条", parent=parent)

    except Exception as e:
        messagebox.showerror("错误", f"程序运行出错: {str(e)}", parent=parent)


# if __name__ == "__main__":
#     main()