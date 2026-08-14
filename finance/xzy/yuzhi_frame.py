# -- coding: utf-8 --
# @File : yuzhi_frame.py
# @Description: 预支处理（手续费拆分）工具 - Frame嵌入版
# @Note: 原始逻辑从 shouxufei_cf.py 提取并重构，去掉sys.exit()和多窗口模式

import os
import pandas as pd
import tempfile
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font, Border, Side
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import queue
import threading
import shutil


class YuzhiCore:
    """预支处理核心逻辑（纯计算，无GUI依赖）"""

    def __init__(self, input_file, output_file=None):
        self.input_file = input_file
        if output_file is None:
            # 未指定输出文件时，生成系统临时文件，最终由 GUI 层另存到用户指定位置
            fd, output_file = tempfile.mkstemp(suffix='.xlsx', prefix='yuzhi_')
            os.close(fd)
        self.output_file = output_file

    def process(self, progress_callback=None):
        """
        执行完整的预支处理流程
        :param progress_callback: (percent, message) -> None 的回调函数
        """
        if progress_callback:
            progress_callback(5, "正在读取 Excel 文件...")

        df = pd.read_excel(self.input_file)

        if progress_callback:
            progress_callback(10, "正在处理列映射...")

        # 定义列字母到索引的映射
        col_map = {'A': 0, 'B': 1, 'C': 2, 'D': 3, 'E': 4, 'F': 5, 'G': 6,
                   'H': 7, 'I': 8, 'J': 9, 'K': 10, 'L': 11}
        col_name = {letter: df.columns[idx] for letter, idx in col_map.items()}

        # 对 L 列去除所有空格（先将NaN填充为空字符串，避免astype(str)产生"nan"）
        if col_name['L'] in df.columns:
            df[col_name['L']] = df[col_name['L']].fillna('').astype(str).str.replace(' ', '')

        if progress_callback:
            progress_callback(15, "正在应用条件规则...")

        # 条件判断规则（向量化操作）
        mask1 = (df[col_name['H']].astype(str).str.contains('招商银行', na=False) &
                 ~df[col_name['C']].astype(str).str.contains('伟明', na=False))
        df.loc[mask1, col_name['E']] = 0

        mask2 = df[col_name['I']].astype(str).str.contains('驻厂', na=False)
        df.loc[mask2, col_name['E']] = 0

        mask3 = (df[col_name['C']].astype(str).str.contains('伟明', na=False) &
                 df[col_name['H']].astype(str).str.contains('中信银行', na=False))
        df.loc[mask3, col_name['E']] = 0

        mask4 = (df[col_name['C']].astype(str).str.contains('伟明', na=False) &
                 df[col_name['H']].astype(str).str.contains('招商银行', na=False) &
                 df[col_name['I']].astype(str).str.contains('垫付', na=False))
        df.loc[mask4, col_name['E']] = 0

        mask5 = (df[col_name['C']].astype(str).str.contains('瑞立', na=False) &
                 (df[col_name['H']].astype(str).str.contains('招商银行', na=False) |
                  df[col_name['H']].astype(str).str.contains('邮政', na=False)))
        df.loc[mask5, col_name['E']] = 0

        mask6 = df[col_name['C']].astype(str).str.contains('中世', na=False)
        df.loc[mask6, col_name['E']] = 0

        # ============ 新增规则 2.1：运营中心 = 温州长江汽车电子 ============
        # 民生银行不扣手续费；其余银行统一扣 5 元
        mask_wz_cj = df[col_name['C']].astype(str).str.contains('温州长江汽车电子', na=False)
        mask_wz_cj_cmb = mask_wz_cj & df[col_name['H']].astype(str).str.contains('民生银行', na=False)
        df.loc[mask_wz_cj_cmb, col_name['E']] = 0
        mask_wz_cj_other = mask_wz_cj & ~df[col_name['H']].astype(str).str.contains('民生银行', na=False)
        df.loc[mask_wz_cj_other, col_name['E']] = 5

        # ============ 新增规则 2.2：运营中心 = 中铭工程 ============
        # 招商银行 / 驻厂 / 垫付 不扣手续费
        # 预支金额 ≤ 1000 扣 10 元；> 1000 按 1% 扣（其余银行且不满足免扣条件时）
        mask_zm = df[col_name['C']].astype(str).str.contains('中铭工程', na=False)
        mask_zm_cmb = mask_zm & df[col_name['H']].astype(str).str.contains('招商银行', na=False)
        df.loc[mask_zm_cmb, col_name['E']] = 0
        mask_zm_df = mask_zm & df[col_name['I']].astype(str).str.contains('垫付', na=False)
        df.loc[mask_zm_df, col_name['E']] = 0
        # 计费规则仅作用于 E 列仍为空的行（即未命中免扣条件）
        mask_zm_le_1000 = mask_zm & (df[col_name['D']] <= 1000) & df[col_name['E']].isna()
        df.loc[mask_zm_le_1000, col_name['E']] = 10
        mask_zm_gt_1000 = mask_zm & (df[col_name['D']] > 1000) & df[col_name['E']].isna()
        df.loc[mask_zm_gt_1000, col_name['E']] = df.loc[mask_zm_gt_1000, col_name['D']] * 0.01

        if progress_callback:
            progress_callback(35, "正在计算 E 列数值...")

        # 计算 E 列
        mask_empty_e = df[col_name['E']].isna()
        df.loc[mask_empty_e & (df[col_name['D']] <= 500), col_name['E']] = 5
        df.loc[mask_empty_e & (df[col_name['D']] > 500), col_name['E']] = (
            df.loc[mask_empty_e, col_name['D']] * 0.01
        )

        if progress_callback:
            progress_callback(45, "正在计算 F 列数值...")

        # F列 = D列 - E列
        df[col_name['F']] = df[col_name['D']] - df[col_name['E']]

        if progress_callback:
            progress_callback(50, "正在保存文件...")

        # 保存中间结果
        df.to_excel(self.output_file, index=False, engine='openpyxl')

        # 添加公式到 F 列
        wb_calc = load_workbook(self.output_file)
        ws_calc = wb_calc.active
        last_data_row = len(df) + 1
        for row in range(2, last_data_row + 1):
            ws_calc[f'F{row}'] = f'=D{row}-E{row}'
        wb_calc.save(self.output_file)

        if progress_callback:
            progress_callback(60, "正在拆分工作表...")

        # 根据 G 列内容读取实际 F 列值
        wb_temp = load_workbook(self.output_file, data_only=True)
        ws_temp = wb_temp.active

        df_for_split = df.copy()
        for idx in range(len(df_for_split)):
            row_num = idx + 2
            f_value = ws_temp[f'F{row_num}'].value
            if f_value is not None:
                try:
                    df_for_split.at[idx, col_name['F']] = float(f_value)
                except (ValueError, TypeError):
                    d_val = df_for_split.at[idx, col_name['D']]
                    e_val = df_for_split.at[idx, col_name['E']]
                    if pd.notna(d_val) and pd.notna(e_val):
                        try:
                            df_for_split.at[idx, col_name['F']] = float(d_val) - float(e_val)
                        except (ValueError, TypeError):
                            df_for_split.at[idx, col_name['F']] = 0
                    else:
                        df_for_split.at[idx, col_name['F']] = 0
            else:
                d_val = df_for_split.at[idx, col_name['D']]
                e_val = df_for_split.at[idx, col_name['E']]
                if pd.notna(d_val) and pd.notna(e_val):
                    try:
                        df_for_split.at[idx, col_name['F']] = float(d_val) - float(e_val)
                    except (ValueError, TypeError):
                        df_for_split.at[idx, col_name['F']] = 0
                else:
                    df_for_split.at[idx, col_name['F']] = 0

        wb_temp.close()

        # 按 G 列拆分为多个工作表
        with pd.ExcelWriter(self.output_file, engine='openpyxl', mode='a',
                            if_sheet_exists='replace') as writer:
            unique_values = df_for_split[df_for_split.columns[col_map['G']]].unique()
            for i, value in enumerate(unique_values):
                if pd.isna(value):
                    continue
                filtered_df = df_for_split[df_for_split[df_for_split.columns[col_map['G']]] == value]
                filtered_df.to_excel(writer, sheet_name=str(value), index=False)

                if progress_callback:
                    progress = 60 + (i / len(unique_values) * 10)
                    progress_callback(progress, f"正在拆分工作表：{value}")

        if progress_callback:
            progress_callback(70, "正在加载工作簿并格式化...")

        # 加载修改后的文件进行格式化
        wb = load_workbook(self.output_file)

        yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        duplicate_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        total_sheets = len(wb.sheetnames)
        # 用于清除表头加粗/边框的样式
        normal_font = Font(bold=False)
        no_border = Border(
            left=Side(style='none'),
            right=Side(style='none'),
            top=Side(style='none'),
            bottom=Side(style='none')
        )
        for sheet_idx, sheet_name in enumerate(wb.sheetnames):
            ws = wb[sheet_name]

            # ★ 清除表头行（第1行）的加粗和边框样式
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=1, column=col)
                cell.font = normal_font
                cell.border = no_border

            # 找到 C 列最后一个有内容的行，添加合计行
            last_row = ws.max_row
            total_row = None
            for row in range(1, last_row + 2):
                if ws[f'C{row}'].value is None:
                    ws[f'C{row}'] = '合计'
                    ws[f'D{row}'] = f'=SUM(D2:D{row - 1})'
                    ws[f'E{row}'] = f'=SUM(E2:E{row - 1})'
                    ws[f'F{row}'] = f'=SUM(F2:F{row - 1})'
                    ws[f'E{row}'].fill = yellow_fill
                    ws[f'F{row}'].fill = yellow_fill
                    total_row = row
                    break

            # 标记特殊条件的行
            for row in range(1, ws.max_row + 1):
                if ws[f'I{row}'].value and '驻厂' in str(ws[f'I{row}'].value):
                    for col in ['D', 'E', 'F', 'G', 'H', 'I']:
                        ws[f'{col}{row}'].fill = yellow_fill

                if ws[f'H{row}'].value and '招商银行' in str(ws[f'H{row}'].value):
                    for col in ['D', 'E', 'F', 'G', 'H']:
                        ws[f'{col}{row}'].fill = yellow_fill

                c_val = ws[f'C{row}'].value
                h_val = ws[f'H{row}'].value
                i_val = ws[f'I{row}'].value
                if c_val and '伟明' in str(c_val) and h_val and '中信银行' in str(h_val):
                    for col in ['D', 'E', 'F', 'G', 'H']:
                        ws[f'{col}{row}'].fill = yellow_fill

                if c_val and '瑞立' in str(c_val) and h_val and (
                        '招商银行' in str(h_val) or '邮政' in str(h_val)):
                    for col in ['D', 'E', 'F', 'G', 'H']:
                        ws[f'{col}{row}'].fill = yellow_fill

                if c_val and '中世' in str(c_val):
                    for col in ['D', 'E', 'F', 'G', 'H']:
                        ws[f'{col}{row}'].fill = yellow_fill

                # ★ 新增：温州长江汽车电子 + 民生银行（规则 2.1 免扣）
                if c_val and '温州长江汽车电子' in str(c_val) and h_val and '民生银行' in str(h_val):
                    for col in ['D', 'E', 'F', 'G', 'H']:
                        ws[f'{col}{row}'].fill = yellow_fill

                # ★ 新增：中铭工程 免扣条件（规则 2.2）
                if c_val and '中铭工程' in str(c_val) and h_val and '招商银行' in str(h_val):
                    for col in ['D', 'E', 'F', 'G', 'H']:
                        ws[f'{col}{row}'].fill = yellow_fill
                if c_val and '中铭工程' in str(c_val) and i_val and '垫付' in str(i_val):
                    for col in ['D', 'E', 'F', 'G', 'H', 'I']:
                        ws[f'{col}{row}'].fill = yellow_fill

            # 标记 J 列和 K 列的重复值
            j_values = [ws[f'J{row}'].value for row in range(1, ws.max_row + 1)]
            k_values = [ws[f'K{row}'].value for row in range(1, ws.max_row + 1)]
            j_duplicates = set([x for x in j_values if j_values.count(x) > 1])
            k_duplicates = set([x for x in k_values if k_values.count(x) > 1])

            for row in range(1, ws.max_row + 1):
                if ws[f'J{row}'].value in j_duplicates:
                    ws[f'J{row}'].fill = duplicate_fill
                if ws[f'K{row}'].value in k_duplicates:
                    ws[f'K{row}'].fill = duplicate_fill

            # 添加账号、户名、金额、汇款备注信息
            start_row = total_row + 5
            ws[f'C{start_row}'] = '账号'
            ws[f'D{start_row}'] = '户名'
            ws[f'E{start_row}'] = '金额'
            ws[f'F{start_row}'] = '汇款备注'

            data_row = start_row + 1
            amount_sum = 0

            for row_idx in range(2, total_row):
                bank_card = ws[f'L{row_idx}'].value
                account_name = ws[f'K{row_idx}'].value

                # 将所有字段的 None/nan 统一转换为空字符串，避免写入 "nan"
                if bank_card is None or (isinstance(bank_card, float) and __import__('math').isnan(bank_card)):
                    bank_card = ''
                if account_name is None or (isinstance(account_name, float) and __import__('math').isnan(account_name)):
                    account_name = ''

                f_cell = ws[f'F{row_idx}']
                if f_cell.value is not None:
                    if isinstance(f_cell.value, str) and f_cell.value.startswith('='):
                        d_value = ws[f'D{row_idx}'].value
                        e_value = ws[f'E{row_idx}'].value
                        if d_value is not None and e_value is not None:
                            try:
                                actual_amount = float(d_value) - float(e_value)
                            except (ValueError, TypeError):
                                actual_amount = 0
                        else:
                            actual_amount = 0
                    else:
                        try:
                            actual_amount = float(f_cell.value)
                        except (ValueError, TypeError):
                            actual_amount = 0
                else:
                    actual_amount = 0

                name = ws[f'J{row_idx}'].value
                # 汇款备注(name)也需要处理 nan
                if name is None or (isinstance(name, float) and __import__('math').isnan(name)):
                    name = ''

                ws[f'C{data_row}'] = bank_card
                ws[f'D{data_row}'] = account_name
                ws[f'E{data_row}'] = actual_amount
                ws[f'F{data_row}'] = name

                if actual_amount is not None:
                    try:
                        amount_sum += float(actual_amount)
                    except (ValueError, TypeError):
                        pass

                data_row += 1

            # 金额合计行
            amount_total_row = data_row
            ws[f'D{amount_total_row}'] = '金额合计'
            ws[f'E{amount_total_row}'] = amount_sum
            ws[f'E{amount_total_row}'].fill = yellow_fill

            if progress_callback:
                progress = 75 + ((sheet_idx + 1) / total_sheets * 20)
                progress_callback(progress, f"正在处理工作表：{sheet_name}")

        if progress_callback:
            progress_callback(95, "正在保存最终文件...")

        wb.save(self.output_file)

        if progress_callback:
            progress_callback(100, "处理完成！")


class YuzhiFrame(ttk.Frame):
    """预支处理功能 - 作为Frame嵌入主界面"""

    def __init__(self, parent, app):
        super().__init__(parent)
        self.app = app
        self._build_ui()

    def _get_toplevel(self):
        return self.winfo_toplevel()

    def _build_ui(self):
        main_frame = ttk.Frame(self, padding=25)
        main_frame.pack(fill=tk.BOTH, expand=True)

        # 标题
        ttk.Label(
            main_frame, text="手续费拆分处理",
            font=("Microsoft YaHei", 16, "bold")
        ).pack(pady=(0, 15))

        # 说明区域
        desc_frame = ttk.LabelFrame(main_frame, text="功能说明", padding=12)
        desc_frame.pack(fill=tk.X, pady=(0, 15))

        desc_text = (
            "本功能对Excel文件执行以下操作：\n"
            "1. 根据特定规则自动计算手续费（E列）\n"
            "2. 计算实发金额（F列 = D列 - E列）\n"
            "3. 按指定列（G列）拆分数据为多个工作表\n"
            "4. 自动标记特殊条件行、重复值\n"
            "5. 在每个工作表中添加汇总信息"
        )
        ttk.Label(desc_frame, text=desc_text, font=("Microsoft YaHei", 10),
                  justify="left").pack(anchor="w")

        # 文件选择区域
        file_frame = ttk.LabelFrame(main_frame, text="文件选择", padding=12)
        file_frame.pack(fill=tk.X, pady=(0, 15))

        # 输入文件选择
        row_in = ttk.Frame(file_frame)
        row_in.pack(fill=tk.X, pady=5)
        ttk.Label(row_in, text='选择"工人预支申请"表：', font=("Microsoft YaHei", 10)).pack(side=tk.LEFT)
        self.lbl_input = ttk.Label(row_in, text="未选择", font=("Microsoft YaHei", 9),
                                    foreground="gray")
        self.lbl_input.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)
        self.btn_select_input = ttk.Button(row_in, text="选择", command=self.select_input_file, width=8)
        self.btn_select_input.pack(side=tk.RIGHT)
        self.input_file_path = None

        # 处理按钮
        btn_area = ttk.Frame(main_frame)
        btn_area.pack(pady=10)

        self.btn_process = ttk.Button(
            btn_area, text="开始处理", command=self.start_process, width=18
        )
        self.btn_process.pack()
        self.btn_process.config(state=tk.DISABLED)

        # 状态信息区域
        log_frame = ttk.LabelFrame(main_frame, text="状态日志", padding=10)
        log_frame.pack(fill=tk.BOTH, expand=True, pady=5)

        self.log_text = tk.Text(log_frame, height=10, wrap=tk.WORD, font=("Consolas", 10))
        vsb = ttk.Scrollbar(log_frame, orient=tk.VERTICAL, command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=vsb.set)
        self.log_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        vsb.pack(side=tk.RIGHT, fill=tk.Y)

        self.log("就绪。请选择输入文件后点击'开始处理'")

    def log(self, msg):
        self.log_text.insert(tk.END, msg + "\n")
        self.log_text.see(tk.END)
        print(msg)

    def select_input_file(self):
        filepath = filedialog.askopenfilename(
            title="选择要处理的Excel文件",
            filetypes=[("Excel 文件", "*.xlsx"), ("所有文件", "*.*")]
        )
        if not filepath:
            return
        self.input_file_path = filepath
        filename = filepath.split("/")[-1].split("\\")[-1]
        self.lbl_input.config(text=filename, foreground="black")
        self.log(f"已选择输入文件：{filepath}")
        self._check_ready()

    def _check_ready(self):
        if self.input_file_path:
            self.btn_process.config(state=tk.NORMAL)
        else:
            self.btn_process.config(state=tk.DISABLED)

    def start_process(self):
        """启动处理（使用进度弹窗 + 多线程）"""
        self.btn_process.config(state=tk.DISABLED)
        self.log_text.delete(1.0, tk.END)
        self.log("开始处理...")

        # 创建进度弹窗
        progress_window = tk.Toplevel(self._get_toplevel())
        progress_window.title("处理进度")
        progress_window.geometry("400x150")
        progress_window.transient(self._get_toplevel())
        progress_window.resizable(False, False)

        # 居中
        progress_window.update_idletasks()
        w = progress_window.winfo_width()
        h = progress_window.winfo_height()
        x = self._get_toplevel().winfo_rootx() + (self._get_toplevel().winfo_width() - w) // 2
        y = self._get_toplevel().winfo_rooty() + (self._get_toplevel().winfo_height() - h) // 2
        progress_window.geometry(f'+{x}+{y}')

        label = ttk.Label(progress_window, text="正在处理数据，请稍候...",
                           font=("Arial", 12))
        label.pack(pady=20)

        progress_bar = ttk.Progressbar(progress_window, length=300, mode='determinate')
        progress_bar.pack(pady=10)

        progress_label = ttk.Label(progress_window, text="0%", font=("Arial", 10))
        progress_label.pack(pady=5)

        # 队列用于线程通信
        q = queue.Queue()

        # 结果容器
        result_container = {'error': None, 'output_file': None}

        def update_from_queue():
            try:
                while True:
                    try:
                        pct, msg = q.get_nowait()
                    except queue.Empty:
                        break
                    progress_bar['value'] = pct
                    progress_label.config(text=f"{pct:.0f}%")
                    if msg:
                        label.config(text=msg)
                        self.log(msg)

                    if pct >= 100:
                        progress_window.after(800, lambda: finish_success(progress_window))
            finally:
                progress_window.after(100, update_from_queue)

        def _cleanup_temp(path):
            if path and os.path.exists(path):
                try:
                    os.remove(path)
                except Exception:
                    pass

        def finish_success(window):
            window.destroy()
            # 处理完成，弹出保存对话框让用户自行选择保存位置
            desktop_path = os.path.join(os.environ.get("USERPROFILE", ""), "Desktop")
            save_path = filedialog.asksaveasfilename(
                title="选择保存位置",
                defaultextension=".xlsx",
                filetypes=[("Excel 文件", "*.xlsx"), ("所有文件", "*.*")],
                initialdir=desktop_path,
                initialfile="预支处理结果.xlsx"
            )
            if not save_path:
                # 用户取消保存，清理临时文件
                _cleanup_temp(result_container.get('output_file'))
                messagebox.showinfo("已取消", "未选择保存位置，处理结果已丢弃")
                self.btn_process.config(state=tk.NORMAL)
                self.app.status_label.config(text="手续费预支处理 - 已取消")
                return
            try:
                shutil.copy2(result_container['output_file'], save_path)
                _cleanup_temp(result_container.get('output_file'))
                messagebox.showinfo("完成", f"文件已保存到：\n{save_path}")
                self.btn_process.config(state=tk.NORMAL)
                self.app.status_label.config(text="手续费预支处理 - 完成")
            except Exception as e:
                import traceback
                err = f"保存文件失败：{str(e)}\n{traceback.format_exc()}"
                self.log(f"[错误] {err}")
                messagebox.showerror("保存失败", str(e))
                self.btn_process.config(state=tk.NORMAL)
                self.app.status_label.config(text="手续费预支处理 - 保存失败")

        def finish_error(window, error_msg):
            window.destroy()
            messagebox.showerror("错误", error_msg)
            self.btn_process.config(state=tk.NORMAL)
            self.app.status_label.config(text="手续费预支处理 - 出错")
            self.log(f"[错误] {error_msg}")
    
        # 后台处理函数
        def process_thread():
            try:
                core = YuzhiCore(self.input_file_path)
                core.process(lambda pct, msg: q.put((pct, msg)))
                result_container['output_file'] = core.output_file
            except Exception as e:
                import traceback
                err = f"处理出错：{str(e)}\n{traceback.format_exc()}"
                result_container['error'] = err
                q.put((0, ""))
                progress_window.after(0, lambda: finish_error(progress_window, err))

        # 启动
        progress_window.after(100, update_from_queue)
        thread = threading.Thread(target=process_thread, daemon=True)
        thread.start()
