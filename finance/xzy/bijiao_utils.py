# -- coding: utf-8 --
# @File : bijiao_utils.py
# @Description: 数据比较工具 - 工具函数

from openpyxl import load_workbook


def get_sheet_names(file_path):
    """获取Excel文件的所有工作表名称"""
    wb = load_workbook(file_path, read_only=True)
    sheet_names = wb.sheetnames
    wb.close()
    return sheet_names
