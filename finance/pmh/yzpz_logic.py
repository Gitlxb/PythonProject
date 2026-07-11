from openpyxl.styles import Font
from datetime import datetime
from finance.pmh.yzpz_core import ExcelProcessor

# ==================== 扣款处理方法 ====================

def process_deduction_data(self, source_worksheet, source_sheet_name, selected_column_type=None):
    """处理需扣回和预支扣款数据并复制到 Sheet4（支持三列：需扣回、预支扣款、预支扣回）

    Args:
        source_worksheet: 源工作表
        source_sheet_name: 源工作表名称
        selected_column_type: 用户选择的列类型 ('需扣回', '预支扣款', '预支扣回')，None表示自动选择
    """
    try:
        if not self.workbook:
            return False, "请先选择主 Excel 文件"

        name_col = self._find_column_by_header(source_worksheet, "姓名", cache_key=f"{source_sheet_name}_姓名")
        deduction_col = self._find_column_contains(source_worksheet, "需扣回",
                                                   cache_key=f"{source_sheet_name}_需扣回")
        advance_deduction_col = self._find_column_contains(source_worksheet, "预支扣款",
                                                           cache_key=f"{source_sheet_name}_预支扣款")
        advance_return_col = self._find_column_contains(source_worksheet, "预支扣回",
                                                        cache_key=f"{source_sheet_name}_预支扣回")

        if not name_col:
            return False, "未找到'姓名'列"
        if not deduction_col and not advance_deduction_col and not advance_return_col:
            return False, "未找到'需扣回'列、'预支扣款'列或'预支扣回'列"

        available_columns = []
        if deduction_col:
            for row in range(3, source_worksheet.max_row + 1):
                val = source_worksheet.cell(row=row, column=deduction_col).value
                if val is not None and str(val).strip() != '':
                    available_columns.append(('需扣回', deduction_col))
                    break

        if advance_deduction_col:
            for row in range(3, source_worksheet.max_row + 1):
                val = source_worksheet.cell(row=row, column=advance_deduction_col).value
                if val is not None and str(val).strip() != '':
                    available_columns.append(('预支扣款', advance_deduction_col))
                    break

        if advance_return_col:
            for row in range(3, source_worksheet.max_row + 1):
                val = source_worksheet.cell(row=row, column=advance_return_col).value
                if val is not None and str(val).strip() != '':
                    available_columns.append(('预支扣回', advance_return_col))
                    break

        if not available_columns:
            return False, "未找到有效的需扣回或预支扣款数据"

        if len(available_columns) > 1 and selected_column_type:
            target_col = None
            for col_name, col_idx in available_columns:
                if col_name == selected_column_type:
                    target_col = col_idx
                    break
            if not target_col:
                return False, f"未找到'{selected_column_type}'列"
        else:
            target_col = available_columns[0][1]

        deduction_rows = []
        for row in range(3, source_worksheet.max_row + 1):
            name_value = source_worksheet.cell(row=row, column=name_col).value
            if not self._is_valid_name(name_value):
                continue
            deduction_value = source_worksheet.cell(row=row, column=target_col).value
            if deduction_value is not None and str(deduction_value).strip() != '':
                deduction_rows.append({
                    'name': name_value,
                    'deduction': deduction_value
                })

        if not deduction_rows:
            return False, "未找到有效的需扣回或预支扣款数据"

        names_num, data_num = self._get_round_sheet_numbers()
        data_sheet_name = f'Sheet{data_num}'

        if data_sheet_name not in self.workbook.sheetnames:
            return False, f"{data_sheet_name} 不存在，请先完成第一步"

        sheet4 = self.workbook[data_sheet_name]
        sheet4.cell(row=1, column=52).value = selected_column_type if selected_column_type else "预支扣款"

        DEDUCTION_NAME_COL = 5
        DEDUCTION_VALUE_COL = 6
        ADVANCE_NAME_COL = 9
        ADVANCE_VALUE_COL = 10
        COMPARE_COL = 11
        SUM_COL = 12
        DEDUCTION_SUM_NAME_COL = 13
        DEDUCTION_SUM_VALUE_COL = 14

        sheet4.cell(row=1, column=DEDUCTION_NAME_COL).value = "姓名"
        column_header = selected_column_type if selected_column_type else "需扣回/预支扣款"
        sheet4.cell(row=1, column=DEDUCTION_VALUE_COL).value = column_header

        for idx, data in enumerate(deduction_rows, start=2):
            sheet4.cell(row=idx, column=DEDUCTION_NAME_COL).value = data['name']
            sheet4.cell(row=idx, column=DEDUCTION_VALUE_COL).value = data['deduction']

        def copy_columns_with_total(src_name_col, src_value_col, dest_name_col, dest_value_col, title_name,
                                    title_value):
            """通用方法：复制两列数据并添加总计"""
            sheet4.cell(row=1, column=dest_name_col).value = title_name
            sheet4.cell(row=1, column=dest_value_col).value = title_value

            max_row_idx = 1
            total_value = 0

            for row in range(2, sheet4.max_row + 1):
                name_value = sheet4.cell(row=row, column=src_name_col).value
                value = sheet4.cell(row=row, column=src_value_col).value

                if name_value and str(name_value).strip() != '总计':
                    sheet4.cell(row=row, column=dest_name_col).value = name_value
                    sheet4.cell(row=row, column=dest_value_col).value = value
                    max_row_idx = row

                    if value is not None:
                        try:
                            total_value += float(value)
                        except (ValueError, TypeError):
                            pass

            return max_row_idx, total_value

        name_advance_dict = {}
        for row in range(2, sheet4.max_row + 1):
            name_value = sheet4.cell(row=row, column=1).value
            if not self._is_valid_name(name_value):
                continue
            name_str = str(name_value).strip()
            name_advance_dict[name_str] = name_advance_dict.get(name_str, 0) + self._to_float(
                sheet4.cell(row=row, column=2).value)

        sheet4.cell(row=1, column=DEDUCTION_SUM_NAME_COL).value = "姓名"
        sheet4.cell(row=1, column=DEDUCTION_SUM_VALUE_COL).value = "求和项：需扣回/预支扣款"
        max_row_13_14 = 1
        total_deduction = 0
        for row in range(2, sheet4.max_row + 1):
            name_value = sheet4.cell(row=row, column=5).value
            value = sheet4.cell(row=row, column=6).value
            if self._is_valid_name(name_value):
                sheet4.cell(row=row, column=DEDUCTION_SUM_NAME_COL).value = name_value
                sheet4.cell(row=row, column=DEDUCTION_SUM_VALUE_COL).value = value
                max_row_13_14 = row
                total_deduction += self._to_float(value)

        sheet4.cell(row=1, column=ADVANCE_NAME_COL).value = "姓名"
        sheet4.cell(row=1, column=ADVANCE_VALUE_COL).value = "求和项：预支数额"

        max_row_9_10 = 1
        total_advance = 0

        for row in range(2, sheet4.max_row + 1):
            name_13 = sheet4.cell(row=row, column=DEDUCTION_SUM_NAME_COL).value
            if not self._is_valid_name(name_13):
                continue
            name_str = str(name_13).strip()
            sheet4.cell(row=row, column=ADVANCE_NAME_COL).value = name_13
            advance_total = name_advance_dict.get(name_str, 0)
            sheet4.cell(row=row, column=ADVANCE_VALUE_COL).value = advance_total
            max_row_9_10 = row
            total_advance += advance_total

        max_data_row = max(max_row_9_10, max_row_13_14)
        total_row = max_data_row + 1

        if max_data_row > 1:
            sheet4.cell(row=total_row, column=ADVANCE_NAME_COL).value = "总计"
            sheet4.cell(row=total_row, column=ADVANCE_VALUE_COL).value = total_advance
            sheet4.cell(row=total_row, column=DEDUCTION_SUM_NAME_COL).value = "总计"
            sheet4.cell(row=total_row, column=DEDUCTION_SUM_VALUE_COL).value = total_deduction

        for row in range(2, total_row):
            name_9 = sheet4.cell(row=row, column=ADVANCE_NAME_COL).value
            name_13 = sheet4.cell(row=row, column=DEDUCTION_SUM_NAME_COL).value
            sheet4.cell(row=row, column=COMPARE_COL).value = (name_9 == name_13)

        for row in range(2, total_row):
            val_10 = sheet4.cell(row=row, column=ADVANCE_VALUE_COL).value
            val_14 = sheet4.cell(row=row, column=DEDUCTION_SUM_VALUE_COL).value

            sum_value = 0
            if val_10 is not None:
                try:
                    sum_value += float(val_10)
                except (ValueError, TypeError):
                    pass
            if val_14 is not None:
                try:
                    sum_value += float(val_14)
                except (ValueError, TypeError):
                    pass

            sheet4.cell(row=row, column=SUM_COL).value = sum_value

        return True, f"已处理 {len(deduction_rows)} 条需扣回/预支扣款数据到 {data_sheet_name}"

    except Exception as e:
        return False, f"处理需扣回数据失败：{str(e)}"


# ==================== 平账处理方法 ====================

def balance_accounts(self, salary_sheet_name, month, selected_units=None):
    """平账功能（批量优化）"""
    try:
        if not self.workbook:
            return False, "请先选择 Excel 文件"

        salary_sheet = self.workbook[salary_sheet_name]

        names_num, data_num = self._get_round_sheet_numbers()
        data_sheet_name = f'Sheet{data_num}'

        if data_sheet_name not in self.workbook.sheetnames:
            return False, f"{data_sheet_name} 不存在，请先完成核对"

        sheet4 = self.workbook[data_sheet_name]

        name_col = self._find_column_by_header(
            salary_sheet, "姓名", cache_key=f"{salary_sheet_name}_姓名")
        unit_col = self._find_column_by_header(
            salary_sheet, "单位", cache_key=f"{salary_sheet_name}_单位")
        advance_col = self._find_column_contains(
            salary_sheet, "预支数额", cache_key=f"{salary_sheet_name}_预支数额")
        balance_col = self._find_column_contains(
            salary_sheet, "平账额", cache_key=f"{salary_sheet_name}_平账额")

        settled_col = self._find_column_contains(
            salary_sheet, "是否平账", cache_key=f"{salary_sheet_name}_平账标记1")
        if not settled_col:
            settled_col = self._find_column_contains(
                salary_sheet, "工资表是否平账", cache_key=f"{salary_sheet_name}_平账标记2")

        month_col = self._find_column_contains(
            salary_sheet, "工资表所属月份", cache_key=f"{salary_sheet_name}_月份1")
        if not month_col:
            month_col = self._find_column_contains(
                salary_sheet, "工资表月份", cache_key=f"{salary_sheet_name}_月份2")

        if not month:
            month = "未知月份"

        marker_col = 52
        column_type_marker = sheet4.cell(row=1, column=marker_col).value

        if column_type_marker == "需扣回":
            source_name_col = 5
            source_value_col = 6
            target_name_col = 1
            target_advance_col = 2
            use_matching = True
        else:
            source_name_col = 1
            source_value_col = 2
            target_name_col = 1
            target_advance_col = 2
            use_matching = False

        SUM_COL = 12
        DEDUCTION_SUM_NAME_COL = 13
        DEDUCTION_SUM_VALUE_COL = 14
        split_data = {}

        for row in range(2, sheet4.max_row + 1):
            name_13 = sheet4.cell(row=row, column=DEDUCTION_SUM_NAME_COL).value
            sum_value = sheet4.cell(row=row, column=SUM_COL).value
            deduction_value = sheet4.cell(row=row, column=DEDUCTION_SUM_VALUE_COL).value

            if self._is_valid_name(name_13) and sum_value is not None and deduction_value is not None:
                sum_num = self._to_float(sum_value)
                ded_num = self._to_float(deduction_value)
                if sum_num > 0:
                    split_data[str(name_13).strip()] = abs(ded_num)

        sheet4_data = {}

        if use_matching:
            deduction_names = set()
            for row in range(2, sheet4.max_row + 1):
                name_5 = sheet4.cell(row=row, column=source_name_col).value
                if self._is_valid_name(name_5):
                    deduction_names.add(str(name_5).strip())

            for row in range(2, sheet4.max_row + 1):
                name = sheet4.cell(row=row, column=target_name_col).value
                advance = sheet4.cell(row=row, column=target_advance_col).value
                if self._is_valid_name(name) and advance is not None:
                    name_str = str(name).strip()
                    if name_str in deduction_names:
                        sheet4_data.setdefault(name_str, []).append(
                            {"advance": advance, "row": row, "used": False})
        else:
            for row in range(2, sheet4.max_row + 1):
                name = sheet4.cell(row=row, column=source_name_col).value
                advance = sheet4.cell(row=row, column=source_value_col).value
                if self._is_valid_name(name) and advance is not None:
                    name_str = str(name).strip()
                    sheet4_data.setdefault(name_str, []).append(
                        {"advance": advance, "row": row, "used": False})

        salary_rows = []
        matched_count = 0
        for row in range(3, salary_sheet.max_row + 1):
            name_value = salary_sheet.cell(row=row, column=name_col).value
            balance_value = salary_sheet.cell(row=row, column=balance_col).value

            if balance_value is not None and str(balance_value).strip() != '':
                continue

            if selected_units and unit_col:
                row_unit = salary_sheet.cell(row=row, column=unit_col).value
                if not row_unit or str(row_unit).strip() not in selected_units:
                    continue

            salary_rows.append({
                'row': row,
                'name': name_value,
                'advance': salary_sheet.cell(row=row, column=advance_col).value
            })
            if name_value and str(name_value).strip() in sheet4_data:
                matched_count += 1

        updated_count = 0
        now = datetime.now()
        today_date = f"{now.year}.{now.month}.{now.day}"
        updated_rows = []
        new_rows_to_add = []
        no_balance_rows = set()

        for name_str, need_balance_amount in split_data.items():
            if name_str not in sheet4_data or need_balance_amount <= 0:
                continue

            matching_salary_rows = [r for r in salary_rows if r['name'] and str(r['name']).strip() == name_str]
            if not matching_salary_rows:
                continue

            matching_salary_rows.sort(key=lambda x: x['row'])
            remaining_balance = need_balance_amount

            for salary_record in matching_salary_rows:
                if remaining_balance <= 0:
                    no_balance_rows.add(salary_record['row'])
                    continue

                original_advance = self._to_float(salary_record['advance'])
                if original_advance <= 0:
                    continue

                if original_advance <= remaining_balance:
                    remaining_balance -= original_advance
                else:
                    balance_part = remaining_balance
                    no_balance_part = original_advance - remaining_balance
                    salary_sheet.cell(row=salary_record['row'], column=advance_col).value = balance_part
                    new_rows_to_add.append({
                        'source_row': salary_record['row'],
                        'advance': no_balance_part,
                        'is_split': True,
                        'should_balance': False
                    })
                    remaining_balance = 0

        for data in salary_rows:
            if data['row'] in no_balance_rows:
                continue

            name_value = data['name']
            if not name_value:
                continue
            name_str = str(name_value).strip()
            if name_str not in sheet4_data:
                continue

            salary_advance = self._to_float(data['advance'], None)
            if salary_advance is None:
                continue

            for record in sheet4_data[name_str]:
                if record["used"]:
                    continue

                sheet4_advance = self._to_float(record["advance"], None)
                if sheet4_advance is None:
                    continue

                if abs(salary_advance - sheet4_advance) < 0.01:
                    record["used"] = True
                    real_advance = salary_sheet.cell(row=data['row'], column=advance_col).value
                    salary_sheet.cell(row=data['row'], column=balance_col).value = real_advance
                    if settled_col:
                        salary_sheet.cell(row=data['row'], column=settled_col).value = today_date
                    if month_col:
                        salary_sheet.cell(row=data['row'], column=month_col).value = month
                    updated_count += 1
                    updated_rows.append(data['row'])
                    break

        if new_rows_to_add:
            rows_to_append = []

            for new_row_data in new_rows_to_add:
                if new_row_data.get('keep_original', False):
                    continue

                source_row = new_row_data['source_row']
                new_advance = new_row_data['advance']
                should_balance = new_row_data.get('should_balance', False)

                row_data = {}
                for col in range(1, salary_sheet.max_column + 1):
                    cell_value = salary_sheet.cell(row=source_row, column=col).value
                    row_data[col] = cell_value

                row_data[advance_col] = new_advance

                source_name = row_data.get(name_col)

                if not should_balance:
                    if balance_col:
                        row_data[balance_col] = None
                    if settled_col:
                        row_data[settled_col] = None
                    if month_col:
                        row_data[month_col] = None

                rows_to_append.append({
                    'data': row_data,
                    'source_row': source_row,
                    'should_balance': should_balance,
                    'name': source_name
                })

            raw_max = salary_sheet.max_row
            actual_last_row = raw_max
            for check_row in range(raw_max, 0, -1):
                row_has_data = False
                for col in range(1, min(salary_sheet.max_column + 1, 20)):
                    val = salary_sheet.cell(row=check_row, column=col).value
                    if val is not None and str(val).strip() != '':
                        row_has_data = True
                        break
                if row_has_data:
                    actual_last_row = check_row
                    break

            for idx, row_info in enumerate(rows_to_append, start=1):
                target_row = actual_last_row + idx
                row_data = row_info['data']
                source_row = row_info['source_row']

                for col, value in row_data.items():
                    target_cell = salary_sheet.cell(row=target_row, column=col)
                    target_cell.value = value
                    # 复制源行的对齐样式到新增行（安全方式，避免 StyleProxy 哈希错误）
                    try:
                        source_cell = salary_sheet.cell(row=source_row, column=col)
                        align = source_cell.alignment
                        if align and align.horizontal:
                            target_cell.alignment = align
                    except (TypeError, ValueError):
                        pass

        if updated_rows:
            red_font = Font(color="FF0000")
            for row_num in updated_rows:
                for col in range(1, salary_sheet.max_column + 1):
                    cell = salary_sheet.cell(row=row_num, column=col)
                    if cell.value is not None and str(cell.value).strip() != '':
                        cell.font = red_font

        return True, f"已平账{updated_count}条记录，月份：{month}"

    except Exception as e:
        return False, f"平账失败：{str(e)}"


# 将方法绑定到 ExcelProcessor 类
ExcelProcessor.process_deduction_data = process_deduction_data
ExcelProcessor.balance_accounts = balance_accounts
