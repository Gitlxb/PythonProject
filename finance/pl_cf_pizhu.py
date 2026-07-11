# -- coding: utf-8 --
# @Time    : 2025-03-22
# @Author  : 贝特利
# @File    : pl_cf_pizhu.py
# @Desc    : 批量批注工具（合二为一：批量添加批注 + 批量拆分批注）

import os
import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import load_workbook
from openpyxl.comments import Comment
from tkinter.filedialog import askopenfilename


class CommentToolApp:
    """批量批注工具主界面 - 合并 pl_pizhu + pl_cfpizhu"""

    def __init__(self, parent=None):
        if parent is None:
            self.root = tk.Tk()
            self.root.title("批量批注工具")
            self.root.geometry("520x380")
            self._build_ui()
            self.root.mainloop()
        else:
            self.root = tk.Toplevel(parent)
            self.root.title("批量批注工具")
            self.root.geometry("520x380")
            self.root.transient(parent)
            self.root.grab_set()
            self._build_ui()
            self.root.lift()
            self.root.focus_force()

    def _build_ui(self):
        self.root.configure(bg="#f0f0f0")

        # 标题
        title_frame = tk.Frame(self.root, bg="#2c3e50", height=50)
        title_frame.pack(fill="x")
        title_frame.pack_propagate(False)

        tk.Label(
            title_frame,
            text="📝 批量批注工具",
            font=("Microsoft YaHei", 16, "bold"),
            bg="#2c3e50",
            fg="white"
        ).pack(expand=True)

        # 主内容区
        main = tk.Frame(self.root, bg="#f0f0f0", padx=30, pady=20)
        main.pack(fill="both", expand=True)

        # 文件选择
        file_frame = tk.Frame(main, bg="#f0f0f0")
        file_frame.pack(fill="x", pady=(0, 15))

        tk.Label(
            file_frame, text="选择Excel文件:", font=("Microsoft YaHei", 10),
            bg="#f0f0f0"
        ).pack(anchor="w")

        path_frame = tk.Frame(file_frame, bg="#f0f0f0")
        path_frame.pack(fill="x", pady=5)

        self.file_path_var = tk.StringVar()
        tk.Entry(
            path_frame, textvariable=self.file_path_var, font=("Microsoft YaHei", 9),
            width=55, state="readonly"
        ).pack(side="left", fill="x", expand=True)

        tk.Button(
            path_frame, text="浏览...", font=("Microsoft YaHei", 9),
            command=self._select_file, bg="#3498db", fg="white",
            width=8, cursor="hand2"
        ).pack(side="right", padx=(10, 0))

        # 分隔线
        ttk.Separator(main, orient="horizontal").pack(fill="x", pady=10)

        # 功能说明
        info_frame = tk.LabelFrame(
            main, text=" 功能说明 ", font=("Microsoft YaHei", 9, "bold"),
            bg="#f0f0f0", fg="#333", padx=10, pady=8
        )
        info_frame.pack(fill="x", pady=(0, 15))

        info_text = (
            "• 【批量添加批注】：将B~K列的内容合并后，作为批注写入A列\n"
            "• 【批量拆分批注】：将A列的批注内容拆开，依次写回右侧单元格"
        )
        tk.Label(
            info_frame, text=info_text, font=("Microsoft YaHei", 9),
            bg="#f0f0f0", fg="#555", justify="left", anchor="w"
        ).pack(fill="x")

        # 按钮区
        btn_frame = tk.Frame(main, bg="#f0f0f0")
        btn_frame.pack(fill="x", pady=10)

        self.add_btn = tk.Button(
            btn_frame,
            text="✅  批量添加批注\n   (B~K列 → A列批注)",
            font=("Microsoft YaHei", 11, "bold"),
            bg="#27ae60", fg="white", activebackground="#2ecc71",
            width=22, height=2, cursor="hand2",
            command=self._add_comments
        )
        self.add_btn.pack(side="left", expand=True, padx=(0, 10))

        self.split_btn = tk.Button(
            btn_frame,
            text="📤  批量拆分批注\n   (A列批注 → 右侧列)",
            font=("Microsoft YaHei", 11, "bold"),
            bg="#e67e22", fg="white", activebackground="#f39c12",
            width=22, height=2, cursor="hand2",
            command=self._split_comments
        )
        self.split_btn.pack(side="right", expand=True, padx=(10, 0))

        # 状态栏
        self.status_var = tk.StringVar(value="就绪 - 请选择Excel文件")
        status = tk.Label(
            self.root, textvariable=self.status_var,
            font=("Microsoft YaHei", 9), bg="#ecf0f1", fg="#7f8c8d",
            anchor="w", padx=15, pady=6
        )
        status.pack(fill="x", side="bottom")

    def _select_file(self):
        path = askopenfilename(
            title="选择Excel文件",
            filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")]
        )
        if path:
            self.file_path_var.set(path)
            self.status_var.set(f"已选择: {os.path.basename(path)}")

    def _add_comments(self):
        """批量添加批注 (原 pl_pizhu.py)"""
        file_path = self.file_path_var.get()
        if not file_path:
            messagebox.showwarning("提示", "请先选择Excel文件！")
            return

        try:
            wb = load_workbook(filename=file_path)
            ws = wb.active
            count = 0

            for row in ws.iter_rows(min_col=1, max_col=1):
                a_cell = row[0]
                if a_cell.value:
                    merged_content = []
                    for col in range(2, 12):  # B~K列
                        cell_value = ws.cell(row=a_cell.row, column=col).value
                        if cell_value:
                            merged_content.append(str(cell_value))

                    comment_text = "\n".join(merged_content)
                    if a_cell.comment:
                        a_cell.comment = None
                    if comment_text:
                        a_cell.comment = Comment(comment_text, "BatchTool")
                        count += 1

            wb.save(file_path)
            self.status_var.set(f"完成：已为 {count} 个单元格添加批注")
            messagebox.showinfo("完成", f"批量批注已完成！\n共处理 {count} 个单元格。\n文件已保存到原路径。")

        except Exception as e:
            self.status_var.set("错误：批量添加批注失败")
            messagebox.showerror("错误", f"处理过程中发生错误:\n{e}")

    def _split_comments(self):
        """批量拆分批注 (原 pl_cfpizhu.py)"""
        file_path = self.file_path_var.get()
        if not file_path:
            messagebox.showwarning("提示", "请先选择Excel文件！")
            return

        try:
            wb = load_workbook(file_path)
            ws = wb.active
            count = 0

            for row in ws.iter_rows(min_col=1, max_col=1):
                for cell in row:
                    if cell.comment is not None:
                        comment_text = cell.comment.text
                        filtered_lines = [
                            line for line in comment_text.split('\n')
                            if not line.strip().startswith("Admin:")
                        ]
                        filtered_text = " ".join(filtered_lines)
                        split_values = " ".join(filtered_text.split()).split(" ")

                        for i, value in enumerate(split_values):
                            ws.cell(row=cell.row, column=cell.column + i + 1, value=value)
                        count += 1

            wb.save(file_path)
            self.status_var.set(f"完成：已拆分 {count} 个批注")
            messagebox.showinfo("完成", f"批量拆分批注已完成！\n共拆分 {count} 个批注。\n文件已保存到原路径。")

        except Exception as e:
            self.status_var.set("错误：批量拆分批注失败")
            messagebox.showerror("错误", f"处理过程中发生错误:\n{e}")


if __name__ == "__main__":
    CommentToolApp()
